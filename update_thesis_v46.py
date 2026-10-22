from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import json
import math

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\VTP")
SRC = ROOT / "Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_da_kiem_chung.docx"
OUT = ROOT / "Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_V46_bo_sung_20260930.docx"
EXPLAIN = ROOT / "Giai_thich_quy_trinh_V46_de_hieu_20260930.docx"
FIG_FLOW = ROOT / "V46_so_do_quy_trinh_de_hieu.png"
FIG_CUT = ROOT / "V46_minh_hoa_cat_dem.png"
FIG_MAP = ROOT / "V46_ban_do_38_o_hieu_chinh.png"
RESEARCH_NOTE = ROOT / "V46_Nghien_cuu_bo_sung_cho_do_an_20260930.md"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


VERIFY = read_json(ROOT / "V46_audit_20260930" / "V46_final_verification.json")
COND = read_json(ROOT / "V46_audit_20260930" / "V46_conditioning_summary.json")
SEED = read_json(ROOT / "V46_seed55_historical_validated_20260930" / "V46_seed55_validation_summary.json")
STABLE = read_json(ROOT / "V46_stable_core_20260930" / "V46_stable_core_summary.json")
TOPO = read_json(ROOT / "V46_topology_20260930" / "V46_topology_summary.json")
HYDRO = read_json(ROOT / "V46_hydro_evidence_20260930" / "V46_hydro_evidence_summary.json")


def find_para(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise ValueError(f"Paragraph not found: {startswith}")


def copy_run_format(src_run, dst_run):
    if src_run is not None and src_run._r.rPr is not None:
        dst_run._r.insert(0, deepcopy(src_run._r.rPr))


def replace_text(p, text: str, *, bold: bool | None = None, italic: bool | None = None):
    src_run = p.runs[0] if p.runs else None
    p.clear()
    r = p.add_run(text)
    copy_run_format(src_run, r)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    return p


def insert_before(anchor, text: str, like=None, *, bold: bool | None = None, italic: bool | None = None):
    p = anchor.insert_paragraph_before()
    src = like or anchor
    if src._p.pPr is not None:
        p._p.insert(0, deepcopy(src._p.pPr))
    r = p.add_run(text)
    src_run = src.runs[0] if src.runs else None
    copy_run_format(src_run, r)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    return p


def set_cell_shading(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, bold=False, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def insert_table_before(doc, anchor, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Chỉ tiêu", "Kết quả V46", "Ý nghĩa khoa học"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
    for a, b, c in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], a, size=10)
        set_cell_text(cells[1], b, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], c, size=10)
    widths = [Cm(4.0), Cm(3.5), Cm(8.0)]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
    anchor._p.addprevious(table._tbl)
    return table


def add_reference(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)


def build_figures():
    # 1) Simplified process flow
    fig, ax = plt.subplots(figsize=(11, 7))
    ax.axis("off")
    boxes = [
        (0.08, 0.80, "1. V41\nNền địa hình sạch"),
        (0.39, 0.80, "2. Khóa đáp án\nSentinel-1 2021"),
        (0.70, 0.80, "3. Bằng chứng độc lập\nOSM lịch sử + HydroRIVERS + JRC"),
        (0.70, 0.48, "4. Dòng chảy từ V41\nD8 + upstream area + HAND"),
        (0.39, 0.48, "5. Lọc ứng viên ổn định\n407 ô stable Tier A"),
        (0.08, 0.48, "6. Kiểm topology\n24 cụm, giới hạn cut"),
        (0.08, 0.16, "7. Áp dụng có kiểm soát\n38 ô / 16 cụm"),
        (0.39, 0.16, "8. Audit cuối\nPASS mọi invariant"),
        (0.70, 0.16, "9. Bước tiếp theo\nChạy BG_Flood + validation SAR"),
    ]
    for x, y, txt in boxes:
        ax.text(
            x, y, txt, transform=ax.transAxes, ha="center", va="center", fontsize=11,
            bbox=dict(boxstyle="round,pad=0.6", facecolor="#EEF5FA", edgecolor="#6B879B", linewidth=1.2),
        )
    arrows = [
        ((0.18, .80), (.29, .80)), ((.49,.80),(.60,.80)), ((.70,.72),(.70,.58)),
        ((.60,.48),(.49,.48)), ((.29,.48),(.18,.48)), ((.08,.40),(.08,.24)),
        ((.18,.16),(.29,.16)), ((.49,.16),(.60,.16)),
    ]
    for a, b in arrows:
        ax.annotate("", xy=b, xytext=a, xycoords=ax.transAxes, textcoords=ax.transAxes,
                    arrowprops=dict(arrowstyle="->", linewidth=1.5, color="#5E6B73"))
    ax.text(.5,.95,"Quy trình V46 theo cách dễ hiểu",ha="center",va="center",fontsize=18,fontweight="bold",transform=ax.transAxes)
    ax.text(.5,.03,"Điểm quan trọng: ảnh Sentinel-1 năm 2021 được giữ riêng để kiểm tra sau cùng.",ha="center",fontsize=11,transform=ax.transAxes)
    fig.tight_layout()
    fig.savefig(FIG_FLOW, dpi=200, bbox_inches="tight")
    plt.close(fig)

    # 2) Simple cross-section before/after
    x = np.arange(9)
    before = np.array([5.0, 4.8, 4.7, 6.4, 6.7, 6.2, 4.6, 4.4, 4.2])
    after = before.copy()
    after[3:6] = [5.25, 5.10, 4.95]
    fig, ax = plt.subplots(figsize=(10, 5.6))
    ax.plot(x, before, marker="o", label="Trước V46")
    ax.plot(x, after, marker="o", label="Sau cut có giới hạn")
    ax.fill_between(x, after, before, where=before>after, alpha=0.18)
    ax.set_title("Ví dụ đơn giản: chỉ hạ một chỗ cản dòng đã có bằng chứng")
    ax.set_xlabel("Các ô DEM liên tiếp (mỗi ô 30 m)")
    ax.set_ylabel("Cao độ minh họa (m)")
    ax.grid(True, alpha=.25)
    ax.legend()
    ax.text(4, 6.85, "'Gờ' địa hình cản nước", ha="center")
    ax.text(4, 5.45, "V46 tạo một đường thấp vừa đủ\nkhông san phẳng cả vùng", ha="center")
    fig.tight_layout()
    fig.savefig(FIG_CUT, dpi=200, bbox_inches="tight")
    plt.close(fig)

    # 3) Actual modified cells over a DEM backdrop
    cells = pd.read_csv(ROOT / "V46_audit_20260930" / "V46_applied_cells.csv")
    dem_path = ROOT / "05-Du-lieu-GIS" / "PhuYen_DEM_30m_V46.asc"
    with rasterio.open(dem_path) as src:
        dem = src.read(1, out_shape=(max(1, src.height//5), max(1, src.width//5)), resampling=rasterio.enums.Resampling.bilinear)
        nodata = src.nodata
        if nodata is not None:
            dem = np.where(dem == nodata, np.nan, dem)
        b = src.bounds
    finite = dem[np.isfinite(dem)]
    vmax = float(np.nanpercentile(finite, 95)) if finite.size else 100.0
    fig, ax = plt.subplots(figsize=(8.5, 8.2))
    ax.imshow(dem, extent=[b.left,b.right,b.bottom,b.top], origin="upper", cmap="terrain", vmin=0, vmax=vmax)
    ax.scatter(cells["x"], cells["y"], s=28, edgecolors="black", linewidths=.4, label="38 ô V46 đã hiệu chỉnh")
    ax.set_title("Vị trí 38 ô DEM thực sự được hiệu chỉnh trong V46")
    ax.set_xlabel("Easting UTM 49N (m)")
    ax.set_ylabel("Northing UTM 49N (m)")
    ax.legend(loc="lower left")
    fig.tight_layout()
    fig.savefig(FIG_MAP, dpi=200, bbox_inches="tight")
    plt.close(fig)


def revise_thesis():
    doc = Document(SRC)

    # Abstract / overview
    replace_text(find_para(doc, "Nghiên cứu tập trung vào ảnh hưởng"),
        "Nghiên cứu tập trung vào ảnh hưởng của địa hình raster và các rào cản nhân tạo đối với kết nối thủy lực giữa lòng sông và đồng bằng. Các phiên bản V1–V44 lần lượt thử nghiệm cấu hình nhám, xử lý hố trũng bằng Priority-Flood và hiệu chỉnh cục bộ cao độ. Kiểm toán V45 cho thấy phương án san gạt theo vùng đệm có thể gây thay đổi địa hình quá lớn nên không được dùng làm DEM khoa học. Từ đó, đồ án xây dựng V46 theo hướng hydrological conditioning độc lập: dùng V41 làm nền, giữ Sentinel-1 năm 2021 làm tập kiểm định, kết hợp OSM lịch sử trước sự kiện, HydroRIVERS, JRC Global Surface Water trước 2021 và topology D8/HAND. V46 cuối cùng chỉ hạ 38 ô thuộc 16 cụm, diện tích 0,0342 km², với độ hạ 0,735–2,9625 m; toàn bộ kiểm tra provenance, topology, giới hạn cut và điều kiện không dùng Sentinel-1 để chọn ô đều PASS. V46 mới chứng minh quy trình điều kiện hóa DEM có kiểm soát; hiệu quả mô phỏng ngập vẫn phải được đánh giá bằng một lần chạy BG_Flood và validation độc lập sau khi khóa DEM.")

    replace_text(find_para(doc, "The study examines how raster topography"),
        "The study examines how raster topography and artificial barriers affect hydraulic connectivity between the river channel and floodplain. Model versions V1–V44 tested roughness configuration, Priority-Flood depression treatment and local terrain conditioning. Audit of the V45 buffer-based carving experiment showed that it could introduce unrealistically large terrain changes, so it is not used as the scientific DEM. The study therefore develops V46 as an independently conditioned DEM using V41 as the clean base, keeping the 2021 Sentinel-1 flood extent as a holdout, and combining pre-event historical OSM, HydroRIVERS, pre-2021 JRC Global Surface Water, and terrain-derived D8/HAND evidence. The final V46 lowers only 38 cells in 16 components over 0.0342 km², with cut depths from 0.735 to 2.9625 m; provenance, topology, cut constraints, and the Sentinel holdout checks all pass. These results establish an auditable terrain-conditioning workflow, while flood-model improvement remains to be tested by a new BG_Flood run and independent validation after V46 is frozen.")

    replace_text(find_para(doc, "- Thực nghiệm số học kế tiếp"),
        "- Thực nghiệm số học kế tiếp (Iterative Computational Experimentation): kế thừa và so sánh các kịch bản V1–V44; dùng V45 như nhánh chẩn đoán để nhận diện rủi ro over-conditioning; sau đó chuyển sang V46 với V41 làm base sạch và chỉ cho phép thay đổi DEM khi có đồng thuận hydrography, topology D8 và giới hạn cut được audit.")
    replace_text(find_para(doc, "- Phân tích hạn chế của cách chọn ô hạ đê"),
        "- Phân tích hạn chế của cách chọn ô hạ đê dựa trên mặt nạ SAR trong V44, kiểm toán rủi ro san gạt rộng của V45, rồi xây dựng V46 theo nguyên tắc tách dữ liệu hiệu chỉnh khỏi dữ liệu kiểm định. Sentinel-1 năm 2021 không được dùng để chọn vị trí cut trong V46 mà chỉ được giữ cho bước validation sau cùng.")
    replace_text(find_para(doc, "Chương 3: Kết quả thực nghiệm và thảo luận."),
        "Chương 3: Kết quả thực nghiệm và thảo luận. Trình bày cơ sở dữ liệu đầu vào, lịch sử các kịch bản V1–V44, đối chiếu V44 với mặt nạ SAR, phân tích HAND, chẩn đoán giới hạn V44–V45 và kết quả xây dựng DEM V46 bằng bằng chứng thủy văn độc lập. Phần V46 báo cáo riêng quá trình chọn ứng viên, kiểm topology, giới hạn độ hạ, các invariant kiểm toán và giới hạn chưa có validation thủy lực sau V46. Các kết quả GeoAI không thuộc miền Phú Yên được loại khỏi phần kết quả định lượng.")

    # Method section: make V45 explicitly exploratory and add V46 method
    replace_text(find_para(doc, "2.3.3. Thuật toán San phẳng hình thái"),
                 "2.3.3. Thử nghiệm V45 và giới hạn của san phẳng hình thái", bold=True)
    replace_text(find_para(doc, "Phân tích V44 cho thấy điều kiện"),
        "Phân tích V44 cho thấy điều kiện `sar_flood == 1` có thể bỏ qua các đỉnh cản khô nằm giữa lòng dẫn và vùng ngập phía sau. V45 vì vậy từng thử morphological carving trong vùng đệm quanh sông. Tuy nhiên, kiểm toán sau đó cho thấy cách hạ đồng mức theo buffer có thể thay đổi cả các ô không có bằng chứng công trình hoặc topology thủy văn trực tiếp; đây là rủi ro over-conditioning của DEM.")
    replace_text(find_para(doc, "- Bước 1: Dùng phép toán hình thái học GIS"),
        "- V45 tạo vùng đệm rộng quanh lòng dẫn và tìm các ô cao vượt ngưỡng để thử hạ cao độ. Cách này hữu ích như một thí nghiệm chẩn đoán nhưng không đủ chặt để xác định vị trí cống/kênh hay cao độ lòng dẫn thực tế.")
    replace_text(find_para(doc, "- Bước 2: Quét toàn bộ các ô lưới"),
        "- Audit V45 ghi nhận 41.907 ô bị thay đổi, mức hạ cực đại khoảng 171,67 m. Quy mô thay đổi này vượt xa mức có thể bảo vệ chỉ bằng bằng chứng thủy văn hiện có, nên V45 không được dùng làm DEM chuẩn cho kiểm định khoa học.")
    replace_text(find_para(doc, "Phương pháp này tương đương với giải pháp hạ thấp đồng mức"),
        "Bài học từ V45 là mọi can thiệp DEM cần có lineage rõ, bằng chứng đa nguồn, kiểm tra topology và giới hạn độ hạ/chiều dài đường cắt. Nguyên tắc này là cơ sở để thiết kế V46.")

    anchor_24 = find_para(doc, "2.4. Phương pháp đánh giá")
    heading_like = find_para(doc, "2.3.3. Thử nghiệm V45")
    body_like = find_para(doc, "Bài học từ V45")
    insert_before(anchor_24, "2.3.4. Điều kiện hóa DEM thủy văn V46 bằng bằng chứng độc lập", like=heading_like, bold=True)
    insert_before(anchor_24,
        "V46 được xây dựng từ `PhuYen_DEM_30m_V41.asc` thay vì kế thừa V44/V45, nhằm tránh việc dùng chính mặt nạ ngập Sentinel-1 năm 2021 vừa để sửa địa hình vừa để kiểm định. V41 được giữ nguyên làm base; các raster dẫn xuất dùng cho routing không ghi ngược vào DEM nền.", like=body_like)
    insert_before(anchor_24,
        "Nguồn bằng chứng gồm: (i) OpenStreetMap waterways tại snapshot trước sự kiện, mốc 25/11/2021; (ii) HydroRIVERS làm mạng thủy văn độc lập; (iii) JRC Global Surface Water occurrence giai đoạn 1984–2020; và (iv) các feature địa hình tính từ V41 gồm D8 flow direction, upstream area, HAND và local relief. Sentinel-1 trận lũ 2021 không tham gia tạo candidate hoặc quyết định cut [22], [36]–[40].", like=body_like)
    insert_before(anchor_24,
        "Quy trình V46 thực hiện theo bốn cổng: (1) tìm candidate ổn định qua nhiều corridor sensitivity; (2) chỉ giữ seed có hỗ trợ từ OSM lịch sử và hydrography độc lập; (3) yêu cầu topology D8 trực tiếp giữa upstream/downstream anchors; (4) tạo constrained profile với giới hạn độ sâu và số ô bị sửa. Nhánh strict cho phép tối đa 1 m và 4 ô/cụm; nhánh bounded review cho phép tối đa 3 m và 8 ô/cụm; đường nối dài quá 20 ô hoặc thiếu bằng chứng bổ sung sẽ bị reject.", like=body_like)
    insert_before(anchor_24,
        "Mỗi ô được sửa phải có audit record gồm tọa độ, cao độ trước/sau, độ hạ, component, OSM way, HydroRIVERS ID, JRC occurrence/upstream-area evidence, decision và reason code. Manifest lưu hash của đầu vào/đầu ra để có thể kiểm tra lại provenance. Đây là điểm khác căn bản so với việc san phẳng toàn vùng đệm về một cao độ cố định.", like=body_like)

    replace_text(find_para(doc, "2. Tiền xử lý dữ liệu:"),
        "2. Tiền xử lý dữ liệu: Nắn chỉnh địa hình ảnh SAR, lọc nhiễu và phân ngưỡng để tạo mặt nạ tham chiếu; xử lý DEM qua V41 và quy trình V46 hydrological conditioning dùng OSM lịch sử, HydroRIVERS, JRC pre-2021 và topology D8/HAND. Sentinel-1 năm 2021 được tách khỏi bước chọn cut của V46 và chỉ dùng ở validation sau khi DEM được khóa. Bản đồ độ nhám Manning và thấm đất được xây dựng độc lập cho mô hình thủy lực.")

    # Results: diagnose V45 and add actual V46 section
    replace_text(find_para(doc, "3.6. Chẩn đoán rào cản địa hình"),
                 "3.6. Chẩn đoán giới hạn V44–V45 và yêu cầu chuyển sang V46", bold=True)
    replace_text(find_para(doc, "Qua phân tích chi tiết trắc dọc"),
        "Qua phân tích mặt cắt địa hình dọc tuyến sông, nghiên cứu nhận diện một cơ chế cản dòng rõ tại một số vị trí: cao trình raster trên bờ cao hơn mặt nước mô phỏng, làm đứt kết nối giữa lòng sông và vùng trũng phía sau. Phát hiện này giải thích được một phần vùng bỏ sót, nhưng không được coi là bằng chứng rằng toàn bộ False Negative đều do đê hoặc rằng mọi đỉnh cao trên DEM là công trình thực địa.")
    replace_text(find_para(doc, "Phương pháp xử lý đề xuất cho Kịch bản V45"),
        "Kiểm toán V45 và quyết định chuyển hướng:", bold=True)
    replace_text(find_para(doc, "Để khắc phục hạn chế này mà không phụ thuộc"),
        "V45 từng thử loại bỏ điều kiện `sar_flood == 1` và san gạt theo vùng đệm để mở kết nối. Cách thử này giúp bộc lộ vấn đề nhưng đồng thời tạo ra quá nhiều thay đổi địa hình, nên không được dùng làm kết quả cuối.")
    replace_text(find_para(doc, "- Bước 1: Thiết lập vùng đệm"),
        "- Audit V45: 41.907 ô thay đổi; mức hạ lớn nhất khoảng 171,67 m. Không có bộ dữ liệu hình học đê/cống đủ để bảo vệ quy mô thay đổi đó.")
    replace_text(find_para(doc, "- Bước 2: Nhận diện và hạ cao trình"),
        "- Quyết định phương pháp: quay về V41 làm base sạch và chỉ condition DEM tại các vị trí có bằng chứng đa nguồn + topology + constraint, hình thành nhánh V46.")
    replace_text(find_para(doc, "Tệp `PhuYen_DEM_30m_V45.asc` hiện chỉ là DEM thử nghiệm"),
        "Tệp `PhuYen_DEM_30m_V45.asc` vì vậy chỉ được giữ như nhánh thử nghiệm/đối chứng về phương pháp. V46 được thiết kế để giải quyết đúng điểm yếu này: số ô sửa phải tối thiểu, có lý do cho từng ô, có nguồn trước sự kiện và không dùng Sentinel-1 2021 trong quyết định conditioning.")

    anchor_ml = find_para(doc, "3.7. Định hướng mô hình học máy")
    replace_text(anchor_ml, "3.8. Định hướng mô hình học máy thay thế (Surrogate Model)", bold=True)
    heading3_like = find_para(doc, "3.6. Chẩn đoán giới hạn")
    body3_like = find_para(doc, "Tệp `PhuYen_DEM_30m_V45.asc`")
    insert_before(anchor_ml, "3.7. Xây dựng và kiểm toán DEM V46 bằng bằng chứng thủy văn độc lập", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "3.7.1. Tách dữ liệu conditioning và dữ liệu validation", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "V46 sử dụng V41 làm base với SHA-256 `9fb70773...e15488`. Sentinel-1 flood extent năm 2021 được khóa hoàn toàn khỏi bước conditioning. Bằng chứng nước mặt dùng trong candidate là JRC occurrence 1984–2020; OSM được truy vấn theo snapshot trước ngày 25/11/2021. Cách tách này cho phép Sentinel-1 2021 tiếp tục đóng vai trò dữ liệu kiểm định độc lập sau khi V46 đã được freeze.", like=body3_like)
    insert_before(anchor_ml,
        "3.7.2. Tạo bằng chứng địa hình và candidate ổn định", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "Từ V41, hệ thống tạo D8 flow direction, upstream area, HAND theo ngưỡng drainage 1 km² và 10 km², cùng fill-depth chỉ phục vụ routing. Có 118.545 ô cần fill tạm khi thiết lập hướng dòng, nhưng độ fill cực đại chỉ 0,123 m và không ghi ngược vào V41. Qua ba mức corridor sensitivity, số Tier A lần lượt là 407, 1.580 và 4.530 ô; giao ổn định giữa các cấu hình là 407 ô. Điều này giúp loại các candidate chỉ xuất hiện khi thay đổi buffer.", like=body3_like)
    insert_before(anchor_ml,
        "3.7.3. Kiểm chứng lịch sử và topology", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "Một seed 55 ô được dùng làm danh sách tọa độ cần kiểm tra, nhưng OSM hiện tại không được coi là bằng chứng. Khi đối chiếu OSM lịch sử, cả 55/55 ô đều có hỗ trợ trước sự kiện; chúng tạo thành 24 component. Tất cả 24 component có đường nối D8 trực tiếp giữa upstream/downstream anchors. Sau khi áp dụng constraint, 1 component được `AUTO_APPROVE_STRICT`, 15 component được `REVIEW_APPROVE_BOUNDED` và 8 component bị `REJECT`.", like=body3_like)
    insert_before(anchor_ml,
        "Các constraint chính gồm: sai lệch OSM lịch sử tối đa 1 cell cho nhánh strict; direct D8 topology bắt buộc; cut strict ≤1 m và ≤4 ô/cụm; bounded review ≤3 m và ≤8 ô/cụm; path span ≤20 ô; đồng thời cần JRC occurrence ≥25% trong 1 cell hoặc upstream area ≥0,5 km² trong 2 cell. Những component không đạt cổng bị loại thay vì ép DEM phải khớp.", like=body3_like)
    insert_before(anchor_ml,
        "3.7.4. DEM V46 sau khi áp dụng", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "V46 cuối cùng sửa 38 ô thuộc 16 component, tương đương 0,0342 km². Mọi thay đổi đều là hạ cao độ; độ hạ nhỏ nhất 0,735 m, trung vị 1,160 m và lớn nhất 2,9625 m. Tổng thể tích hạ theo lưới là khoảng 49.290 m³. So với V45 thử nghiệm (41.907 ô và có cut cực đại khoảng 171,67 m), V46 thu hẹp can thiệp xuống những vị trí có thể truy vết bằng bằng chứng và topology.", like=body3_like)

    rows = [
        ("Base DEM", "V41", "Giữ base sạch; không kế thừa sửa đổi SAR-based của V44/V45"),
        ("JRC dùng cho conditioning", "1984–2020", "Kết thúc trước sự kiện 2021 để giảm leakage"),
        ("Seed lịch sử", "55/55 ô đạt", "OSM hiện tại chỉ là tọa độ seed, không phải bằng chứng"),
        ("Component topology", "24", "24/24 có direct D8 topology"),
        ("Quyết định", "1 strict + 15 bounded + 8 reject", "Fail closed khi không đạt constraint"),
        ("Ô DEM thực sự sửa", "38 ô / 16 component", "Diện tích 0,0342 km²"),
        ("Độ hạ", "0,735–2,9625 m", "Median 1,160 m; bounded cap 3 m"),
        ("Sentinel-1 2021", "Không dùng", "Giữ làm validation holdout"),
        ("Final audit", "PASS", "Hash, provenance, topology, reason code và cut caps đều đạt"),
    ]
    insert_table_before(doc, anchor_ml, rows)
    insert_before(anchor_ml,
        "Bảng 3.7: Tóm tắt các cổng kiểm soát và kết quả DEM V46", like=body3_like, italic=True)
    insert_before(anchor_ml,
        "3.7.5. Ý nghĩa và giới hạn", like=heading3_like, bold=True)
    insert_before(anchor_ml,
        "Kết quả trên cho phép gọi V46 là một DEM đã được hydrologically conditioned theo quy trình có audit. Tuy nhiên, đây chưa phải bằng chứng rằng mô phỏng ngập V46 tốt hơn V44. Cần chạy lại BG_Flood với forcing đã khóa provenance, sau đó mới mở holdout Sentinel-1 2021 và số liệu mực nước/lưu lượng để đánh giá CSI, Recall, Precision, cân bằng nước và tính hợp lý thời gian. Báo cáo không gán CSI cho V46 trước bước này.", like=body3_like)

    # Conclusions and recommendations
    replace_text(find_para(doc, "5. Xác định hạn chế logic của V44"),
        "5. Từ hạn chế của V44 và rủi ro over-conditioning của V45, xây dựng được DEM V46 theo quy trình độc lập và có audit: 55 ô seed được xác minh bằng OSM lịch sử, 24 component có direct D8 topology, 8 component bị reject theo constraint và chỉ 38 ô thuộc 16 component được áp dụng vào DEM. V46 giữ nguyên grid/nodata, chỉ hạ cao độ trong giới hạn ≤3 m, không sử dụng Sentinel-1 năm 2021 để quyết định cut và toàn bộ kiểm tra cuối đều PASS.")
    concl_anchor = find_para(doc, "2. Kiến nghị")
    insert_before(concl_anchor,
        "6. V46 hiện là kết quả tiền xử lý địa hình đã kiểm toán, chưa phải kết quả kiểm định ngập. Do chưa có lượt chạy BG_Flood dài hạn và đối chiếu Sentinel-1 sau khi khóa V46, đồ án không tuyên bố V46 làm tăng CSI/Recall/Precision. Đây là bước thực nghiệm tiếp theo cần hoàn thành để đánh giá lợi ích thủy lực thực sự của DEM mới.", like=find_para(doc, "5. Từ hạn chế của V44"))
    replace_text(find_para(doc, "1. Về dữ liệu địa hình:"),
        "1. Về dữ liệu địa hình và validation V46: dùng V46 làm nhánh DEM chính cho thí nghiệm kế tiếp; chạy BG_Flood với forcing đã khóa provenance, kiểm tra mass balance và ổn định số học, rồi mới mở Sentinel-1 2021 để đánh giá độc lập. Tại các component quan trọng, nên tiếp tục đối chiếu LiDAR/UAV/mặt cắt/hồ sơ cống nhằm xác nhận hình học thực địa và datum cao độ.")

    # References
    add_reference(doc, "[36] Lindsay, J. B. (2016). Efficient hybrid breaching-filling sink removal methods for flow path enforcement in digital elevation models. Hydrological Processes, 30(6), 846–857. https://doi.org/10.1002/hyp.10648")
    add_reference(doc, "[37] Lindsay, J. B. (2016). The practice of DEM stream burning revisited. Earth Surface Processes and Landforms, 41(5), 658–668. https://doi.org/10.1002/esp.3888")
    add_reference(doc, "[38] Yamazaki, D., Ikeshima, D., Sosa, J., Bates, P. D., Allen, G. H., & Pavelsky, T. M. (2019). MERIT Hydro: A high-resolution global hydrography map based on latest topography dataset. Water Resources Research, 55, 5053–5073. https://doi.org/10.1029/2019WR024873")
    add_reference(doc, "[39] Chen, Q., Mudd, S. M., Attal, M., & Hancock, S. (2024). Extracting an accurate river network: Stream burning re-revisited. Remote Sensing of Environment, 312, 114333. https://doi.org/10.1016/j.rse.2024.114333")
    add_reference(doc, "[40] Hawker, L., Neal, J., Savage, J., Kirkpatrick, T., Lord, R., Zylberberg, Y., Groeger, A., Thuy, T. D., Fox, S., Agyemang, F., & Nam, P. K. (2024). Assessing LISFLOOD-FP with the next-generation digital elevation model FABDEM using household survey and remote sensing data in the Central Highlands of Vietnam. Natural Hazards and Earth System Sciences, 24, 539–566. https://doi.org/10.5194/nhess-24-539-2024")

    props = doc.core_properties
    props.comments = "Bổ sung V46 ngày 30/09/2026: hydrological conditioning độc lập, audit topology/provenance, số liệu 38 ô/16 component và giới hạn validation sau V46."
    doc.save(OUT)
    return OUT


def set_doc_defaults(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0,0,0)
    for name, size in [("Heading 1",16),("Heading 2",13),("Heading 3",11)]:
        st = styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0,0,0)


def add_body(doc, text, *, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        r = p.add_run(text)
        r.italic = italic
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_simple_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i], val, size=9, align=WD_ALIGN_PARAGRAPH.CENTER if i==1 else None)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def build_explainer():
    doc = Document()
    set_doc_defaults(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Giải thích toàn bộ quy trình V46 cho người mới bắt đầu")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Đồ án ngập lụt Phú Yên BG Flood và Sentinel 1").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Mục tiêu của tài liệu này là giúp bạn hiểu V46 từ con số 0, đủ để tự giải thích lại khi bảo vệ đồ án.")

    doc.add_heading("1. Hiểu trong 30 giây", level=1)
    add_body(doc, "Hãy tưởng tượng DEM giống một sân trường được lát bằng 4,41 triệu viên gạch vuông. Mỗi viên gạch rộng 30 m x 30 m và có một con số cho biết nó cao bao nhiêu. Nước muốn chảy từ chỗ cao xuống chỗ thấp. Nếu một vài viên gạch bị ghi cao sai, chúng có thể giống như một bức tường vô hình làm nước không đi qua được.")
    add_body(doc, "V46 không san cả sân trường. V46 chỉ tìm những chỗ rất nhỏ có đủ bằng chứng rằng đường nước đáng lẽ phải nối qua đó, kiểm tra đường chảy, rồi hạ đúng vài ô trong giới hạn cho phép. Kết quả cuối chỉ sửa 38 ô trên 4.410.000 ô hợp lệ, tức là can thiệp rất nhỏ và có hồ sơ giải thích cho từng ô.")
    add_body(doc, "Điểm quan trọng nhất: ảnh Sentinel-1 của trận lũ 2021 được cất riêng như đáp án kiểm tra. V46 không nhìn đáp án đó khi quyết định chỗ nào cần sửa. Sau khi V46 đã khóa xong mới được dùng Sentinel-1 để chấm xem mô hình có tốt hơn hay không.")
    doc.add_picture(str(FIG_FLOW), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("2. Từ điển 10 từ cần nhớ", level=1)
    rows = [
        ("DEM", "Bản đồ cao độ dạng lưới", "Mỗi ô giống một viên gạch có độ cao."),
        ("V41", "Nền địa hình sạch dùng làm base", "Là tờ giấy gốc để V46 bắt đầu."),
        ("Hydrological conditioning", "Chỉnh DEM để đường thoát nước hợp lý hơn", "Chỉ chỉnh khi có bằng chứng; không chỉnh để ép khớp đáp án."),
        ("OSM lịch sử", "Mạng sông/kênh OpenStreetMap trước trận lũ", "Giúp chứng minh đường nước đã tồn tại trước 2021."),
        ("HydroRIVERS", "Mạng sông độc lập", "Nguồn thứ hai để đối chiếu, tránh tin một bản đồ duy nhất."),
        ("JRC occurrence", "Tần suất từng nơi từng có nước", "V46 dùng giai đoạn 1984–2020 để tránh nhìn sự kiện 2021."),
        ("D8", "Mỗi ô chảy sang một trong 8 ô lân cận", "Dùng để kiểm tra đường nước có nối thật hay không."),
        ("HAND", "Độ cao so với drainage gần nhất", "Cho biết một ô cao hơn đường thoát nước bao nhiêu."),
        ("Topology", "Quan hệ nối trước–sau của mạng nước", "Có đường sông trên bản đồ chưa đủ; phải nối đúng hướng."),
        ("Audit", "Nhật ký kiểm tra có thể truy lại", "Biết ô nào sửa, vì sao sửa, dữ liệu nào đã dùng và hash là gì."),
    ]
    add_simple_table(doc,["Từ","Nghĩa","Cách nhớ"],rows,[3.2,5.0,8.0])

    doc.add_heading("3. Vì sao phải bỏ cách V45", level=1)
    add_body(doc, "V45 là một thử nghiệm hữu ích vì nó cho thấy nếu chỉ tạo buffer quanh sông rồi hạ mọi ô cao về một mốc cố định, mô hình có thể bị chỉnh quá tay. Audit V45 ghi nhận 41.907 ô bị thay đổi và có ô bị hạ tới khoảng 171,67 m. Con số đó không thể bảo vệ bằng bằng chứng đê/cống hiện có.")
    add_body(doc, "Bài học rút ra rất đơn giản: muốn sửa DEM, phải trả lời được ba câu hỏi cho từng chỗ: (1) Tại sao biết ở đây có đường nước? (2) Tại sao biết hai phía phải nối với nhau? (3) Hạ bao nhiêu là mức nhỏ nhất đủ để nối, và giới hạn tối đa là bao nhiêu?")

    doc.add_heading("4. Quy trình V46 từng bước", level=1)
    steps = [
        ("Bước 1", "Chọn V41 làm tờ giấy gốc", "Không kế thừa các chỉnh sửa có dùng SAR của V44/V45. Hash base được khóa để sau này biết có ai thay file hay không."),
        ("Bước 2", "Cất riêng đáp án Sentinel-1 2021", "Không dùng ảnh ngập 2021 để chọn ô sửa. Như làm bài kiểm tra: chưa làm xong thì không được mở đáp án."),
        ("Bước 3", "Gom bằng chứng có trước sự kiện", "Dùng OSM lịch sử trước 25/11/2021, HydroRIVERS và JRC water occurrence 1984–2020."),
        ("Bước 4", "Tính cách nước muốn chảy trên V41", "Tạo D8, upstream area và HAND. Raster fill chỉ dùng để routing; không ghi ngược vào V41."),
        ("Bước 5", "Tìm ứng viên ổn định", "Thử nhiều corridor. Tier A là 407 / 1.580 / 4.530 ô tùy mức, nhưng phần giao ổn định chỉ còn 407 ô."),
        ("Bước 6", "Kiểm 55 ô seed bằng dữ liệu lịch sử", "55/55 ô seed đều có hỗ trợ OSM lịch sử; chúng tạo thành 24 component."),
        ("Bước 7", "Kiểm đường nối topology", "24/24 component có direct D8 topology. Mỗi component phải có upstream/downstream anchors và bằng chứng bổ sung."),
        ("Bước 8", "Áp giới hạn hạ cao độ", "Strict: tối đa 1 m, 4 ô/cụm. Bounded review: tối đa 3 m, 8 ô/cụm. Path span tối đa 20 ô."),
        ("Bước 9", "Cho phép hoặc từ chối", "1 component auto strict, 15 component bounded review, 8 component reject. Không đạt là bỏ, không ép DEM."),
        ("Bước 10", "Tạo V46 và audit lại", "38 ô thuộc 16 component được hạ; mọi hash/invariant/reason code đều PASS."),
    ]
    for a,b,c in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r=p.add_run(f"{a}. {b}: "); r.bold=True
        p.add_run(c)

    doc.add_picture(str(FIG_CUT), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, "Hình trên chỉ là ví dụ minh họa. V46 thật không hạ một dải cố định như vậy; mỗi component có profile, anchor, cut depth và reason code riêng.", italic=True)

    doc.add_heading("5. Các con số thật phải nhớ", level=1)
    actual_rows = [
        ("Số ô hợp lệ trong miền", "4.410.000", "Toàn bộ lưới V41/V46 dùng cho kiểm tra"),
        ("Routing fill cells", "118.545", "Chỉ fill tạm để tính hướng dòng; không đổi V41"),
        ("Fill depth cực đại", "0,123 m", "Nhỏ và chỉ dùng cho routing"),
        ("Stable Tier A", "407 ô", "Giao giữa các mức sensitivity"),
        ("Seed lịch sử", "55/55 ô", "Tất cả có historical support"),
        ("Topology components", "24", "24/24 có direct D8"),
        ("Reject", "8 component", "Fail closed khi không đạt cổng"),
        ("Applied", "38 ô / 16 component", "Thực sự thay đổi DEM"),
        ("Diện tích sửa", "0,0342 km²", "38 x 900 m²"),
        ("Cut depth", "0,735–2,9625 m", "Median khoảng 1,160 m"),
        ("Thể tích hạ", "~49.290 m³", "Tính từ cut depth x diện tích cell"),
        ("Sentinel-1 2021 dùng để condition?", "Không", "Giữ làm holdout"),
        ("Final verification", "PASS", "Tất cả invariant đã kiểm tra đều đạt"),
    ]
    add_simple_table(doc,["Chỉ tiêu","Giá trị","Hiểu đơn giản"],actual_rows,[4.4,3.5,8.0])
    doc.add_picture(str(FIG_MAP), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, "Các điểm trên bản đồ là 38 ô thực sự khác V41. Đây là thay đổi rất thưa, không phải một vùng lớn bị san phẳng.", italic=True)

    doc.add_heading("6. V46 chứng minh được gì và chưa chứng minh được gì", level=1)
    add_simple_table(doc,["Đã chứng minh","Chưa được phép kết luận"],[
        ("Có một quy trình chọn vị trí sửa DEM có provenance, historical evidence, topology và cut constraint.", "V46 chắc chắn làm mô hình ngập chính xác hơn V44."),
        ("38 ô được sửa có reason code, OSM + HydroRIVERS và direct D8 evidence.", "CSI/Recall/Precision của V46 cao hơn bao nhiêu."),
        ("Sentinel-1 2021 không được dùng trong conditioning.", "38 ô đó đều là cống/đê thật ngoài thực địa."),
        ("Grid, NoData và base lineage được bảo toàn; mọi change là lowering.", "Mức cut là cao độ lòng sông/cống đo được ngoài hiện trường."),
        ("Final audit của pipeline PASS.", "Mô hình thủy lực đã pass mass balance hoặc hydrograph validation sau V46."),
    ],[8.0,8.0])

    doc.add_heading("7. Tại sao 'không dùng Sentinel-1 để sửa V46' lại quan trọng", level=1)
    add_body(doc, "Ví dụ: giáo viên cho bạn đề toán và đáp án. Nếu bạn nhìn đáp án rồi sửa từng bước tính cho đến khi ra đúng đáp án, sau đó dùng chính đáp án đó để nói rằng cách tính của bạn chính xác, thì phép kiểm tra không còn độc lập. V44 có phần logic dùng SAR để chọn nơi hạ DEM nên độ khớp với SAR không phải validation độc lập hoàn toàn.")
    add_body(doc, "V46 cố tình không nhìn Sentinel-1 2021 khi chọn cut. Vì vậy, nếu sau này chạy BG_Flood với V46 và kết quả gần SAR hơn, bằng chứng đó mạnh hơn nhiều: DEM đã được khóa bằng dữ liệu khác rồi mới mở đáp án vệ tinh để chấm.")

    doc.add_heading("8. Nếu thầy cô hỏi khi bảo vệ", level=1)
    qa = [
        ("V46 khác V45 ở đâu?", "V45 hạ theo buffer rộng và threshold, tạo 41.907 thay đổi với cực đại khoảng 171,67 m. V46 chỉ sửa 38 ô có bằng chứng lịch sử + mạng sông độc lập + direct D8 topology và cap tối đa 3 m."),
        ("Tại sao dùng V41 làm base?", "V41 là base trước các sửa đổi dùng SAR ở V44/V45, nên phù hợp hơn để giữ Sentinel-1 2021 làm dữ liệu kiểm định độc lập."),
        ("55 ô seed có phải ground truth không?", "Không. Seed chỉ là danh sách tọa độ cần kiểm tra. OSM hiện tại không được dùng làm bằng chứng; từng seed phải được xác nhận bằng OSM lịch sử và nguồn khác."),
        ("Tại sao cần HydroRIVERS nếu đã có OSM?", "Để không phụ thuộc một nguồn duy nhất. Hai mạng độc lập cùng hỗ trợ làm bằng chứng mạnh hơn."),
        ("JRC 1984–2020 dùng để làm gì?", "Để biết một vị trí có lịch sử xuất hiện nước hay không, đồng thời tránh sử dụng trực tiếp sự kiện 2021."),
        ("D8 topology là gì?", "Là kiểm tra rằng theo địa hình, có một đường chảy trực tiếp hợp lý từ phía thượng lưu sang hạ lưu của component."),
        ("Tại sao 8 component bị reject?", "Vì V46 fail closed: nếu lệch centerline, thiếu cổng bằng chứng hoặc vượt constraint thì bỏ, không cố sửa cho đủ."),
        ("V46 có tốt hơn V44 chưa?", "Chưa được phép kết luận. V46 mới là DEM đã condition và audit. Phải chạy BG_Flood rồi validation bằng Sentinel-1/mực nước sau khi khóa V46."),
        ("38 ô có ít quá không?", "Mục tiêu không phải sửa nhiều. Mục tiêu là sửa tối thiểu những điểm cản dòng có thể chứng minh. Ít nhưng có bằng chứng tốt hơn nhiều và khó giải thích."),
        ("Điểm mới khoa học của đồ án là gì?", "Chuyển từ chỉnh DEM để đuổi theo SAR sang quy trình condition DEM có holdout, historical evidence, multi-source consensus, direct topology, constrained cut và audit per-cell."),
    ]
    for q,a in qa:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5)
        p.add_run("Hỏi: "+q+" ").bold=True
        p.add_run("Trả lời: "+a)

    doc.add_heading("9. Bài nói 2 phút để bạn học thuộc ý", level=1)
    speech = (
        "Sau khi kiểm tra V44 và V45, em nhận thấy nếu dùng chính ảnh SAR 2021 để chọn vị trí sửa DEM thì phép đối chiếu sau đó không còn độc lập hoàn toàn. V45 còn có rủi ro chỉnh địa hình quá mạnh, vì audit cho thấy hơn 41 nghìn ô thay đổi và có cut rất lớn. Vì vậy em quay về V41 làm base sạch và xây dựng V46 theo nguyên tắc chỉ sửa khi có bằng chứng độc lập. Em dùng OSM lịch sử trước trận lũ, HydroRIVERS, JRC water occurrence trước 2021 và các chỉ báo địa hình D8, upstream area, HAND. Sau sensitivity, historical validation và topology check, 24 component được xem xét; 8 bị loại, chỉ 16 component được áp dụng. DEM V46 cuối cùng chỉ hạ 38 ô, độ hạ từ 0,735 đến 2,9625 m, toàn bộ có reason code và manifest kiểm toán. Sentinel-1 2021 không tham gia conditioning. Vì vậy điểm em có thể kết luận hiện tại là V46 có quy trình địa hình chặt và có thể audit; em chưa kết luận V46 làm tăng CSI cho đến khi chạy lại BG_Flood và mở tập Sentinel-1 holdout để validation."
    )
    add_body(doc, speech)

    doc.add_heading("10. Bạn cần mở file nào khi muốn kiểm tra lại", level=1)
    file_rows = [
        ("05-Du-lieu-GIS/PhuYen_DEM_30m_V41.asc", "Base DEM sạch"),
        ("V46_hydro_evidence_20260930/V46_hydro_evidence_summary.json", "D8, HAND, JRC, fill-depth và invariant Sentinel holdout"),
        ("V46_seed55_historical_validated_20260930/V46_seed55_validation_summary.json", "55 seed và OSM lịch sử"),
        ("V46_stable_core_20260930/V46_stable_core_summary.json", "Sensitivity và 407 stable Tier A"),
        ("V46_topology_20260930/V46_topology_summary.json", "24 component, decision strict/bounded/reject, constraint"),
        ("V46_audit_20260930/V46_applied_cells.csv", "Từng ô thực sự bị hạ và reason code"),
        ("V46_audit_20260930/V46_conditioning_summary.json", "38 ô, 0,0342 km², cut và volume"),
        ("V46_audit_20260930/V46_final_verification.json", "Bảng kiểm cuối PASS/FAIL"),
        ("05-Du-lieu-GIS/PhuYen_DEM_30m_V46.asc", "DEM V46 cuối cùng"),
    ]
    add_simple_table(doc,["File","Dùng để hiểu gì"],file_rows,[8.0,8.0])

    doc.add_heading("11. Nguồn khoa học chính để đọc khi cần", level=1)
    refs = [
        "Lindsay (2016), Efficient hybrid breaching-filling sink removal methods, Hydrological Processes, DOI 10.1002/hyp.10648.",
        "Lindsay (2016), The practice of DEM stream burning revisited, Earth Surface Processes and Landforms, DOI 10.1002/esp.3888.",
        "Yamazaki et al. (2019), MERIT Hydro, Water Resources Research, DOI 10.1029/2019WR024873.",
        "Pekel et al. (2016), High-resolution mapping of global surface water and its long-term changes, Nature, DOI 10.1038/nature20584.",
        "Nobre et al. (2011), Height Above the Nearest Drainage, Journal of Hydrology, DOI 10.1016/j.jhydrol.2011.03.051.",
        "Chen et al. (2024), Extracting an accurate river network: Stream burning re-revisited, Remote Sensing of Environment 312, 114333, DOI 10.1016/j.rse.2024.114333.",
        "Hawker et al. (2024), Assessing LISFLOOD-FP with FABDEM in the Central Highlands of Vietnam, NHESS 24, 539–566, DOI 10.5194/nhess-24-539-2024.",
    ]
    add_bullets(doc, refs)

    doc.add_heading("12. Checklist trước khi nói 'V46 tốt hơn'", level=1)
    add_bullets(doc,[
        "Khóa forcing mưa, inflow, infiltration và datum; ghi hash/version.",
        "Chạy BG_Flood đủ thời gian với V46 và lưu log ổn định/CFL/NaN.",
        "Kiểm tra cân bằng nước trước khi nhìn SAR.",
        "Chỉ sau khi run hoàn tất mới mở Sentinel-1 2021 holdout để tính CSI/Recall/Precision.",
        "Đối chiếu thêm mực nước/lưu lượng hoặc vết lũ độc lập nếu có.",
        "Báo cáo cả cải thiện lẫn suy giảm; không chọn chỉ frame đẹp nhất để kết luận.",
        "Nếu một component V46 gây hành vi bất thường, quay lại audit record và bằng chứng địa hình thay vì chỉnh theo kết quả SAR.",
    ])

    props=doc.core_properties
    props.title="Giải thích toàn bộ quy trình V46 cho người mới bắt đầu"
    props.subject="Tài liệu học và chuẩn bị bảo vệ đồ án Phú Yên BG Flood"
    props.comments="Giải thích đơn giản từ dữ liệu audit V46 ngày 30/09/2026."
    doc.save(EXPLAIN)
    return EXPLAIN


def write_research_note_if_missing():
    if RESEARCH_NOTE.exists() and RESEARCH_NOTE.stat().st_size > 1000:
        return
    note = f"""# Nghiên cứu bổ sung V46 cho đồ án Phú Yên — 30/09/2026

## Kết luận có thể đưa vào luận văn

V46 chuyển trọng tâm từ hiệu chỉnh DEM theo kết quả Sentinel-1 sang hydrological conditioning có dữ liệu holdout. Base là V41; Sentinel-1 2021 không tham gia chọn cell hoặc quyết định cut. Hydrography/evidence gồm OSM lịch sử trước 25/11/2021, HydroRIVERS, JRC occurrence 1984–2020 và các feature D8/HAND/upstream area tính từ V41.

Kết quả audit nội bộ: 55/55 seed có historical OSM support; 24 component có direct D8 topology; 1 component strict auto-approved, 15 bounded review-approved, 8 rejected. DEM cuối sửa {COND['selection']['modified_cells']} ô, diện tích {COND['selection']['modified_area_km2']} km², cut {COND['selection']['cut_depth_min_m']:.3f}–{COND['selection']['cut_depth_max_m']:.3f} m, median {COND['selection']['cut_depth_median_m']:.3f} m. Final verification status: {VERIFY['status']}.

Điểm phải giữ trong luận văn: V46 là DEM conditioned đã audit, chưa phải bằng chứng mô phỏng ngập tốt hơn. CSI/Recall/Precision chỉ được báo sau một lượt BG_Flood mới và validation độc lập bằng Sentinel-1/gauge sau khi V46 được freeze.

## Cơ sở khoa học chính

- Lindsay (2016), constrained/hybrid breaching-filling: https://doi.org/10.1002/hyp.10648
- Lindsay (2016), stream burning và lỗi topology do scale mismatch: https://doi.org/10.1002/esp.3888
- Yamazaki et al. (2019), MERIT Hydro kết hợp DEM, water body data và OSM để tạo hydrography/hydrologically adjusted elevation: https://doi.org/10.1029/2019WR024873
- Pekel et al. (2016), JRC Global Surface Water: https://doi.org/10.1038/nature20584
- Nobre et al. (2011), HAND: https://doi.org/10.1016/j.jhydrol.2011.03.051
- Chen et al. (2024), stream burning re-revisited: https://doi.org/10.1016/j.rse.2024.114333
- Hawker et al. (2024), đánh giá FABDEM trong mô hình ngập tại Tây Nguyên Việt Nam: https://doi.org/10.5194/nhess-24-539-2024

## Không được viết quá mức bằng chứng

- Không gọi 38 ô là 38 cống thực địa nếu chưa có khảo sát công trình.
- Không gọi cut depth là độ sâu lòng sông thực.
- Không nói V46 tăng CSI khi chưa chạy mô hình sau V46.
- Không dùng OSM hiện tại như historical evidence; seed hiện tại chỉ dùng định vị.
- Không gọi HAND là validation độc lập về mực nước; HAND là terrain plausibility feature.
"""
    RESEARCH_NOTE.write_text(note, encoding="utf-8")


def main():
    build_figures()
    write_research_note_if_missing()
    thesis = revise_thesis()
    explain = build_explainer()
    print(f"THESIS={thesis} size={thesis.stat().st_size}")
    print(f"EXPLAIN={explain} size={explain.stat().st_size}")
    print(f"FIG_FLOW={FIG_FLOW} size={FIG_FLOW.stat().st_size}")
    print(f"FIG_CUT={FIG_CUT} size={FIG_CUT.stat().st_size}")
    print(f"FIG_MAP={FIG_MAP} size={FIG_MAP.stat().st_size}")


if __name__ == "__main__":
    main()
