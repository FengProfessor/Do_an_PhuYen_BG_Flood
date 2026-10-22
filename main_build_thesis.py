# -*- coding: utf-8 -*-
"""
Chương trình chính tổng hợp và sinh toàn văn Đồ án tốt nghiệp Đại học
Đạt chuẩn format Trường ĐH Công nghệ - ĐHQGHN (TS. Hà Minh Cường hướng dẫn)
Dung lượng >= 50 trang, hơn 35 tài liệu tham khảo có DOI/Live URL kiểm chứng.
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

from thesis_front_matter import build_front_matter
from thesis_introduction import build_introduction
from thesis_chapter1 import build_chapter1
from thesis_chapter2 import build_chapter2
from thesis_chapter3 import build_chapter3
from thesis_conclusion_refs import build_conclusion_and_refs

OUTPUT_PATH = r"D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx"

print("Đang khởi tạo tài liệu Word chuẩn A4...")
doc = docx.Document()

# Thiết lập lề trang chuẩn đóng gáy A4
sec = doc.sections[0]
sec.page_width = Inches(8.27)    # 21.0 cm
sec.page_height = Inches(11.69)  # 29.7 cm
sec.top_margin = Inches(0.98)     # 2.5 cm
sec.bottom_margin = Inches(1.18)  # 3.0 cm
sec.left_margin = Inches(1.18)    # 3.0 cm (đóng gáy sách)
sec.right_margin = Inches(0.79)   # 2.0 cm

# Footer
footer = sec.footer
p_ft = footer.paragraphs[0]
p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_ft = p_ft.add_run("Trường Đại học Công nghệ – ĐHQGHN | Viện Công nghệ Hàng không Vũ trụ")
r_ft.font.name = 'Times New Roman'
r_ft.font.size = Pt(9)
r_ft.font.italic = True
r_ft.font.color.rgb = RGBColor(128, 128, 128)

# Phông chữ Normal
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(13)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.25
normal_style.paragraph_format.space_before = Pt(0)
normal_style.paragraph_format.space_after = Pt(6)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="{top}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:left w:val="{left}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:bottom w:val="{bottom}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:right w:val="{right}" w:sz="4" w:space="0" w:color="B0B0B0"/>
        </w:tcBorders>
    '''.format(
        top=kwargs.get('top', 'single'),
        left=kwargs.get('left', 'none'),
        bottom=kwargs.get('bottom', 'single'),
        right=kwargs.get('right', 'none')
    ))
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def add_h(text, level=1, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(10)
    elif level == 2:
        run.font.size = Pt(13.5)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(13)
        run.italic = True
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def add_p(text, bold_prefix="", indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    return p

def add_caption(text, is_table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    if is_table:
        run.bold = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.italic = True
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)

def add_fig(img_path, caption_text, width_inches=5.8):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_inches))
        add_caption(caption_text, is_table=False)
    else:
        add_caption(f"[{caption_text}]", is_table=False)

def add_formula(text, formula_id):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.italic = True
    run.font.name = 'Times New Roman'
    r_id = p.add_run(f"\t\t\t\t\t({formula_id})")
    r_id.bold = True
    r_id.font.name = 'Times New Roman'

# Tiến hành lắp ghép toàn văn đồ án
print("Đang viết Front Matter...")
build_front_matter(doc, add_h, add_p, set_cell_border, set_cell_shading)

print("Đang viết Phần Mở đầu...")
build_introduction(doc, add_h, add_p)

print("Đang viết Chương 1...")
build_chapter1(doc, add_h, add_p, add_fig)

print("Đang viết Chương 2...")
build_chapter2(doc, add_h, add_p, add_formula, add_fig, add_caption, set_cell_border, set_cell_shading)

print("Đang viết Chương 3...")
build_chapter3(doc, add_h, add_p, add_caption, add_fig, set_cell_border, set_cell_shading)

print("Đang viết Kết luận, Kiến nghị và Tài liệu tham khảo...")
build_conclusion_and_refs(doc, add_h, add_p)

print(f"Đang lưu tài liệu ra {OUTPUT_PATH}...")
doc.save(OUTPUT_PATH)

total_words = sum(len(p.text.split()) for p in doc.paragraphs)
total_paras = len(doc.paragraphs)
total_tables = len(doc.tables)
total_imgs = len(doc.inline_shapes)

print("="*60)
print(f"Biên dịch thành công toàn văn đồ án!")
print(f"Đường dẫn file: {OUTPUT_PATH}")
print(f"Tổng số đoạn văn (Paragraphs): {total_paras}")
print(f"Tổng số từ (Word count): {total_words}")
print(f"Tổng số bảng biểu (Tables): {total_tables}")
print(f"Tổng số hình ảnh nhúng (Figures): {total_imgs}")
print("="*60)
