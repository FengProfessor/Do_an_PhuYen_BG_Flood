# -*- coding: utf-8 -*-
"""
Empirical Verification Script for De-AI Keyword Audit
Author: challenger_m6_1 (De-AI Keyword Adversarial Challenger)
Target:
- D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx
- D:\VTP\thesis_front_matter.py
- D:\VTP\thesis_introduction.py
- D:\VTP\thesis_chapter1.py
- D:\VTP\thesis_chapter2.py
- D:\VTP\thesis_chapter3.py
- D:\VTP\thesis_conclusion_refs.py
- D:\VTP\main_build_thesis.py
"""

import os
import sys
import re
import zipfile
import xml.etree.ElementTree as ET

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

import docx

# Explicit list of strictly forbidden AI clichés and sensational phrases
STRICT_FORBIDDEN_TERMS = [
    # 1. Phrases from user prompt & Section 2026-09-28T08:04:48Z
    "đanh thép",
    "ngoạn mục",
    "không thể chối cãi",
    "bức tường thành",
    "thảm khốc",
    "chí mạng",
    "kỷ lục",
    "xe ủi",
    
    # 2. Phrases from DISPATCH.md
    "minh chứng đanh thép",
    "bằng chứng đanh thép",
    "bước đột phá ngoạn mục",
    "bứt phá ngoạn mục",
    "bức tường đê khóa chặt",
    "vô vàn thiệt hại",
    "vô vàn",
    "thảm họa khốc liệt",
    "tàn khốc",
    "khốc liệt",
    "nút thắt cổ chai chí mạng",
    "mức kỷ lục",
    "xe ủi công trình",
    "vựa lúa nuôi sống",
    "bất lực hoàn toàn",
    "vũ khí đắc lực",
    
    # 3. Expanded AI clichés & sensational buzzwords
    "vô tiền khoáng hậu",
    "như một phép màu",
    "thần kỳ",
    "bảo chứng",
    "ngoan cường",
    "khuất phục",
    "đầu hàng",
    "cứu tinh",
    "song kiếm hợp bích",
    "bảo bối",
    "kinh ngạc",
    "choáng ngợp",
    "khủng khiếp",
    "kinh hoàng",
    "tuyệt hảo",
    "bất khả xâm phạm",
    "ổn định tuyệt đối",
    "an toàn tuyệt đối",
    "chính xác tuyệt đối"
]

# Legitimate scientific terms containing "tuyệt đối"
ALLOWED_TUYET_DOI_PATTERNS = [
    r"giá\s+trị\s+tuyệt\s+đối",      # absolute value |x|
    r"sai\s+số\s+tuyệt\s+đối",       # MAE - mean absolute error
    r"cao\s+trình\s+tuyệt\s+đối",     # absolute elevation above sea level
    r"cao\s+độ\s+tuyệt\s+đối",        # absolute elevation
    r"độ\s+cao\s+tuyệt\s+đối"         # absolute altitude
]

# Unicode Emoji regex
EMOJI_PATTERN = re.compile(
    r"[\U00010000-\U0010ffff"
    r"\u2600-\u26ff"
    r"\u2700-\u27bf"
    r"\u2300-\u23ff"
    r"\u2b50\u2b55"
    r"]",
    flags=re.UNICODE
)

def check_text_segment(text, source_name, location_info, is_narrative=True):
    findings = []
    text_lower = text.lower()

    # 1. Strictly forbidden AI clichés
    for term in STRICT_FORBIDDEN_TERMS:
        t_term = term.lower()
        if t_term in text_lower:
            start = 0
            while True:
                pos = text_lower.find(t_term, start)
                if pos == -1:
                    break
                snippet = text[max(0, pos-40):min(len(text), pos+len(t_term)+40)].strip()
                findings.append({
                    "severity": "CRITICAL",
                    "category": "STRICT_FORBIDDEN_CLICHE",
                    "term": term,
                    "source": source_name,
                    "location": location_info,
                    "snippet": snippet
                })
                start = pos + len(t_term)

    # 2. Sensational vs Scientific 'tuyệt đối'
    if "tuyệt đối" in text_lower:
        start = 0
        while True:
            pos = text_lower.find("tuyệt đối", start)
            if pos == -1:
                break
            snippet = text[max(0, pos-40):min(len(text), pos+len("tuyệt đối")+40)].strip()
            
            is_scientific = any(re.search(pat, snippet.lower()) for pat in ALLOWED_TUYET_DOI_PATTERNS)
            if is_scientific:
                findings.append({
                    "severity": "BENIGN_SCIENTIFIC",
                    "category": "SCIENTIFIC_TUYET_DOI",
                    "term": "tuyệt đối (khoa học: trắc địa/toán học)",
                    "source": source_name,
                    "location": location_info,
                    "snippet": snippet
                })
            else:
                findings.append({
                    "severity": "CRITICAL",
                    "category": "SENSATIONAL_TUYET_DOI",
                    "term": "tuyệt đối (cảm tính/phi khoa học)",
                    "source": source_name,
                    "location": location_info,
                    "snippet": snippet
                })
            start = pos + len("tuyệt đối")

    # 3. Unicode Emojis
    for m in EMOJI_PATTERN.finditer(text):
        pos = m.start()
        snippet = text[max(0, pos-20):min(len(text), pos+20)].strip()
        findings.append({
            "severity": "CRITICAL",
            "category": "EMOJI",
            "term": m.group(0),
            "source": source_name,
            "location": location_info,
            "snippet": snippet
        })

    # 4. Exclamation marks in narrative scientific prose
    if is_narrative and "!" in text:
        start = 0
        while True:
            pos = text.find("!", start)
            if pos == -1:
                break
            # Ignore code '!=', markdown syntax '!['
            if (pos + 1 < len(text) and text[pos+1] in ['=', '[']) or (pos > 0 and text[pos-1] in ['!', '/']):
                start = pos + 1
                continue
            
            snippet = text[max(0, pos-30):min(len(text), pos+30)].strip()
            findings.append({
                "severity": "WARNING_EXCLAMATION",
                "category": "NARRATIVE_EXCLAMATION",
                "term": "!",
                "source": source_name,
                "location": location_info,
                "snippet": snippet
            })
            start = pos + 1

    return findings

def audit_docx(docx_path):
    print("="*80)
    print(f"AUDITING DOCX FILE: {docx_path}")
    print("="*80)
    
    doc = docx.Document(docx_path)
    findings = []
    
    # 1. Body Paragraphs
    total_words = 0
    ref_count = 0
    for idx, p in enumerate(doc.paragraphs):
        p_text = p.text
        words = len(p_text.split())
        total_words += words
        if p_text.strip().startswith("[") and "]" in p_text and ("http" in p_text or "doi" in p_text.lower()):
            ref_count += 1
        
        # Front matter acknowledgements (Lời cảm ơn) can have polite exclamation
        is_ack = "cảm ơn" in p_text.lower()
        is_narrative = not is_ack
        
        if p_text.strip():
            f = check_text_segment(p_text, "DOCX:Paragraph", f"Para #{idx+1}", is_narrative=is_narrative)
            findings.extend(f)

    # 2. Tables
    table_words = 0
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    p_text = p.text
                    table_words += len(p_text.split())
                    if p_text.strip():
                        f = check_text_segment(p_text, "DOCX:Table", f"T{t_idx+1}:R{r_idx+1}C{c_idx+1}:P{p_idx+1}", is_narrative=True)
                        findings.extend(f)

    # 3. Headers and Footers
    for s_idx, sec in enumerate(doc.sections):
        for p in sec.header.paragraphs:
            if p.text.strip():
                f = check_text_segment(p.text, "DOCX:Header", f"Sec #{s_idx+1} Header", is_narrative=True)
                findings.extend(f)
        for p in sec.footer.paragraphs:
            if p.text.strip():
                f = check_text_segment(p.text, "DOCX:Footer", f"Sec #{s_idx+1} Footer", is_narrative=True)
                findings.extend(f)

    # 4. Raw XML scan of word/document.xml and other parts
    with zipfile.ZipFile(docx_path, 'r') as z:
        for fname in z.namelist():
            if fname.startswith("word/") and fname.endswith(".xml"):
                content = z.read(fname).decode('utf-8', errors='ignore')
                try:
                    root = ET.fromstring(content)
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    texts = [elem.text for elem in root.findall('.//w:t', ns) if elem.text]
                    full_xml_text = " ".join(texts)
                    xml_f = check_text_segment(full_xml_text, f"XML:{fname}", "XML Root", is_narrative=False)
                    for xf in xml_f:
                        if xf["severity"] == "CRITICAL":
                            findings.append(xf)
                except Exception as e:
                    print(f"Warning parsing XML node {fname}: {e}")

    total_all_words = total_words + table_words
    print(f"DOCX Physical Metrics:")
    print(f"  Total Paragraphs   : {len(doc.paragraphs)}")
    print(f"  Total Tables       : {len(doc.tables)}")
    print(f"  Total Inline Shapes: {len(doc.inline_shapes)}")
    print(f"  Total Word Count   : {total_all_words} (Paragraphs: {total_words}, Tables: {table_words})")
    print(f"  Verified References: {ref_count}")
    
    return findings, total_all_words, len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes), ref_count

def audit_python_files(py_files):
    print("\n" + "="*80)
    print("AUDITING PYTHON SOURCE FILES")
    print("="*80)
    all_findings = []
    for fpath in py_files:
        basename = os.path.basename(fpath)
        print(f"Scanning {basename}...")
        with open(fpath, 'r', encoding='utf-8', errors='replace') as f:
            lines = f.readlines()
        
        for l_idx, line in enumerate(lines, 1):
            line_str = line.strip()
            if not line_str:
                continue
            is_narrative = ('add_p(' in line or 'add_h(' in line) and not ("cảm ơn" in line.lower())
            f = check_text_segment(line_str, basename, f"Line {l_idx}", is_narrative=is_narrative)
            all_findings.extend(f)

    return all_findings

def main():
    docx_file = r"D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx"
    py_files = [
        r"D:\VTP\thesis_front_matter.py",
        r"D:\VTP\thesis_introduction.py",
        r"D:\VTP\thesis_chapter1.py",
        r"D:\VTP\thesis_chapter2.py",
        r"D:\VTP\thesis_chapter3.py",
        r"D:\VTP\thesis_conclusion_refs.py",
        r"D:\VTP\main_build_thesis.py"
    ]

    docx_findings, total_words, total_paras, total_tables, total_figs, ref_count = audit_docx(docx_file)
    py_findings = audit_python_files(py_files)

    all_findings = docx_findings + py_findings

    critical_items = [f for f in all_findings if f["severity"] == "CRITICAL"]
    warning_items = [f for f in all_findings if f["severity"] == "WARNING_EXCLAMATION"]
    benign_items = [f for f in all_findings if f["severity"] == "BENIGN_SCIENTIFIC"]

    print("\n" + "="*80)
    print("EMPIRICAL VERIFICATION SCORECARD")
    print("="*80)

    target_terms = [
        ("đanh thép / minh chứng đanh thép", "đanh thép"),
        ("ngoạn mục / bước đột phá ngoạn mục", "ngoạn mục"),
        ("không thể chối cãi", "không thể chối cãi"),
        ("bức tường thành / bức tường đê", "bức tường"),
        ("thảm khốc / thảm họa khốc liệt / tàn khốc", "khốc"),
        ("chí mạng / nút thắt cổ chai chí mạng", "chí mạng"),
        ("kỷ lục / mức kỷ lục", "kỷ lục"),
        ("tuyệt đối (cảm tính/phi khoa học)", "tuyệt đối (cảm tính/phi khoa học)"),
        ("xe ủi / xe ủi công trình", "xe ủi"),
        ("vô vàn / vô vàn thiệt hại", "vô vàn"),
        ("vựa lúa nuôi sống", "vựa lúa"),
        ("bất lực hoàn toàn", "bất lực"),
        ("vũ khí đắc lực", "vũ khí"),
        ("Emoji cảm xúc (🏆, 🔥, 🚀, v.v.)", "EMOJI"),
        ("Dấu chấm than (!) trong văn luận khoa học", "NARRATIVE_EXCLAMATION")
    ]

    print(f"{'Target Term / Cliché':<45} | {'Occurrences':<12} | {'Status'}")
    print("-" * 75)
    for label, key in target_terms:
        if key == "EMOJI":
            cnt = len([c for c in critical_items if c['category'] == "EMOJI"])
        elif key == "NARRATIVE_EXCLAMATION":
            cnt = len(warning_items)
        elif key == "tuyệt đối (cảm tính/phi khoa học)":
            cnt = len([c for c in critical_items if c['category'] == "SENSATIONAL_TUYET_DOI"])
        else:
            cnt = len([c for c in critical_items if key in c['term'].lower() or key in c['snippet'].lower()])
        
        status = "PASSED" if cnt == 0 else "FAILED"
        print(f"{label:<45} | {cnt:<12} | {status}")

    print("\n" + "="*80)
    print("DOCUMENT QUALITY & SIZE ACCEPTANCE CRITERIA")
    print("="*80)
    print(f"Word count: {total_words:,} words (Target: >= 18,000 words) -> {'PASSED' if total_words >= 18000 else 'FAILED'}")
    print(f"Estimated page count: ~{round(total_words/310)} pages (Target: 65-70 pages, >= 50 pages) -> {'PASSED' if total_words/310 >= 50 else 'FAILED'}")
    print(f"Major tables count: {total_tables} tables (Target: >= 6 tables) -> {'PASSED' if total_tables >= 6 else 'FAILED'}")
    print(f"Verified references: {ref_count} citations (Target: 35 citations with DOI/URL) -> {'PASSED' if ref_count >= 35 else 'FAILED'}")

    print("\n" + "="*80)
    print("BENIGN SCIENTIFIC TERMINOLOGY AUDITED")
    print("="*80)
    for b in benign_items:
        print(f"  [BENIGN APPROVED] {b['source']} ({b['location']}): \"{b['snippet']}\"")

    print("\n" + "="*80)
    if len(critical_items) == 0 and len(warning_items) == 0:
        verdict = "APPROVE"
        print(f"FINAL EMPIRICAL VERDICT: {verdict}")
        print("Empirical verification CONFIRMED: 0 occurrences of forbidden AI clichés,")
        print("0 sensational phrases, 0 emojis, 0 narrative exclamation marks.")
    else:
        verdict = "REQUEST_CHANGES"
        print(f"FINAL EMPIRICAL VERDICT: {verdict}")
        print(f"Detected {len(critical_items)} critical violations and {len(warning_items)} warnings.")
    print("="*80)

if __name__ == "__main__":
    main()
