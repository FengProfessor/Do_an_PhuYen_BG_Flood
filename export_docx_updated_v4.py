import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('Báo cáo Tổng hợp Tiến độ Dự án', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Mở bài + Bối cảnh
doc.add_heading('Mở bài và Bối cảnh: Kiến trúc Hybrid (BG_Flood + GeoAI)', level=1)
doc.add_paragraph('Báo cáo chi tiết tiến độ triển khai dự án mô phỏng lũ lụt BG_Flood và các bước chuẩn bị dữ liệu cho mô hình GeoAI, được tổng hợp theo từng tuần.')
doc.add_paragraph('Theo định hướng bài báo về ứng dụng AI trong dự báo ngập lụt, phương pháp tiếp cận Hybrid (Vật lý + AI) bao gồm các ý chính sau:')
doc.add_paragraph('Vấn đề mô hình truyền thống: BG_Flood rất chính xác nhưng tốn quá nhiều tài nguyên tính toán (chạy mất 24h).', style='List Bullet')
doc.add_paragraph('Giải pháp Hybrid: Kết hợp độ chính xác của vật lý và tốc độ của AI.', style='List Bullet')
doc.add_paragraph('Bước 1: Chạy mô hình vật lý BG_Flood hàng trăm lần để sinh dữ liệu (Training Set).', style='List Bullet 2')
doc.add_paragraph('Bước 2: Dùng dữ liệu đó để huấn luyện mạng nơ-ron (GeoAI).', style='List Bullet 2')
doc.add_paragraph('Bước 3: Dùng GeoAI để suy luận và xuất bản đồ ngập lụt theo thời gian thực (chỉ mất 1-2 phút).', style='List Bullet 2')

# Tiến độ
doc.add_heading('Tiến độ Dự án theo Tuần', level=1)

# Week 0
doc.add_heading('Tuần 0: Tiền chuẩn bị & Khởi tạo dự án', level=2)
doc.add_picture(r'D:\VTP\img_week0.png', width=Inches(5.5))
p = doc.add_paragraph('Những việc đã hoàn thành tốt:', style='List Bullet')
doc.add_paragraph('Thu thập dữ liệu nền tảng: Chuẩn bị thành công các file DEM 90m phủ trọn lưu vực Đắk Lắk - Phú Yên (18,000 km2).', style='List Bullet 2')
doc.add_paragraph('Cấu hình môi trường: Thiết lập xong thư mục dự án D:\VTP, tải và cài đặt thành công phần mềm lõi mô phỏng thủy lực BG_Flood (chạy trên GPU CUDA).', style='List Bullet 2')
doc.add_paragraph('Xác định kiến trúc & ranh giới: Chốt phương án đưa toàn bộ lượng mưa từ Đắk Lắk dồn về hệ thống Sông Ba và thiết lập cơ chế xả lũ vào Phú Yên.', style='List Bullet 2')
p = doc.add_paragraph('Khó khăn / Vấn đề gặp phải:', style='List Bullet')
doc.add_paragraph('Sự bỡ ngỡ ban đầu về cách cấu hình các tham số (parameters) phức tạp của mô hình BG_Flood (nhám Manning, biên xả, lượng mưa).', style='List Bullet 2')

# Week 1
doc.add_heading('Tuần 1: Khởi tạo mô hình và xử lý kịch bản diện rộng (V22)', level=2)
doc.add_picture(r'D:\VTP\img_week1.png', width=Inches(5.5))
doc.add_paragraph('Chi tiết Kịch bản V22 (Mô phỏng 18,000 km2):', style='List Bullet')
doc.add_paragraph('Vùng mô phỏng (Domain): Bao gồm Đắk Lắk (đóng vai trò là vùng sinh thủy - Catchment area, nhận nước mưa dồn vào sông Ba) và Phú Yên (vùng ngập lụt - Floodplain ở hạ lưu).', style='List Bullet 2')
doc.add_paragraph('Độ phân giải bản đồ: Lưới DEM 90m để có thể chạy được mô hình trên một diện tích khổng lồ mà không bị tràn RAM GPU.', style='List Bullet 2')
doc.add_paragraph('Thời gian mô phỏng: Chạy liên tục dòng chảy trong 120 giờ (5 ngày) để bao quát toàn bộ chu kỳ hình thành đỉnh lũ và nước rút.', style='List Bullet 2')
p = doc.add_paragraph('Những việc đã hoàn thành tốt:', style='List Bullet')
doc.add_paragraph('Hoàn thiện cấu hình tham số mô hình BG_Flood cho kịch bản V22 chạy trên GPU thành công.', style='List Bullet 2')
p = doc.add_paragraph('Khó khăn / Vấn đề gặp phải:', style='List Bullet')
doc.add_paragraph('File output định dạng .nc (NetCDF) ban đầu gặp lỗi rỗng dữ liệu do thiếu thư viện netCDF4 khi chạy Python script.', style='List Bullet 2')
doc.add_paragraph('Khó khăn trong việc đọc và đánh giá kết quả trực quan ngay lập tức trên phần mềm GIS do file quá nặng và lạ lẫm.', style='List Bullet 2')

# Week 2
doc.add_heading('Tuần 2: Trực quan hóa GIS và khắc phục lỗi tọa độ', level=2)
doc.add_picture(r'D:\VTP\img_week2.png', width=Inches(5.5))
doc.add_paragraph('Chi tiết các bước xử lý trực quan hóa:', style='List Bullet')
doc.add_paragraph('Xử lý File NetCDF (.nc): Viết script Python (export_v22_daily.py) để tự động đọc mảng dữ liệu mô phỏng 120 giờ (5 ngày).', style='List Bullet 2')
doc.add_paragraph('Trích xuất bản đồ tĩnh GeoTIFF (.tif): Nhóm dữ liệu theo từng block 24 giờ để lọc ra giá trị ngập sâu nhất (Max Depth) của mỗi ngày. Từ đó sinh ra 5 file bản đồ tĩnh tương ứng từ ngày 28/11/2021 đến 02/12/2021.', style='List Bullet 2')
doc.add_paragraph('Mục đích kép: File .tif nhẹ hơn rất nhiều, giúp QGIS render mượt mà không bị treo máy. Quan trọng hơn, đây chính là định dạng ma trận chuẩn (Image Matrix) bắt buộc phải có để nạp vào huấn luyện mạng nơ-ron tích chập (CNN GeoAI) trong tương lai.', style='List Bullet 2')
p = doc.add_paragraph('Giải quyết triệt để lỗi không khớp tọa độ (Coordinate Mismatch):', style='List Bullet')
doc.add_paragraph('Nguyên nhân: File DEM 90m lưu tọa độ mỏ neo ở "Góc" (Corner) của ô lưới, trong khi BG_Flood xuất file NetCDF lấy tọa độ mỏ neo ở "Tâm" (Center). Điều này gây ra hiện tượng bản đồ ngập lụt bị lệch 45m (nửa pixel) và lộn ngược trục Y khi đưa vào QGIS.', style='List Bullet 2')
doc.add_paragraph('Khắc phục: Bổ sung thuật toán lật ngược ma trận trục Y (np.flipud) và khai báo ép cứng hệ tọa độ Affine Transform (tịnh tiến về đúng hệ EPSG:32648). Kết quả: các file .tif đã khớp 100% từng pixel với bản đồ DEM gốc.', style='List Bullet 2')

# Week 3
doc.add_heading('Tuần 3: Đánh giá kiểm định mô hình, Thu nhỏ độ phân giải và Tự động hóa', level=2)
doc.add_paragraph('Những việc đang thực hiện (Dọn đường cho mô hình AI):', style='List Bullet')
doc.add_paragraph('Mô phỏng hạ lưu chi tiết (30m): Đang cấu hình kịch bản 30m cực kỳ chi tiết cho hạ lưu Phú Yên (từ đập Sông Ba Hạ ra Biển Đông) sử dụng Copernicus DEM 30m và ESA WorldCover 10m.', style='List Bullet 2')
doc.add_paragraph('Tự động hóa: Chuẩn bị các kịch bản chạy vòng lặp để sinh 100 kịch bản ngập lụt khác nhau, xây dựng kho Training Set cho GeoAI.', style='List Bullet 2')

doc.add_paragraph('PHÂN TÍCH ĐÁNH GIÁ CHẤT LƯỢNG MÔ HÌNH V23 SO VỚI THỰC TẾ (SAR):', style='Heading 3')
doc.add_picture(r'D:\VTP\BG_Flood_Project\BG_Flood\results\SAR_vs_Sim.png', width=Inches(5.5))
doc.add_paragraph('Giải nghĩa hình ảnh so sánh:', style='List Bullet')
doc.add_paragraph('Hình bên TRÁI (SAR - Vệ tinh): Thể hiện vết ngập lụt thực tế được chụp từ vệ tinh radar lúc 05:36 sáng ngày 30/11/2021. Vùng màu xanh là diện tích thực sự có nước ngập.', style='List Bullet 2')
doc.add_paragraph('Hình bên PHẢI (Sim - Mô phỏng): Thể hiện vết ngập do mô hình BG_Flood dự báo (mức nước > 0.3m).', style='List Bullet 2')
doc.add_paragraph('Bằng mắt thường, ta thấy mô hình (hình phải) đang cho ra vệt nước hẹp và "khô" hơn rất nhiều so với sự lan rộng của nước trên thực tế (hình trái).', style='List Bullet 2')

doc.add_paragraph('Phân tích Dữ liệu Kỹ thuật (Confusion Matrix):', style='List Bullet')
doc.add_paragraph('Kỹ thuật viên đã viết script bóc tách đối chiếu đối xứng từng điểm ảnh (Pixel-by-pixel) trên tổng số 4.5 triệu pixel của ma trận.', style='List Bullet 2')
doc.add_paragraph('Đúng dương tính (True Positive - TP = 5,177 pixels): Là vùng vệ tinh thấy ngập và mô hình cũng báo ngập. Tỷ lệ trùng khớp này quá nhỏ.', style='List Bullet 2')
doc.add_paragraph('Sai âm tính (False Negative - FN = 111,277 pixels): Là vùng thực tế bị ngập nhưng mô hình lại báo khô ráo (bỏ sót). Con số khổng lồ này chỉ ra rủi ro dự báo hụt (Under-prediction).', style='List Bullet 2')
doc.add_paragraph('Chỉ số CSI (Critical Success Index - Tới hạn thành công): 3.29% - Quá thấp so với chuẩn.', style='List Bullet 2')
doc.add_paragraph('Độ nhạy (Recall): 4.45% - Nghĩa là mô hình chỉ phát hiện được chưa tới 5% diện tích ngập thực tế.', style='List Bullet 2')
doc.add_paragraph('Lưu ý về Accuracy (Độ chính xác toàn cục = 96.68%): Đây là một cái bẫy thống kê (ảo giác). Accuracy cao ngất ngưởng hoàn toàn là do diện tích đồi núi khô ráo (True Negative) quá lớn đã lấn át toàn bộ kết quả, do đó trong bài toán ngập lụt chúng ta không dùng Accuracy để đánh giá.', style='List Bullet 2')

doc.add_paragraph('Đề xuất Hiệu chỉnh (Tuning) cho Vòng lặp tới:', style='List Bullet')
doc.add_paragraph('Lưu lượng xả: Kiểm tra lại biên Inflow, có thể dòng xả lũ hoặc lưu vực gom nước chưa đủ lớn.', style='List Bullet 2')
doc.add_paragraph('Lượng mưa: Đánh giá lại phân bổ ERA5.', style='List Bullet 2')
doc.add_paragraph('Hệ số cản (Manning): Tăng nhám ở hạ lưu để nước thoát ra biển chậm lại, giúp mực nước dâng cao hơn và lan rộng hơn.', style='List Bullet 2')

doc.save(r'D:\VTP\Bao_cao_Tien_do_v4.docx')
