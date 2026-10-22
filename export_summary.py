import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('Tóm tắt Bài báo Nghiên cứu: Ứng dụng Hybrid GeoAI & Mô hình Vật lý (BG_Flood)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Dưới đây là tóm tắt nội dung cốt lõi của bài báo gốc từ NIWA (New Zealand) về việc kết hợp sức mạnh của Trí tuệ Nhân tạo (Machine Learning) và Mô hình Thủy lực (Physics-based) để dự báo ngập lụt.', style='Normal')

# 1. Vấn đề giải quyết
doc.add_heading('1. Vấn đề Đặt ra', level=1)
doc.add_paragraph('Các mô hình thủy lực vật lý truyền thống (như BG_Flood) tốn quá nhiều thời gian tính toán. Thường mất đến 24 giờ để chạy xong một kịch bản ngập lụt diện rộng. Tốc độ này quá chậm, không đáp ứng được yêu cầu cảnh báo sớm và ứng phó khẩn cấp khi thiên tai thực sự xảy ra.', style='List Bullet')

# 2. Giải pháp
doc.add_heading('2. Giải pháp Hybrid (Vật lý + AI)', level=1)
doc.add_paragraph('Cách tiếp cận: Các nhà khoa học NIWA sử dụng mô hình Mạng Nơ-ron Tích chập (Convolutional Neural Network - CNN).', style='List Bullet')
doc.add_paragraph('Quy trình huấn luyện (Training): Đầu tiên, họ tạo ra một "thư viện" khổng lồ các kịch bản lũ lụt mẫu bằng cách chạy mô hình vật lý BG_Flood (phần việc nặng nhọc này làm trước từ lúc rảnh rỗi). Sau đó, họ dùng kho dữ liệu này để "dạy" cho AI cách nước chảy.', style='List Bullet')
doc.add_paragraph('Giai đoạn Thực tế (Real-time Inference): Khi có bão thực tế, mô hình AI chỉ cần nạp dữ liệu dự báo về mực nước biển và lưu lượng sông của 8-12 giờ trước đó. AI sẽ ngay lập tức "vẽ" ra bản đồ ngập lụt.', style='List Bullet')

# 3. Số liệu Đột phá
doc.add_heading('3. Những Con số Đột phá', level=1)
doc.add_paragraph('Tốc độ: Thời gian để có được bản đồ ngập lụt giảm đột phá từ 24 giờ xuống chỉ còn 1 - 2 phút.', style='List Bullet')
doc.add_paragraph('Sai số độ sâu (Error): Trung bình chỉ sai lệch dưới 0.1 mét (10cm) tại mỗi ô lưới bản đồ.', style='List Bullet')
doc.add_paragraph('Độ chính xác vùng ngập (Accuracy): AI có khả năng dự đoán đúng diện tích vùng bị ngập với độ chính xác lên tới 95%.', style='List Bullet')

# 4. Bước tiếp theo
doc.add_heading('4. Hướng phát triển tiếp theo', level=1)
doc.add_paragraph('Vận hành thực tế (Operationalise): Áp dụng hệ thống AI này vào dự báo trực tiếp cho các cơn lũ thực tế tại Westport, đồng thời nhân rộng mô hình ra các khu vực có nguy cơ cao khác tại New Zealand.', style='List Bullet')
doc.add_paragraph('Đánh giá rủi ro (Impact Forecasting): Tích hợp bản đồ ngập lụt vừa xuất ra vào phần mềm RiskScape để tự động tính toán ra con số thiệt hại ước tính về tài sản, số người cần sơ tán và các tuyến đường giao thông bị chia cắt.', style='List Bullet')

doc.save(r'D:\VTP\Tom_tat_Bai_bao_GeoAI.docx')
