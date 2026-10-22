import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('Báo cáo Tổng hợp Tiến độ Dự án', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Mở bài + Bối cảnh
doc.add_heading('Mở bài và Bối cảnh: Kiến trúc Hybrid (BG_Flood + GeoAI)', level=1)
doc.add_paragraph('Báo cáo chi tiết tiến độ triển khai dự án mô phỏng lũ lụt BG_Flood và các bước chuẩn bị dữ liệu cho mô hình GeoAI, được tổng hợp theo từng tuần.')
doc.add_paragraph('Theo định hướng bài báo về ứng dụng AI trong dự báo ngập lụt, phương pháp tiếp cận Hybrid (Vật lý + AI) bao gồm các ý chính sau:')
doc.add_paragraph('Vấn đề mô hình truyền thống: BG_Flood rất chính xác nhưng tốn quá nhiều tài nguyên tính toán (chạy mất 24h).', style='List Bullet')
doc.add_paragraph('Giải pháp Hybrid: Kết hợp độ chính xác của vật lý và tốc độ của AI.', style='List Bullet')
doc.add_paragraph('Bước 1: Chạy mô hình vật lý BG_Flood hàng trăm lần để sinh dữ liệu (Training Set).', style='List Bullet 2')
doc.add_paragraph('Bước 2: Dùng dữ liệu đó để huấn luyện mạng nơ-ron (GeoAI).', style='List Bullet 2')
doc.add_paragraph('Bước 3: Dùng GeoAI để suy luận và xuất bản đồ ngập lụt theo thời gian thực (chỉ mất 1-2 phút).', style='List Bullet 2')

# Tiến độ
doc.add_heading('Tiến độ Dự án theo Tuần', level=1)

# Week 0
doc.add_heading('Tuần 0: Tiền chuẩn bị & Khởi tạo dự án', level=2)
doc.add_picture(r'D:\VTP\img_week0.png', width=Inches(5.5))
p = doc.add_paragraph('Những việc đã hoàn thành tốt:', style='List Bullet')
doc.add_paragraph('Thu thập dữ liệu nền tảng: Chuẩn bị thành công các file DEM 90m phủ trọn lưu vực Đắk Lắk - Phú Yên (18,000 km2).', style='List Bullet 2')
doc.add_paragraph('Cấu hình môi trường: Thiết lập xong thư mục dự án D:\VTP, tải và cài đặt thành công phần mềm lõi mô phỏng thủy lực BG_Flood (chạy trên GPU CUDA).', style='List Bullet 2')
doc.add_paragraph('Xác định kiến trúc & ranh giới: Chốt phương án đưa toàn bộ lượng mưa từ Đắk Lắk dồn về hệ thống Sông Ba và thiết lập cơ chế xả lũ vào Phú Yên.', style='List Bullet 2')
p = doc.add_paragraph('Khó khăn / Vấn đề gặp phải:', style='List Bullet')
doc.add_paragraph('Sự bỡ ngỡ ban đầu về cách cấu hình các tham số (parameters) phức tạp của mô hình BG_Flood (nhám Manning, biên xả, lượng mưa).', style='List Bullet 2')

# Week 1
doc.add_heading('Tuần 1: Khởi tạo mô hình và xử lý kịch bản diện rộng (V22)', level=2)
doc.add_picture(r'D:\VTP\img_week1.png', width=Inches(5.5))
doc.add_paragraph('Chi tiết Kịch bản V22 (Mô phỏng 18,000 km2):', style='List Bullet')
doc.add_paragraph('Vùng mô phỏng (Domain): Bao gồm Đắk Lắk (đóng vai trò là vùng sinh thủy - Catchment area, nhận nước mưa dồn vào sông Ba) và Phú Yên (vùng ngập lụt - Floodplain ở hạ lưu).', style='List Bullet 2')
doc.add_paragraph('Độ phân giải bản đồ: Lưới DEM 90m để có thể chạy được mô hình trên một diện tích khổng lồ mà không bị tràn RAM GPU.', style='List Bullet 2')
doc.add_paragraph('Thời gian mô phỏng: Chạy liên tục dòng chảy trong 120 giờ (5 ngày) để bao quát toàn bộ chu kỳ hình thành đỉnh lũ và nước rút.', style='List Bullet 2')
p = doc.add_paragraph('Những việc đã hoàn thành tốt:', style='List Bullet')
doc.add_paragraph('Hoàn thiện cấu hình tham số mô hình BG_Flood cho kịch bản V22 chạy trên GPU thành công.', style='List Bullet 2')
p = doc.add_paragraph('Khó khăn / Vấn đề gặp phải:', style='List Bullet')
doc.add_paragraph('File output định dạng .nc (NetCDF) ban đầu gặp lỗi rỗng dữ liệu do thiếu thư viện netCDF4 khi chạy Python script.', style='List Bullet 2')
doc.add_paragraph('Khó khăn trong việc đọc và đánh giá kết quả trực quan ngay lập tức trên phần mềm GIS do file quá nặng và lạ lẫm.', style='List Bullet 2')

# Week 2
doc.add_heading('Tuần 2: Trực quan hóa GIS và khắc phục lỗi tọa độ', level=2)
doc.add_picture(r'D:\VTP\img_week2.png', width=Inches(5.5))
p = doc.add_paragraph('Những việc đã hoàn thành tốt:', style='List Bullet')
doc.add_paragraph('Đã viết thành công script tự động (export_v22_daily.py) để trích xuất file .nc thành 5 file .tif tĩnh (thể hiện mức ngập sâu nhất của từng ngày).', style='List Bullet 2')
doc.add_paragraph('Việc dùng file .tif giúp QGIS render mượt mà, nhẹ máy, và là định dạng chuẩn để đưa vào huấn luyện AI sau này.', style='List Bullet 2')
doc.add_paragraph('Giải quyết triệt để lỗi lệch tọa độ bằng cách thêm hàm lật trục Y và ép cứng affine transform theo hệ tọa độ EPSG:32648.', style='List Bullet 2')
p = doc.add_paragraph('Khó khăn / Vấn đề gặp phải:', style='List Bullet')
doc.add_paragraph('Bị nhầm lẫn ban đầu về đặc tính của file .tif (tưởng .tif có thể play xem dạng video time-series như file .nc).', style='List Bullet 2')
doc.add_paragraph('Lỗi bản đồ: Khi ném file .nc vào QGIS, bản đồ bị lệch khoảng 45m (nửa pixel) và lộn ngược so với file DEM 90m (do đặc thù file .nc lấy tọa độ ở tâm ô lưới).', style='List Bullet 2')

# Week 3
doc.add_heading('Tuần 3: (Hiện tại) Đánh giá mô hình, Thu nhỏ độ phân giải và Tự động hóa', level=2)
doc.add_paragraph('Những việc đang thực hiện:', style='List Bullet')
doc.add_paragraph('Mô phỏng hạ lưu chi tiết (30m): Đang cấu hình kịch bản 30m cực kỳ chi tiết cho hạ lưu Phú Yên sử dụng Copernicus DEM 30m và ESA WorldCover 10m.', style='List Bullet 2')
doc.add_paragraph('Chuẩn bị Dữ liệu GeoAI: Viết các kịch bản sinh dữ liệu tự động cho quá trình huấn luyện máy học tương lai (chạy 100 kịch bản).', style='List Bullet 2')
doc.add_paragraph('Đánh giá kiểm định mô hình: Đã chạy script so sánh đối soát độ chính xác của mô phỏng (V23) với ảnh vệ tinh thực tế (SAR). Chi tiết như sau:', style='List Bullet 2')

# Evaluation moved here under week 3
doc.add_picture(r'D:\VTP\BG_Flood_Project\BG_Flood\results\SAR_vs_Sim.png', width=Inches(5.5))
doc.add_paragraph('Đánh giá độ chính xác (SAR vs BG_Flood):', style='List Bullet')
doc.add_paragraph('Đúng dương tính (TP): 5,177 pixels', style='List Bullet 2')
doc.add_paragraph('Sai âm tính (FN - bỏ sót): 111,277 pixels', style='List Bullet 2')
doc.add_paragraph('Chỉ số CSI (Critical Success Index): 3.29%', style='List Bullet 2')
doc.add_paragraph('Độ nhạy (Recall): 4.45%', style='List Bullet 2')
doc.add_paragraph('Phân tích và Đề xuất:', style='List Bullet')
doc.add_paragraph('Mô hình đang dự báo trượt (Under-prediction) khá nhiều so với vệ tinh.', style='List Bullet 2')
doc.add_paragraph('Cần hiệu chỉnh lại: lưu lượng xả biên, lượng mưa ERA5, và tăng hệ số nhám Manning ở vùng hạ lưu.', style='List Bullet 2')

doc.save(r'D:\VTP\Bao_cao_Tien_do_v3.docx')
