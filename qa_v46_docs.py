from pathlib import Path
import re
from PIL import Image, ImageChops, ImageStat, ImageDraw
from docx import Document

ROOT = Path(r"D:\VTP")

def page_num(p: Path):
    m = re.search(r"(\d+)", p.stem)
    return int(m.group(1)) if m else 0

def inspect_doc(path: Path):
    d = Document(path)
    print(path.name, "paras", len(d.paragraphs), "tables", len(d.tables), "images", len(d.inline_shapes))
    for p in d.paragraphs:
        txt = p.text.strip()
        if txt.startswith(("2.3.4", "3.7.", "3.8.", "[36]", "[37]", "[38]", "[39]", "[40]")):
            print("  ", txt[:240])

def make_sheets(render_dir: Path, out_dir: Path, pages_per_sheet=4):
    out_dir.mkdir(parents=True, exist_ok=True)
    pages = sorted(render_dir.glob("page-*.png"), key=page_num)
    if not pages:
        raise RuntimeError(f"No pages in {render_dir}")
    first = Image.open(pages[0]).convert("RGB")
    w, h = first.size
    print(render_dir.name, "pages", len(pages), "page_size", (w, h))
    for idx in range(0, len(pages), pages_per_sheet):
        group = pages[idx:idx+pages_per_sheet]
        sheet = Image.new("RGB", (w*2, h*2), "white")
        draw = ImageDraw.Draw(sheet)
        for j, p in enumerate(group):
            im = Image.open(p).convert("RGB")
            x = (j % 2) * w
            y = (j // 2) * h
            sheet.paste(im, (x, y))
            draw.rectangle([x+8, y+8, x+170, y+42], fill="white")
            draw.text((x+14, y+12), f"PAGE {page_num(p)}", fill="black")
        start = page_num(group[0]); end = page_num(group[-1])
        out = out_dir / f"sheet_{start:02d}_{end:02d}.jpg"
        sheet.save(out, quality=88, optimize=True)
        print(out)

if __name__ == "__main__":
    inspect_doc(ROOT / "Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_V46_bo_sung_20260930.docx")
    inspect_doc(ROOT / "Giai_thich_quy_trinh_V46_de_hieu_20260930.docx")
    make_sheets(ROOT / "_render_v46_thesis", ROOT / "_qa_v46_thesis_sheets")
    make_sheets(ROOT / "_render_v46_explain", ROOT / "_qa_v46_explain_sheets")
