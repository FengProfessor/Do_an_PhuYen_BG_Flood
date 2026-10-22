"""Apply the reviewed corrections to a new thesis copy using bundled Python."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(r"D:\VTP")
WORK = ROOT / "Do_an_edit_20261005"
SOURCE = ROOT / "Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx"
DEST = ROOT / "Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_chinh_sua_20261005.docx"
EVIDENCE = json.loads((ROOT / "Do_an_research_20261005/project_evidence_audit.json").read_text(encoding="utf-8"))
SOURCE_HASH = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
pp = list(doc.paragraphs)
tt = list(doc.tables)
changes = []
removed = set()

def setp(index, text):
    pp[index].text = text
    changes.append(index)

def drop(index):
    if index not in removed:
        pp[index]._p.getparent().remove(pp[index]._p)
        removed.add(index)

def replace_range(start, end, texts):
    assert len(texts) <= end-start+1
    for j, text in enumerate(texts):
        setp(start+j, text)
    for j in range(start+len(texts), end+1):
        drop(j)

def clean_md(text):
    def link(m):
        url = m.group(2)
        number = 5 if "4310076" in url else 6 if "2011.03.051" in url else 10 if "JRC119812" in url else None
        return f"[{number}]" if number else m.group(1)
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, text).replace("**", "").replace("`", "")

draft = (ROOT / "Do_an_research_20261005/Ban_viet_lai_cac_doan_trong_yeu_20261005.md").read_text(encoding="utf-8")
sections = {}
for match in re.finditer(r"^## (\d+)\. .+\n([\s\S]*?)(?=^## |\Z)", draft, re.M):
    sections[int(match.group(1))] = match.group(2).strip()
summary = [clean_md(p) for p in sections[1].split("\n\n") if not p.startswith("**Từ khóa")]
abstract = [clean_md(p) for p in sections[2].split("\n\n") if not p.startswith("**Keywords")]
replace_range(22,24,summary)
replace_range(28,30,abstract)
setp(25,"Từ khóa: mô hình thủy lực hai chiều; BG_Flood; SAR; Phú Yên; phạm vi ngập; hiệu chỉnh địa hình.")
setp(31,"Keywords: two-dimensional hydraulic modelling; BG_Flood; SAR; Phu Yen; flood extent; terrain modification.")

replacements = {
1:"",
3:"ỨNG DỤNG VIỄN THÁM GIS VÀ MÔ HÌNH THỦY LỰC SỐ\nĐỂ MÔ PHỎNG VÀ ĐÁNH GIÁ PHẠM VI NGẬP LỤT\nTẠI HẠ LƯU SÔNG BA PHÚ YÊN",
4:"Sinh viên thực hiện: [Họ và tên sinh viên]\nMã số sinh viên: [Mã số sinh viên]\nKhóa đào tạo: QH-2022-I/CQ\nCán bộ hướng dẫn: TS. Hà Minh Cường\nCơ quan công tác: Viện Công nghệ Hàng không Vũ trụ",
8:"Tôi xin cam đoan đồ án này được thực hiện dưới sự hướng dẫn của TS. Hà Minh Cường. Các kết quả mô phỏng, phép đối chiếu và giới hạn phương pháp được trình bày theo hồ sơ nghiên cứu; những dữ liệu và nhận định kế thừa được dẫn nguồn trong danh mục tài liệu tham khảo.",
9:"Tôi chịu trách nhiệm về nội dung đồ án và việc sử dụng các tài liệu, dữ liệu được trích dẫn. Các kết quả thực nghiệm được phân biệt với giả thuyết và hướng nghiên cứu tiếp theo.",
13:"Trong quá trình thực hiện đồ án, em đã nhận được sự hướng dẫn và hỗ trợ của thầy cô, gia đình và bạn bè.",
14:"Em xin trân trọng cảm ơn TS. Hà Minh Cường, Viện Công nghệ Hàng không Vũ trụ, Trường Đại học Công nghệ, Đại học Quốc gia Hà Nội, đã hướng dẫn và góp ý về phương pháp nghiên cứu, xử lý dữ liệu viễn thám và mô hình hóa thủy lực.",
16:"Em cảm ơn các thầy cô và các bạn đã trao đổi chuyên môn, chia sẻ tài liệu và hỗ trợ trong quá trình nghiên cứu.",
64:"Thiên tai liên quan đến thời tiết, khí hậu và nước gây tổn thất lớn về người và tài sản. Theo WMO, giai đoạn 1970–2019 ghi nhận hơn 11.000 thiên tai thuộc nhóm này, hơn 2 triệu người tử vong và thiệt hại khoảng 3,64 nghìn tỷ USD [1]. Đây là thống kê tổng hợp nhiều loại thiên tai, không riêng lũ lụt. Trong nghiên cứu ngập lụt, việc mô tả phạm vi nước tràn và độ bất định của kết quả có ý nghĩa đối với đánh giá rủi ro.",
65:"Việt Nam có nhiều lưu vực ven biển chịu ảnh hưởng của mưa lớn và lũ sông. Địa hình chuyển nhanh từ vùng núi sang đồng bằng, cùng các yếu tố vận hành hồ chứa và tiêu thoát nước hạ lưu, đặt ra yêu cầu lựa chọn dữ liệu và mô hình phù hợp cho từng khu vực.",
66:"Khu vực hạ lưu sông Ba và đồng bằng Tuy Hòa được lựa chọn để khảo sát khả năng mô phỏng phạm vi ngập bằng mô hình hai chiều trên lưới 30 m. Miền tính còn chứa các lưu vực lân cận; vì vậy, kết quả cần được phân biệt với phạm vi riêng của lưu vực sông Ba và phạm vi hành chính toàn tỉnh.",
67:"Đợt mưa lũ cuối tháng 11/2021 là sự kiện nghiên cứu của đồ án. Báo cáo Thiên tai Việt Nam 2021 ghi nhận lưu lượng xả hồ Sông Ba Hạ đạt 9.400 m³/s tại 15 giờ ngày 30/11 và mực nước tại Phú Lâm đạt 4,63 m [2]. Những số liệu này là bối cảnh quan trắc, cần được phân biệt với các hydrograph đã điều chỉnh trong bộ kịch bản mô hình.",
68:"Bài toán đặt ra là thiết lập quy trình mô phỏng có thể truy vết dữ liệu, đối chiếu phạm vi ngập tại thời điểm phù hợp và phân tích nguyên nhân sai khác. Mô hình hai chiều mô tả dòng tràn trên đồng bằng, nhưng yêu cầu dữ liệu địa hình, forcing và phương pháp số nhất quán. Chi phí tính toán phụ thuộc miền tính, độ phân giải, phần cứng và bộ giải; đồ án không có phép chạy CPU đối chứng để xác định hệ số gia tốc GPU.",
69:"BG_Flood là mô hình mã nguồn mở phục vụ tính toán thủy lực hai chiều và hỗ trợ thực thi trên GPU [3]. Ảnh Sentinel-1 SAR cung cấp thông tin bề mặt trong điều kiện có mây [9]. Việc kết hợp hai nguồn cho phép đối chiếu không gian, với điều kiện ghi rõ sai số quan sát và mức độc lập giữa dữ liệu hiệu chỉnh với dữ liệu đánh giá.",
70:"Trên lưới 30 m, kết nối dòng chảy qua cầu, cống hoặc những dải địa hình hẹp có thể được biểu diễn chưa đầy đủ. Hiệu chỉnh cao độ là một giả thuyết kỹ thuật để khảo sát vấn đề này. Cần kiểm tra phạm vi và mức thay đổi, thay vì mặc nhiên coi mọi ô nhô cao là vật cản giả hoặc quy toàn bộ bỏ sót ngập cho DEM.",
71:"Đề tài nghiên cứu quy trình ứng dụng viễn thám, GIS và mô hình thủy lực số để mô phỏng và đánh giá phạm vi ngập lụt tại hạ lưu sông Ba, Phú Yên, với trọng tâm là kết quả V44 và các giới hạn của phép đối chiếu.",
73:clean_md(sections[3].split("\n\n")[0]),
74:"Các mục tiêu cụ thể gồm:",
75:"1. Hệ thống hóa dữ liệu địa hình, lớp phủ, mưa, lưu lượng và SAR; ghi rõ nguồn, đơn vị, hệ tọa độ và cách xử lý.",
76:"2. Thiết lập mô hình, thực hiện mô phỏng 120 giờ và kiểm tra cấu trúc, thời gian, phạm vi không gian của kết quả.",
77:"3. Đối chiếu phạm vi ngập với SAR tại thời điểm phù hợp, sử dụng cùng miền hợp lệ và các chỉ số CSI, Recall, Precision.",
78:"4. Phân tích phân bố sai khác và vai trò có thể có của kết nối địa hình; đề xuất thí nghiệm có đối chứng và dữ liệu bổ sung.",
84:"Các đặc trưng địa hình, tham số nhám và tổn thất, cùng ảnh hưởng có thể có của kết nối qua công trình tới dòng tràn.",
85:"Phạm vi không gian: miền raster có mép ngoài X=[270.000; 333.000] m, Y=[1.415.000; 1.478.000] m trong hệ WGS84 UTM vùng 49 Bắc (EPSG:32649), diện tích hình chữ nhật khoảng 3.969 km². Miền bao gồm khu vực hạ lưu sông Ba và các vùng lân cận; đây không phải diện tích ngập hoặc diện tích lưu vực sông Ba.",
86:"Phạm vi thời gian: kịch bản V44 có thời gian mô hình 432.000 s, tương đương 120 giờ. Bộ đánh giá sử dụng lớp SAR mang mốc 05:36 ngày 30/11/2021 và chọn frame 33 tại 367.200,28125 s, khoảng 102 giờ. Nếu lấy epoch 00:00 ngày 26/11 như quy ước đồ án, ảnh ứng với 101,6 giờ và lệch khung khoảng 24 phút; cần xác nhận bằng metadata scene và epoch trước khi diễn giải thời gian tuyệt đối.",
89:"Mô hình dựa trên vật lý: sử dụng hệ phương trình nước nông hai chiều để mô tả dòng chảy trung bình theo chiều sâu. Độ đúng của phương trình lý thuyết cần được phân biệt với chất lượng dữ liệu và triển khai số.",
90:"Đối chiếu quan sát: sử dụng sản phẩm SAR làm lớp ngập tham chiếu. Khi SAR đã tham gia hiệu chỉnh địa hình, phép chấm điểm với cùng lớp chưa phải kiểm định độc lập.",
91:"Phân tích không gian: kiểm tra đồng đăng ký raster, miền hợp lệ và phân bố ngập theo chênh cao tới mạng thoát nước.",
92:"Thực nghiệm số: so sánh các nhánh chỉ khi cùng forcing, runtime, thời điểm và quy tắc đánh giá, hoặc ghi rõ những khác biệt làm hạn chế diễn giải.",
95:"Thiết lập quy trình mô phỏng và đối chiếu có hồ sơ dữ liệu, cấu hình và các chỉ số có thể tái lập.",
96:"Nhận diện giới hạn của hiệu chỉnh địa hình dựa trên SAR và xây dựng giả thuyết về kết nối dòng chảy để thử nghiệm có đối chứng.",
97:"Phân biệt độ khớp phạm vi ngập với kiểm định độc lập, đồng thời sử dụng chẩn đoán địa hình như bằng chứng bổ trợ.",
99:"Cung cấp bộ kết quả mô phỏng và phân tích sai khác để phục vụ nghiên cứu tiếp theo về ngập lụt tại khu vực.",
100:"Khả năng ứng dụng trong quy hoạch hoặc cảnh báo cần được đánh giá thêm bằng quan trắc độc lập, nhiều sự kiện và dữ liệu dự báo; chưa xem kết quả hiện tại là công cụ vận hành đã nghiệm thu.",
104:"Chương 1 trình bày tổng quan mô hình ngập lụt, cơ sở SAR và dữ liệu địa hình, cùng các nghiên cứu liên quan tại khu vực.",
105:"Chương 2 mô tả hệ phương trình nước nông, phương pháp thể tích hữu hạn và nhánh central-upwind Kurganov–Petrova thực dùng, dữ liệu đầu vào, hiệu chỉnh địa hình và quy tắc đối chiếu.",
106:"Chương 3 trình bày cấu hình V44, kết quả tái lập tại 102 giờ và chuỗi 120 giờ, chẩn đoán địa hình, các giới hạn và thiết kế nghiên cứu V45/V48.",
108:"Tài liệu tham khảo gồm các nguồn được sử dụng trong thân bài; phụ lục ghi các tệp cấu hình, dữ liệu và kết quả phục vụ truy vết.",
113:"Mô hình ngập lụt mô tả dòng chảy trong lòng dẫn và sự trao đổi nước với bãi tràn. Các cách tiếp cận thường dùng gồm mô hình một chiều, mô hình kết nối một–hai chiều và mô hình hai chiều. Lựa chọn phụ thuộc mục tiêu, đặc điểm địa hình và dữ liệu; không có một cấu hình ưu việt cho mọi bài toán [4], [31].",
114:"Mô hình thủy lực một chiều",
115:"Mô hình một chiều giải dòng chảy dọc mạng sông và sử dụng dữ liệu mặt cắt để biểu diễn tiết diện. Cách tiếp cận này phù hợp với bài toán truyền lũ trong lòng dẫn. Khi nước tràn theo nhiều hướng trên đồng bằng, cần bổ sung cấu trúc bãi tràn hoặc kết nối với miền hai chiều; giới hạn này không tự xác định mức sai số của một mô hình cụ thể [40].",
116:"Mô hình kết nối một và hai chiều",
117:"Mô hình kết nối dùng miền một chiều cho sông và miền hai chiều cho bãi tràn, trao đổi nước qua các đường hoặc công trình liên kết. Chất lượng kết quả phụ thuộc dữ liệu mặt cắt, cao trình trao đổi và việc thiết lập liên kết. Đây là một lựa chọn có thể giảm chi phí mô tả lòng dẫn so với việc đưa mọi cấu trúc vào lưới hai chiều [40].",
118:"Mô hình thủy lực hai chiều",
119:"Mô hình hai chiều mô tả vận tốc theo hai phương ngang và có thể khảo sát dòng tràn trên địa hình phức tạp. Chi phí tính tăng khi số ô lớn hoặc bước thời gian bị giới hạn; giá trị thời gian chạy cần được đo trên từng cấu hình. Không suy ra một số ngày tính toán chung cho mọi bộ giải CPU [4], [31].",
120:"Thực thi trên GPU",
121:"GPU cung cấp khả năng xử lý song song nhiều phép tính ô lưới. Vacondio và cộng sự trình bày bộ giải thể tích hữu hạn tăng tốc bằng GPU [24]; BG_Flood cũng hỗ trợ cách tổ chức tính toán trên GPU [3]. Hiệu suất phụ thuộc thuật toán và phần cứng. Đồ án báo cáo thời gian solver của V44, không chuyển hệ số tăng tốc của công trình khác thành kết quả đo của dự án.",
123:"Các nghiên cứu tại Việt Nam cho thấy cần kết hợp mô phỏng với dữ liệu quan trắc phù hợp khu vực. Đối với hạ lưu sông Ba, công trình có dữ liệu nguồn cụ thể là nghiên cứu của Lương Hữu Dũng và cộng sự năm 2024 [40].",
124:"Nghiên cứu này sử dụng bộ mô hình HEC để mô phỏng hạ lưu sông Ba, kết nối mô hình một–hai chiều và đánh giá bằng mực nước quan trắc cùng vết lũ [40]. Cách tổ chức dữ liệu đó là một điểm tham khảo khi xây dựng kế hoạch kiểm định độc lập cho BG_Flood.",
125:"Đồ án tập trung vào kết quả trên lưới 30 m và đối chiếu SAR trong bộ dữ liệu dự án. Việc so với một nghiên cứu khác cần cùng mục tiêu, miền và chỉ số; không so trực tiếp sai số mực nước tại trạm với CSI của mặt nạ ngập để xếp hạng mô hình.",
126:"Khoảng trống cần xử lý là tính nhất quán của dữ liệu và runtime, quan hệ giữa hiệu chỉnh với đánh giá, cùng khả năng phân tích nguyên nhân sai khác. Các nội dung này quyết định phạm vi kết luận có thể rút ra từ thí nghiệm hiện có.",
134:"Tại khu vực nghiên cứu, mưa, lưu lượng sông và điều kiện tiêu thoát nước có thể cùng ảnh hưởng tới ngập. Tuy nhiên, việc định lượng riêng tác động của triều hoặc vận hành hồ cần dữ liệu tương ứng và thiết kế mô hình; không suy ra từ một ảnh phạm vi ngập.",
136:"Sông Ba có vùng thượng lưu và trung lưu trước khi chảy qua đồng bằng Tuy Hòa ra cửa Đà Diễn. Khi nghiên cứu hạ lưu, cần xác định rõ mạng sông và các vị trí nguồn. Sông Hinh nối vào hệ thống sông Ba, trong khi Bàn Thạch là hệ thống lân cận chảy tới cửa Đà Nông; không đồng nhất hai nguồn này [40].",
138:"Địa hình chuyển từ vùng cao tới đồng bằng có thể làm thay đổi độ dốc và hình thái dòng tràn. Đối với mô hình, những thay đổi này được biểu diễn thông qua cao độ ô lưới và kết nối giữa các ô; sai số cao độ có thể ảnh hưởng cả hướng chảy và độ sâu.",
139:"Mưa là đầu vào biến đổi theo không gian và thời gian. Chuỗi ERA5 thuộc dữ liệu tái phân tích, không phải tập trạm đo độc lập. Việc tổng hợp mưa trên miền hoặc nhân hệ số cần được mô tả và kiểm tra bằng dữ liệu phù hợp [11].",
140:"Lưu lượng sau hồ chứa cần được mô tả bằng chuỗi theo thời gian và đúng vị trí trong mạng sông. Cần phân biệt lưu lượng đến hồ, qua tràn, qua tua-bin và lưu lượng tổng, đồng thời ghi rõ những bước điều chỉnh khi tạo hydrograph mô hình.",
143:"Viễn thám thu nhận thông tin bề mặt từ xa. Ảnh quang học phụ thuộc điều kiện chiếu sáng và mây; SAR là quan sát vi ba chủ động, hữu ích khi theo dõi bề mặt trong nhiều điều kiện thời tiết [9].",
145:"SAR kết hợp tín hiệu biên độ và pha thu tại nhiều vị trí dọc quỹ đạo để cải thiện độ phân giải theo phương vị. Độ phân giải thực và kích thước pixel lấy mẫu là hai thuộc tính khác nhau; cần sử dụng thông số của đúng chế độ và sản phẩm [9].",
149:"Mặt nước tương đối nhẵn có thể phản xạ phần lớn năng lượng ra xa cảm biến, tạo tín hiệu tán xạ ngược thấp. Tuy nhiên, giá trị dB phụ thuộc góc tới, phân cực, gió và điều kiện bề mặt; không áp dụng một khoảng tín hiệu cố định cho mọi vùng ngập [10].",
151:"Độ nhám và cấu trúc lớp phủ làm thay đổi tán xạ ngược của đất cạn. Một số bề mặt cạn cũng có tín hiệu thấp, trong khi nước có thể có tín hiệu cao; vì vậy, sự tương phản nước–đất cần được đánh giá trong từng ảnh [10].",
153:"Tán xạ kép giữa mặt nước và thân cây hoặc công trình có thể làm tăng tín hiệu tại vùng ngập. Cơ chế này là một lý do ngập dưới thực vật hoặc trong đô thị có thể khó nhận diện bằng một ngưỡng cường độ thấp [10].",
155:"SAR có thể bỏ sót hoặc phân loại nhầm ngập do tán thực vật, bóng radar, công trình và mặt nước bị gió làm nhám [10]. Trong ma trận mô hình–SAR, bỏ sót nước của SAR có thể tạo ô mô hình ngập/SAR cạn; cần phân biệt với tên FN khi xem SAR là nhãn tham chiếu. Không dùng một hạn chế quan sát để giải thích mọi nhóm bất đồng.",
157:"Ở thời kỳ sự kiện năm 2021, nhiệm vụ Sentinel-1 cung cấp dữ liệu SAR băng C với các chế độ thu nhận và chu kỳ lặp phụ thuộc cấu hình quan sát. Khi lựa chọn ảnh, cần kiểm tra vệ tinh, quỹ đạo, thời gian UTC và chế độ của đúng scene [9].",
158:"Chế độ IW và các kênh VV/VH thường được sử dụng cho quan sát đất liền. Sản phẩm GRD biểu diễn dữ liệu cường độ; các thông số lấy mẫu, độ phân giải và hiệu chỉnh cần được đọc từ hồ sơ sản phẩm, không suy ra từ kích thước raster đã đưa về lưới mô hình [9].",
159:"Lựa chọn VV hoặc VH cần căn cứ dữ liệu, mặt phủ và phép đánh giá. Đồ án chưa có so sánh độc lập hai phân cực để kết luận VV luôn tốt hơn VH. Metadata scene và chuỗi xử lý của sản phẩm SAR được xem là nội dung cần hoàn thiện.",
161:"Otsu là phương pháp chọn ngưỡng từ histogram bằng cách tối đa hóa phương sai giữa các lớp [5]. Với SAR, kết quả có thể được dùng để xây dựng một lớp nước tham chiếu; độ đúng phân loại cần được đánh giá riêng.",
162:"Thuật toán chia histogram thành hai lớp theo từng ngưỡng ứng viên và chọn ngưỡng có phương sai giữa lớp lớn nhất. Phương pháp không đòi hỏi hai phân phối Gaussian và không bảo đảm rằng hai lớp tìm được tương ứng hoàn toàn với nước ngập và đất cạn [5].",
170:"Otsu có thể tính tự động sau khi đã lựa chọn miền dữ liệu và quy tắc tiền xử lý. Tuy vậy, histogram, NODATA, mặt phủ và mặt nạ nước thường trực đều ảnh hưởng kết quả. Ngưỡng tối ưu thống kê cần được phân biệt với ngưỡng có độ chính xác ngập cao nhất ngoài thực địa.",
175:"DEM là thuật ngữ chung cho biểu diễn độ cao dạng số; tùy sản phẩm, nó có thể phản ánh bề mặt hoặc địa hình. DTM hướng tới mặt đất, còn DSM bao gồm ảnh hưởng của vật thể trên bề mặt. Không suy ra dữ liệu trần chỉ từ tên DEM.",
176:"Copernicus GLO-30 là dữ liệu bề mặt; FABDEM là sản phẩm phát triển từ dữ liệu Copernicus để giảm ảnh hưởng rừng và công trình [8], [37]. Đây là các sản phẩm khác nhau. Việc chọn nguồn cần đi cùng đánh giá sai số địa hình và chuỗi xử lý của đúng raster.",
178:"Ngay cả khi giảm ảnh hưởng công trình và thực vật, lưới 30 m có thể chưa thể hiện các đường dẫn hẹp. Độ phân giải ngang, sai số cao độ và cách hiệu chỉnh DEM đều có thể ảnh hưởng kết quả [15], [16].",
179:"Cống hoặc cầu có thể nhỏ hơn kích thước ô lưới. Nếu kết nối này không được biểu diễn, đường hoặc đê trong raster có thể hạn chế trao đổi nước. Ngược lại, hạ cả một ô rộng 30 m cũng có thể tạo tiết diện truyền nước quá lớn so với công trình thực; cần dữ liệu hình học để kiểm tra.",
180:"Địa hình là một giả thuyết giải thích sai khác, cần phân tích cùng forcing, điều kiện biên và quan sát. Hiệu chỉnh cao độ phải ghi rõ tiêu chí chọn ô, mức thay đổi và ảnh hưởng tới độ dốc, thể tích chứa cũng như kết nối dòng chảy [15].",
184:"HAND biểu diễn chênh cao tới kênh nhận nước theo đường dòng chảy [6]. Chỉ số này hỗ trợ mô tả địa hình tương đối; việc quy đổi thành nguy cơ hoặc độ sâu ngập cần thêm giả thiết thủy lực.",
185:"Trong đồ án, nhóm chỉ số ≤3,5 m được dùng để thống kê các ô có chênh cao nhỏ theo quy trình đang triển khai.",
186:"Nhóm 3,5–10 m dùng để thống kê phần chênh cao trung gian trong cùng quy trình.",
187:"Nhóm >10 m dùng để xem xét phần chênh cao lớn; giá trị này không tự chứng minh ô nằm trên sườn dốc hoặc không thể ngập.",
188:"Mã dự án gán cao độ kênh gần nhất theo khoảng cách Euclid sau khi tạo mạng thoát nước. Cách gán này khác HAND theo đường dòng chảy, nên chỉ số hiện có được dùng như chẩn đoán bổ trợ và được gọi rõ là chênh cao tới mạng thoát nước theo quy trình dự án.",
190:"BG_Flood là phần mềm mã nguồn mở phục vụ tính toán nước nông hai chiều, với tài liệu và mã nguồn do nhóm tác giả công bố [3]. Đồ án sử dụng một nhánh triển khai Python; thuật toán và tham số được mô tả theo mã, cấu hình và log thực thi của nhánh này.",
191:"Nhánh V44 dùng tổ chức ô lưới theo khối và bộ giải central-upwind Kurganov–Petrova [36]. Không đồng nhất mọi biến thể BG_Flood với một bộ giải duy nhất: thuật toán trong tài liệu upstream và nhánh thực chạy cần được phân biệt [3].",
192:"Kết quả được lưu dạng NetCDF với các trường độ sâu h, mực nước zs, địa hình zb và vận tốc u, v. Thời gian frame và tọa độ phải được đọc từ tệp, kiểm tra với lưới đầu vào trước khi tính diện tích hoặc chồng xếp SAR.",
211:"R(x,y,t) là nguồn cường độ mưa trước khi trừ tổn thất thấm, đơn vị m/s; không gọi là mưa hữu hiệu nếu I được trừ riêng.",
220:"2.1.3. Phương pháp thể tích hữu hạn và lược đồ central upwind",
222:"Phương pháp thể tích hữu hạn cập nhật biến trung bình ô từ thông lượng qua các mặt và số hạng nguồn [31]. Nhánh V44 dùng lược đồ central-upwind Kurganov–Petrova [36]. Dạng cập nhật tổng quát trên lưới vuông được viết như sau:",
224:"F và G là thông lượng số qua các mặt ô. Công trình Kurganov–Petrova trình bày lược đồ cân bằng tốt và giữ độ sâu không âm [36]; những tính chất lý thuyết này cần được kiểm tra cho triển khai cụ thể. Đồ án không mô tả nhánh V44 như bộ giải HLLC hoặc suy ra bảo toàn nghiêm ngặt từ tên phương pháp.",
226:"Bước thời gian của lược đồ hiện số được lựa chọn theo giới hạn lan truyền thông tin và điều kiện CFL [31]. Một biểu thức kiểm soát thường dùng là:",
229:"Log V44 ghi CFL=0,5. Bước thời gian thích nghi thay đổi theo trạng thái thủy lực và được đọc từ log; không coi một giá trị dt là đặc trưng cố định cho cả sự kiện. Cấu hình, mặc định của mã và giá trị thực thi cần được ghi riêng.",
233:"Mặc định param.eps của mã là 0,0001 m; nhánh GPU dùng ngưỡng 0,001 m khi triệt tiêu vận tốc ở vùng nước rất nông. Hai ngưỡng có vai trò khác nhau.",
234:"Khi tính vận tốc, ma sát và cập nhật trạng thái, cần tránh chia cho độ sâu gần 0. Đồng thời phải ghi nhận ảnh hưởng của các phép đặt lại hoặc chặn giá trị tới động lượng và cân bằng nước.",
235:"Trạng thái hồ tĩnh, độ sâu không âm và cân bằng thể tích là các kiểm tra cần thiết cho triển khai. Mã V44 hiện lưu có phép chặn độ sâu và vận tốc; việc hoàn tất mô phỏng không thay thế các kiểm tra này.",
237:"Sản phẩm SAR được sử dụng như lớp ngập tham chiếu suy đoán, không mặc nhiên là dữ liệu thực địa. Quy trình dưới đây phân biệt các bước cần có trong hồ sơ sản phẩm với những thao tác đã được chứng minh ở bộ dữ liệu dự án [9], [10].",
238:"Thu nhận dữ liệu: bộ đánh giá dùng raster mang mốc 05:36 ngày 30/11/2021. Cần bổ sung scene ID, thời gian UTC, chế độ, phân cực và loại sản phẩm từ metadata. Tên tệp không đủ để xác nhận thời điểm đỉnh lũ.",
239:"Quỹ đạo và hiệu chuẩn: cần ghi nguồn quỹ đạo cùng cấu hình hiệu chuẩn nếu đã áp dụng. Chưa có hồ sơ đủ để xác nhận POD, độ chính xác vị trí 5 cm hoặc một bộ tham số hiệu chuẩn cụ thể cho đúng raster sử dụng.",
240:"Đối với dữ liệu cường độ, hệ số tán xạ ngược dạng tuyến tính có thể biểu diễn ở dB theo:",
241:"σ⁰_dB = 10 log₁₀(σ⁰_linear)",
242:"Việc chuyển DN sang hệ số tán xạ ngược phải tuân theo sản phẩm và metadata hiệu chuẩn. Không tự bổ sung một công thức LUT hoặc chuỗi xử lý khi chưa có dữ liệu nguồn.",
243:"Hiệu chỉnh hình học và đưa về lưới: cần xác định raster gốc, hệ tọa độ, phương pháp hiệu chỉnh địa hình và phép tái lấy mẫu. Lưới mô hình 30 m không chứng minh SAR gốc có độ phân giải 30 m hoặc đã dùng FABDEM để nắn ảnh.",
244:"Giảm nhiễu đốm: bộ lọc và kích thước cửa sổ chỉ được mô tả là đã thực hiện khi có script hoặc log. Không đồng nhất một mô tả lý thuyết của Lee với bằng chứng đã áp dụng Refined Lee 5×5.",
245:"Phân loại và mặt nạ đánh giá",
246:"Tên raster đánh giá cho biết sản phẩm phân ngưỡng Otsu; thuật toán được mô tả ở mục 1.3.4 [5]. Giá trị ngưỡng và histogram của đúng sản phẩm cần được lưu để tái lập; chưa khẳng định −14,85 dB là giá trị đã kiểm chứng trong đợt đối soát.",
247:"Phép đối chiếu lịch sử loại NODATA, các ô có water density>0,5 và các ô chỉ số địa hình không hợp lệ. Cần ghi nguồn của lớp nước thường trực; nghiên cứu của Pekel và cộng sự là cơ sở về dữ liệu mặt nước toàn cầu [22], không phải bằng chứng rằng mọi raster cục bộ đều được tạo đúng theo công trình đó.",
250:"Xử lý vùng trũng DEM có thể thay đổi đường thoát nước và phân bố thể tích chứa. Trước khi áp dụng cần phân biệt vùng trũng thật với lỗi dữ liệu, ghi rõ mức điều chỉnh và kiểm tra ảnh hưởng tới độ dốc. Không quy một lần NaN cho hố trũng chỉ từ diễn giải địa hình.",
251:"Priority-Flood dùng hàng đợi ưu tiên để xử lý vùng trũng theo kết nối từ biên [14]. Phần dưới mô tả nguyên lý thuật toán; cấu hình xử lý thực tế phải được xác định từ script và raster đầu vào/đầu ra.",
252:"Bước 1: đưa các ô biên hợp lệ vào hàng đợi theo cao độ, với cách xử lý NODATA và cửa thoát được quy định. Không gán mọi biên lưu vực hoặc cửa thoát cao độ 0 m.",
253:"Bước 2: lấy ô thấp nhất, xét các ô lân cận chưa xử lý và cập nhật mức cao độ theo quy tắc lan truyền của thuật toán. Một phép nâng ô là thao tác của thuật toán, không tự chứng minh đó là hố giả ngoài thực địa.",
254:"Bước 3: lặp đến khi hoàn tất miền cần xử lý và ghi chênh cao, phạm vi thay đổi, thể tích cùng các tham số của lần xử lý.",
255:"DEM sau xử lý cần được đánh giá lại cho mục tiêu thủy lực. Lấp vùng trũng có thể cải thiện kết nối nhưng cũng loại bỏ dung tích chứa thật; không xem đây là bảo đảm loại bỏ mọi bất ổn định số học [15].",
256:"2.3.2. Hiệu chỉnh địa hình cục bộ để khảo sát kết nối dòng chảy",
257:"Đồ án khảo sát hạ cao độ tại những ô ứng viên dựa trên quy tắc raster. Các chỉnh sửa được xem là giả thuyết về kết nối, cần hồ sơ hình học và thử nghiệm đối chứng để liên hệ với công trình thực [15].",
258:"Quy trình chọn ô cần ghi rõ bản đồ địa hình, lớp ngập và phép tính độ nhô cao cục bộ được sử dụng, cùng mọi ngưỡng và bán kính lân cận.",
259:"Mã V44 dùng lớp SAR trong điều kiện chọn ô. Vì vậy, hiệu chỉnh và đánh giá với cùng lớp ngập có quan hệ phụ thuộc; chỉ số sau hiệu chỉnh chưa phải kiểm định độc lập.",
260:"Một thao tác hạ cao độ ô 30 m có thể mở đường truyền nước ở quy mô raster, nhưng không tự tương đương cống thật hoặc quá trình vỡ đê. Cần ghi cao độ trước/sau và tránh gọi ô ứng viên là vị trí cầu cống đã khảo sát.",
262:"San phẳng hình thái là một hướng thử nghiệm nhằm thay đổi kết nối dọc hành lang sông. Việc bỏ điều kiện chọn theo SAR có thể thay đổi phạm vi ô được tác động, nhưng cũng có nguy cơ mở rộng dòng chảy quá mức. Cần đặt giới hạn và đánh giá cả mức bỏ sót lẫn báo ngập ngoài tham chiếu.",
263:"Bước 1: xác định hành lang và vùng đệm từ mạng sông có provenance, ghi rõ kích thước và lưới.",
264:"Bước 2: xác định quy tắc chọn ô và cao độ mục tiêu theo nhánh thí nghiệm; khóa DEM đầu ra, số ô và phân bố mức thay đổi. Không dùng 3,5 m như cao độ bãi bồi đúng cho mọi vị trí.",
265:"Thử nghiệm BASE/CONDITIONED cần giữ cùng runtime, forcing, nhám, thấm, biên và quy tắc đánh giá. Khác biệt kết quả chỉ được diễn giải trong phạm vi những yếu tố đã kiểm soát.",
266:"2.4. Phương pháp đối chiếu và đánh giá kết quả",
268:"Phạm vi ngập mô phỏng được xác định bằng h>0,3 m. Phép chồng xếp lịch sử đưa lớp SAR về lưới đầu ra, xác định miền hợp lệ rồi tính ma trận 2×2. Cần kiểm tra transform và mặt nạ trước khi so các nhánh; tái lập số học không chứng nhận phương pháp căn chỉnh địa lý cũ.",
271:"TP: ô mô hình ngập và SAR ngập trong miền hợp lệ. Đây là trạng thái trùng khớp giữa hai lớp tham chiếu.",
272:"FP: ô mô hình ngập và SAR không ngập. Đây là bất đồng với SAR, chưa tự chứng minh ngập giả ngoài thực địa.",
273:"FN: ô mô hình không ngập và SAR ngập. Đây là phần bỏ sót so với SAR; nguyên nhân cần được phân tích theo vị trí và dữ liệu.",
274:"TN: ô mô hình và SAR đều không ngập trong miền hợp lệ. NODATA và nước thường trực bị loại không được tính như TN.",
276:"Các chỉ số được tính từ cùng ma trận và cùng miền đánh giá. Định nghĩa CSI, POD/Recall và FAR có thể đối chiếu với tài liệu NOAA/NWS [39].",
282:"Recall đo tỷ lệ ô thuộc lớp ngập SAR được mô hình bao phủ. Giá trị cao hơn cho thấy tăng bao phủ lớp tham chiếu, nhưng cần xem cùng Precision và sai số quan sát.",
283:"3. Độ chính xác dương tính Precision",
285:"Precision đo tỷ lệ ô ngập mô phỏng trùng lớp SAR. Đây không phải độ chính xác tổng thể hoặc xác suất đã được kiểm chứng ngoài thực địa.",
292:"MCC tổng hợp cả bốn nhóm trong ma trận nhị phân, có miền giá trị từ −1 đến +1. Cần báo cáo giá trị và điều kiện tính, không tự gán nhãn chất lượng khi chưa có tiêu chí phù hợp.",
293:"2.4.3. Chẩn đoán phân bố ngập theo chênh cao tới mạng thoát nước",
294:"Một lớp SAR tại một thời điểm chưa xác nhận diễn biến ngập của toàn chuỗi. Chẩn đoán địa hình được dùng để xem phần ngập tính toán phân bố thế nào theo chỉ số tương đối, đồng thời nêu giới hạn chia sẻ dữ liệu với mô hình.",
295:"Mã dự án tạo mạng kênh bằng tích lũy dòng chảy và các chỉ thị nước, sau đó dùng distance_transform_edt để gán ô kênh gần nhất theo khoảng cách Euclid. Đây là khác biệt so với HAND gốc truy theo đường dòng chảy [6].",
296:"P₃,₅ = N(ô ngập hợp lệ có chỉ số ≤3,5 m) / N(ô ngập hợp lệ) × 100%",
297:"P₃,₅ là thống kê phân bố trong thí nghiệm này. Không đặt P₃,₅≥85–90% như tiêu chuẩn đạt chung và không dùng nó để xác nhận độ sâu, vận tốc hoặc bảo toàn nước.",
299:"Quy trình nghiên cứu gồm chuẩn bị dữ liệu, cấu hình và thực thi mô hình, kiểm tra đầu ra, đối chiếu phạm vi ngập và phân tích giới hạn.",
300:"Dữ liệu: DEM đã xử lý trong chuỗi Copernicus, lớp phủ, raster nhám và tổn thất, mưa đồng nhất, bốn hydrograph kịch bản và lớp SAR tham chiếu.",
301:"Tiền xử lý: kiểm tra hệ tọa độ, đơn vị, nguồn dữ liệu và các thay đổi DEM; chỉ ghi thao tác SAR đã có hồ sơ chứng minh.",
302:"Mô phỏng: nạp cấu hình vào nhánh Python V44, tính nghiệm và ghi log các tham số cùng thời gian thực thi.",
303:"Đầu ra: lưu các trường thủy lực trong NetCDF theo bước xuất danh định 3 giờ; đọc thời gian frame thực tế khi hậu xử lý.",
304:"Đánh giá: khóa miền hợp lệ và ngưỡng h>0,3 m, chọn khung theo thời gian quan sát, tính chỉ số và chẩn đoán địa hình bổ trợ.",
306:"CHƯƠNG 3 KẾT QUẢ THỰC NGHIỆM VÀ THẢO LUẬN",
309:"Khu vực trọng tâm là hạ lưu sông Ba và đồng bằng Tuy Hòa trong địa giới Phú Yên tại thời kỳ sự kiện năm 2021. Miền raster rộng hơn vùng trọng tâm; cần giữ rõ phạm vi khi diễn giải diện tích và nguồn lưu lượng.",
310:"Sông Ba chảy ra cửa Đà Diễn. Sông Hinh nối vào hệ thống sông Ba, còn Bàn Thạch chảy tới cửa Đà Nông [40]. Những quan hệ này cần được thể hiện đúng khi đặt nguồn và mô tả mạng sông; không gán nguồn Bàn Thạch cho xả Sông Hinh chỉ từ tên tệp kịch bản.",
312:"Đồ án sử dụng sự kiện cuối tháng 11/2021 làm bối cảnh thiết lập mô phỏng. Nguồn số liệu khí tượng thủy văn được tham chiếu từ báo cáo Thiên tai Việt Nam 2021 [2].",
313:"Mưa lớn trong khu vực và dòng chảy sông là các yếu tố liên quan tới sự kiện. Bộ mưa đang nạp là một chuỗi kịch bản tổng hợp; chưa có hồ sơ đầy đủ để đồng nhất với các trị số mưa tại từng trạm hoặc khẳng định hệ số nhân 2 đã được hiệu chỉnh.",
314:"Báo cáo quốc gia ghi lưu lượng xả Ba Hạ 9.400 m³/s tại 15 giờ ngày 30/11 [2]. Tệp nguồn Ba Hạ trong cấu hình V44 có đỉnh 11.000 m³/s. Cần giữ riêng số liệu bối cảnh quan trắc và giá trị đã điều chỉnh của kịch bản.",
315:"Nguồn quốc gia ghi mực nước Phú Lâm 4,63 m, cao hơn báo động III 0,93 m [2]. Đồ án chưa đối chiếu chuỗi mực nước tại trạm với đầu ra V44 theo cùng datum; vì vậy chưa báo cáo sai số mực nước trạm như kết quả kiểm định.",
317:"DEM đầu vào V44 có 2.100×2.100 ô, bước ô 30 m. NetCDF lịch sử có 2.099×2.099 ô; cần phân biệt kích thước thiết kế với kích thước đầu ra cũ. Bảng 3.1 tổng hợp những thông số đã đối chiếu từ cấu hình và log.",
319:"Địa hình: PhuYen_DEM_30m_V44.asc, 2.100×2.100 ô, gốc mép X=270.000 m, Y=1.415.000 m, bước ô 30 m. Chuỗi provenance cục bộ dẫn tới Copernicus và các xử lý kế thừa; chưa đồng nhất với DTM đo mặt đất hoặc FABDEM chỉ từ tên phiên bản [37].",
320:"Nhám: cấu hình nạp PhuYen_manning_30m_V43.asc theo mô hình Manning. Mã lớp WorldCover và tên lớp phải được kiểm tra theo bảng phân loại chính thức [21]. Hệ số n trong raster là tham số mô hình, không phải phép đo do ESA cung cấp.",
321:"Bảng 3.2 Các mã lớp WorldCover liên quan tới mô tả bề mặt",
322:"Tổn thất: cấu hình nạp il_30m_V36.asc và cl_30m_V36.asc. IL là lượng mất ban đầu, CL là tốc độ mất liên tục theo quy ước engine. Cần xác nhận đơn vị và cách tạo raster trước khi diễn giải trị số của từng loại đất.",
323:"Mưa: rain_v34_era5_x2.txt gồm 144 bản ghi, hai cột thời gian–giá trị, được nạp ở chế độ uniform. Header ghi m/s và trung bình 16 điểm; script V34 nhân chuỗi tiền nhiệm với 2,0. Đây không phải 145 trạm mưa quan trắc độc lập. Mã cũ hiện lưu tiếp tục chia nguồn mưa cho 3.600.000; chưa có hash runtime tại lúc V44 chạy để xác nhận tuyệt đối tác động lên đầu ra lịch sử.",
324:"Lưu lượng: bốn vùng nguồn trong cấu hình được gắn nhãn Ba Hạ, Bàn Thạch, Kỳ Lộ và Tuy An. Đây là các hydrograph kịch bản, không tự là chuỗi quan trắc đã kiểm định.",
325:"Ba Hạ: xa_lu_BaHa_Nov2021_V40_sync.txt; vùng X=[270.030;270.630] m, Y=[1.441.000;1.443.000] m; đỉnh tệp 11.000 m³/s.",
326:"Bàn Thạch: xa_lu_BanThach_V43.txt; X=[276.000;277.000] m, Y=[1.424.500;1.425.500] m; đỉnh tệp 4.007,14 m³/s.",
327:"Kỳ Lộ: xa_lu_KyLo_V40_sync.txt; X=[289.000;290.000] m, Y=[1.476.500;1.477.500] m; đỉnh 5.000 m³/s. Tuy An: xa_lu_TuyAn_V40_sync.txt; X=[301.000;302.000] m, Y=[1.473.500;1.474.500] m; đỉnh 2.000 m³/s.",
328:"3.3. Các nhánh thực nghiệm và phạm vi so sánh",
329:"Các nhánh có thay đổi dữ liệu, runtime và quy tắc đánh giá. Bảng 3.3 chỉ tổng hợp những trạng thái có bằng chứng tại đợt đối soát ngày 05/10/2026, không xem chúng như chuỗi thí nghiệm chỉ khác một yếu tố.",
330:"Bảng 3.3 Trạng thái và phạm vi đánh giá của các nhánh thực nghiệm",
331:"Hai nhóm vấn đề cần tách khi diễn giải kết quả là tính nhất quán thực thi và ảnh hưởng của thay đổi dữ liệu.",
332:"Tính nhất quán thực thi",
333:"Phép đối soát xác nhận V44 hoàn thành 120 giờ và các chỉ số tại frame 33 tái lập đúng theo phương pháp cũ. Tuy nhiên, kích thước lưới xuất và mã xử lý đơn vị mưa hiện lưu có bất nhất. Hoàn tất chạy không chứng minh mọi đầu vào đã được thực thi đúng đơn vị.",
334:"Ảnh hưởng của địa hình",
335:"Việc nâng hoặc hạ cao độ có thể thay đổi kết nối và thể tích chứa. Muốn quy chênh lệch chỉ số riêng cho DEM, cần cùng forcing, runtime, điều kiện biên và phép đánh giá. Các phiên bản lịch sử chưa đáp ứng đồng thời các điều kiện này nên không dùng để khẳng định một nguyên nhân duy nhất.",
337:"V44 sử dụng DEM đã xử lý, bốn nguồn lưu lượng kịch bản và mưa đồng nhất. NetCDF có 40 khung đến 432.000 s. Phép đánh giá được tái tạo theo transform và mask lịch sử; những bất nhất của phương pháp cũ được giữ rõ khi diễn giải.",
338:"Bảng 3.4 trình bày 40 khung đối chiếu với cùng mặt nạ SAR. Các khung ngoài thời điểm quan sát là chẩn đoán độ khớp với một lớp cố định, không phải chuỗi kiểm định diễn biến lũ.",
339:"Bảng 3.4 Độ khớp các khung V44 với mặt nạ SAR cố định",
340:"Từ 3 tới 18 giờ, CSI tăng từ 2,85% tới 9,96%. Đây là xu hướng của chỉ số đối chiếu; chưa đủ để xác nhận pha truyền lũ hoặc trạng thái tràn bờ ngoài thực địa.",
341:"Trong khoảng 21–36 giờ, CSI tăng từ 11,82% tới 17,57%. Việc liên hệ mức tăng với một nguồn lưu lượng hoặc vị trí cống cần kiểm tra riêng chuỗi forcing và trường dòng chảy.",
342:"Tại khoảng 102 giờ, frame 33 cho CSI=20,64%, Recall=24,97%, Precision=54,32%. Ma trận và diện tích tương ứng được trình bày ở Bảng 3.5 và Hình 3.1.",
343:"Đến 120 giờ, CSI≈28,76%, Recall≈40,78%, Precision≈49,39% khi so với cùng lớp SAR. Không chọn khung cuối vì CSI cao hơn để thay thời điểm ảnh. Hình 3.2 thể hiện chuỗi chỉ số đã đối soát.",
344:"Tại frame 33, khoảng 75,03% lớp ngập SAR chưa được mô hình bao phủ và 45,68% lớp ngập mô hình nằm ngoài SAR. Kết quả cho thấy độ khớp hạn chế; Precision không duy trì trên 52,2% trong toàn chuỗi. Cần quan trắc độc lập để đánh giá trạng thái thực địa và phân tích nguyên nhân.",
345:"Bảng 3.5 Ma trận và chỉ số đối chiếu V44 tại frame 33",
347:"Hình 3.1 Phân bố các nhóm đối chiếu V44 với SAR tại khoảng 102 giờ",
349:"Hình 3.2 Chuỗi chỉ số V44 đối chiếu với cùng mặt nạ SAR",
350:"3.5. Chẩn đoán phân bố ngập theo địa hình",
351:"Đầu ra frame 33 được chồng xếp với chỉ số chênh cao tới mạng thoát nước theo quy trình Euclid của dự án. Quy trình này khác HAND theo đường dòng chảy và có liên hệ dữ liệu DEM với mô hình; kết quả được dùng như chẩn đoán bổ trợ.",
352:"Bảng 3.6 Phân bố ô ngập theo chênh cao tới mạng thoát nước",
353:"Trong 132.866 ô ngập hợp lệ, 120.153 ô có chỉ số ≤3,5 m, bằng 90,43%; 11.384 ô nằm trong nhóm 3,5–10 m và 1.329 ô trong nhóm >10 m. Hình 3.3 biểu diễn tỷ lệ các nhóm. Không suy ra độ sâu, độ dốc, cân bằng nước hoặc độ đúng thủy lực từ các tỷ lệ này.",
354:"3.6. Giới hạn của kết quả và thiết kế thực nghiệm tiếp theo",
355:"Hình 3.1 cho thấy phần mô hình không ngập/SAR ngập lớn trong miền đánh giá. Cần xem xét đồng thời địa hình, mưa và lưu lượng, nhám/thấm, biên, căn chỉnh raster và sai số SAR; các chỉ số tổng hợp không phân rã được các nguyên nhân.",
357:"Hình 3.3 Tỷ lệ ô ngập V44 theo các nhóm chỉ số địa hình",
358:"Một giả thuyết cần kiểm tra là kết nối qua các dải cao độ hẹp chưa được biểu diễn đủ trên lưới 30 m. Việc hạ cao độ có thể tăng truyền nước, nhưng cũng có thể làm tăng phần ngập ngoài lớp tham chiếu.",
359:"Cần đối chiếu mặt cắt và cao trình với dữ liệu công trình, đồng thời kiểm tra datum. Không xem một mặt cắt DEM và mực nước tính toán là phép đo thực địa hoặc bằng chứng đủ để xác định thân đê thật.",
360:"Trong mã chọn ô V44 có điều kiện phụ thuộc SAR:",
361:"barrier_mask = (sar_flood == 1) & plain_mask & (local_relief >= relief_threshold)",
362:"Điều kiện sar_flood==1 loại các ô được SAR phân loại cạn khỏi tập ứng viên. Nó có thể bỏ qua một dải cao độ khô, nhưng chưa chứng minh dải ấy là đê thật hay nguyên nhân chính của mọi FN. Đồng thời, dùng lại lớp SAR để chấm điểm làm hạn chế tính độc lập của đánh giá.",
363:"Thiết kế đánh giá V45 và V48",
364:"Đối soát V45 ngày 04/10 ghi CSI≈32,40%, Recall≈72,95%, Precision≈36,83% tại 102 giờ theo engine và phép đánh giá của nhánh đó. Mức bao phủ tăng đi kèm phần báo ngập ngoài SAR; không so trực tiếp như hiệu quả riêng của DEM với V44 cũ.",
365:"Snapshot V45 ngày 05/10 cho thấy đầu ra exploratory tới 105 giờ. Đây là trạng thái có thời điểm, chưa phải bằng chứng hoàn thành 120 giờ. Số ô sửa địa hình phải được khóa theo đúng DEM/manifest thay vì coi một con số của prototype là kết quả công trình thực.",
366:"V48 R1 đã sửa xử lý đơn vị mưa và lưới. Cặp BASE/CONDITIONED giữ các đầu vào khác như nhau và sử dụng DEM đã khóa; nhánh CONDITIONED khác BASE ở 66 ô hạ cao độ, mức hạ lớn nhất khoảng 0,996 m. Các chi tiết thuộc hồ sơ cấu hình, chưa chứng minh cải thiện phạm vi ngập.",
367:"Snapshot 12:49 ngày 05/10/2026 UTC+7 ghi BASE đang chạy tới khoảng 26,43/120 giờ, có 8 frame đến khoảng 24 giờ; chưa có output CONDITIONED. Chỉ bổ sung kết luận về tác động DEM sau khi hai nhánh hoàn tất và được đánh giá theo cùng quy tắc đã khóa trước.",
368:"3.7. Hướng nghiên cứu mô hình học máy thay thế",
369:"Mô hình học máy có thể được khảo sát như công cụ xấp xỉ đầu ra thủy lực nhằm giảm thời gian tính khi xét nhiều kịch bản. Hồ sơ hiện có chưa đủ để xác nhận Bảng 3.7 cũ thuộc đúng thí nghiệm V44, nên phần này trình bày thiết kế nghiên cứu tiếp theo.",
370:"Các thuật toán có thể khảo sát gồm XGBoost [33], ANN và LSTM [34]. Phải xác định dữ liệu, biến đầu vào và thời điểm chúng có sẵn; nếu dùng độ sâu ở thời điểm trước, cần mô tả nguồn và điều kiện vận hành của biến đó.",
371:"Bảng 3.7 Nội dung cần kiểm tra khi xây dựng mô hình học máy thay thế",
372:"Các giới hạn cần được đưa vào thiết kế đánh giá",
373:"Thời gian suy luận phải được đo với số mẫu, phần cứng, cách nạp mô hình và phạm vi tiền/hậu xử lý rõ ràng. Chưa giữ các giá trị 2,8/4,1/7,6 s hoặc tỷ số tăng tốc khi không có hồ sơ đo tương ứng.",
374:"Chia ngẫu nhiên các pixel có thể tạo quan hệ không gian gần giữa tập học và tập kiểm tra. Nếu mục tiêu là dự báo một sự kiện hoặc vùng khác, cần chia theo sự kiện hoặc khối không gian và kiểm tra rò rỉ thông tin theo thời gian.",
375:"Sai số đối với nhãn do solver tạo ra đo khả năng xấp xỉ solver; độ chính xác thực địa cần đánh giá riêng. Học máy cũng cần kiểm tra giá trị âm, tính liên tục và cân bằng nước, cùng khả năng ngoại suy sang kịch bản chưa có trong dữ liệu học.",
376:"Chỉ báo cáo kết quả học máy khi có chuỗi tái lập gồm dữ liệu, mã, split, mô hình và phép đo. Trong phạm vi bản đồ án này, học máy được đặt là hướng phát triển, chưa coi là công cụ dự báo đã được xác nhận.",
391:"Các kiến nghị tập trung vào dữ liệu, kiểm tra số học và thiết kế đánh giá độc lập trước khi chuyển sang ứng dụng vận hành.",
392:"Dữ liệu địa hình: bổ sung khảo sát cao độ, mặt cắt và công trình; kiểm tra datum, phạm vi sửa DEM và ảnh hưởng tới dung tích chứa. Dữ liệu LiDAR/UAV có thể được khảo sát khi phù hợp nguồn lực.",
393:"Dữ liệu thủy văn: hoàn thiện chuỗi mưa, lưu lượng, mực nước và biên biển; phân biệt quan trắc với kịch bản, ghi mọi nội suy và hệ số điều chỉnh.",
394:"Kiểm tra và học máy: đánh giá cân bằng nước, hồ tĩnh và các bài toán chuẩn; xây dựng tập dữ liệu nhiều sự kiện và split phù hợp mục tiêu khái quát hóa trước khi huấn luyện surrogate.",
395:"Ứng dụng: chỉ đánh giá khả năng cảnh báo sau khi có quan trắc độc lập, đầu vào dự báo và thời gian toàn quy trình. Kết quả hiện tại phục vụ nghiên cứu phương pháp và kế hoạch cải thiện.",
}
for index,text in replacements.items():
    setp(index,text)
setp(107, "Kết luận và kiến nghị tổng hợp kết quả, giới hạn và hướng phát triển. Tài liệu tham khảo cùng phụ lục cung cấp hồ sơ để truy vết.")
drop(107)
drop(108)
setp(163, "Giả sử ảnh có N pixel với mức xám từ 1 đến L. Số pixel ở mức xám i là nᵢ và xác suất tương ứng là pᵢ=nᵢ/N. Ngưỡng T chia ảnh thành lớp C₀ với mức xám từ 1 đến T và lớp C₁ với mức xám từ T+1 đến L. Xác suất của hai lớp được tính như sau:")
setp(167, "Mức xám trung bình của ảnh là μ=ω₀μ₀+ω₁μ₁. Phương sai giữa hai lớp được định nghĩa như sau:")
setp(174, "DSM (Digital Surface Model — mô hình bề mặt số) biểu diễn cao độ bề mặt được thu nhận, có thể bao gồm ảnh hưởng của công trình và tán thực vật. Mức ảnh hưởng phụ thuộc cảm biến, phương pháp tạo sản phẩm và xử lý dữ liệu.")
setp(219, "Hệ số nhám Manning n có thứ nguyên thời gian nhân độ dài lũy thừa −1/3. Giá trị n phụ thuộc cách mô tả lòng dẫn và lớp phủ, nên cần được xác định bằng nguồn tham số hoặc hiệu chỉnh phù hợp. Đồ án dùng trường n phân bố theo không gian; nguồn raster, quy tắc gán từ WorldCover và đơn vị phải được lưu trong hồ sơ thiết lập.")
setp(224, "Trong đó ΔFᵢⱼ=F(i+½,j)−F(i−½,j) và ΔGᵢⱼ=G(i,j+½)−G(i,j−½) là chênh lệch thông lượng số qua hai mặt đối diện của ô. Công trình Kurganov–Petrova trình bày lược đồ cân bằng tốt và giữ độ sâu không âm [36]; các tính chất này cần được kiểm tra cho triển khai cụ thể. Nhánh V44 sử dụng lược đồ central-upwind; kết quả bảo toàn được đánh giá bằng kiểm tra cân bằng nước.")
setp(228, "Trong đó c=√(gh) là vận tốc sóng trọng trường trong nước nông; |u|+c và |v|+c xác định vận tốc truyền thông tin theo hai phương. Hệ số Courant (CFL) điều khiển bước thời gian và phải phù hợp với lược đồ số. Log V44 ghi CFL=0,5; riêng giá trị này không chứng minh tính ổn định của toàn bộ lần chạy.")
setp(231, "Ranh giới giữa vùng khô và vùng ngập thay đổi theo thời gian. Khi h tiến về 0, phép tính u=(hu)/h và các số hạng ma sát có thể trở nên nhạy hoặc không xác định do mẫu số nhỏ. Bộ giải cần ngưỡng độ sâu, quy tắc cập nhật vận tốc và xử lý khô–ướt phù hợp để hạn chế giá trị không hữu hạn và duy trì độ sâu không âm.")
setp(279, "CSI nhận giá trị từ 0% đến 100% và đánh giá sự trùng khớp giữa phạm vi mô hình với lớp tham chiếu. Chỉ số tính cả báo động sai FP và bỏ sót FN, không sử dụng số lượng lớn các ô không ngập TN. Trong đồ án, lớp tham chiếu là sản phẩm SAR đã phân loại; độ đúng của SAR và phép căn chỉnh vẫn ảnh hưởng đến kết quả.")
setp(280, "2. Xác suất phát hiện so với lớp tham chiếu (Probability of Detection — POD / Recall):")
setp(342, "Tại khoảng 102 giờ, khung 33 cho CSI 20,64%, Recall 24,97% và Precision 54,32%. Ma trận và diện tích tương ứng được trình bày ở Bảng 3.5 và Hình 3.1.")
setp(369, "Mô hình học máy có thể được khảo sát như công cụ xấp xỉ đầu ra thủy lực để giảm thời gian tính khi xét nhiều kịch bản. Hồ sơ hiện có chưa đủ để xác nhận kết quả học máy của thí nghiệm V44; phần này đề xuất thiết kế nghiên cứu tiếp theo.")
setp(370, "Các thuật toán có thể khảo sát gồm XGBoost [33], mạng nơ-ron nhân tạo (ANN) và mạng bộ nhớ dài–ngắn hạn (LSTM) [34]. Thiết kế phải xác định biến đầu vào và thời điểm chúng có sẵn; nếu dùng độ sâu ở thời điểm trước, cần mô tả nguồn và điều kiện vận hành của biến đó.")
setp(373, "Thời gian suy luận cần được đo với số mẫu, phần cứng, cách nạp mô hình và phạm vi tiền/hậu xử lý rõ ràng. Các phép so sánh tốc độ phải sử dụng cùng phạm vi tác vụ và dữ liệu; hồ sơ hiện tại chưa đủ để công bố thời gian suy luận hoặc hệ số tăng tốc.")
drop(79)
conclusion = [clean_md(p) for p in sections[7].split("\n\n")]
replace_range(380,389,conclusion)

# Keep only actual figures and let Word populate lists from the caption fields.
for i in range(36,51): drop(i)
for i in range(53,61): drop(i)

def field(p, instruction, cached=""):
    run = p.add_run()
    start = OxmlElement("w:fldChar"); start.set(qn("w:fldCharType"),"begin")
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"),"preserve"); code.text = " "+instruction+" "
    separator = OxmlElement("w:fldChar"); separator.set(qn("w:fldCharType"),"separate")
    result = OxmlElement("w:t"); result.text = cached
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    for element in [start,code,separator,result,end]: run._r.append(element)

toc_title = pp[35].insert_paragraph_before("MỤC LỤC",style="Heading 1")
toc_title.paragraph_format.page_break_before = True
toc_p = pp[35].insert_paragraph_before("")
field(toc_p,'TOC \\o "1-3" \\h \\z \\u')
fig_p = pp[51].insert_paragraph_before("")
field(fig_p,'TOC \\h \\z \\c "Hinh"')
tab_p = pp[61].insert_paragraph_before("")
field(tab_p,'TOC \\h \\z \\c "Bang"')

def new_rows(table, rows):
    for row in list(table.rows): table._tbl.remove(row._tr)
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for _ in rows[0]:
        column=OxmlElement('w:gridCol');column.set(qn('w:w'),str(int(16/len(rows[0])*567)));grid.append(column)
    for values in rows:
        assert len(values)==len(rows[0])
        row=table.add_row()
        for cell,text in zip(row.cells,values): cell.text=str(text)

new_rows(tt[1],[["Mô hình", "SAR ngập", "SAR không ngập"],
               ["Ngập h>0,3 m", "TP", "FP"],
               ["Không ngập h≤0,3 m", "FN", "TN"]])
new_rows(tt[2],[["Thông số", "Giá trị", "Ghi chú"],
 ["Raster đầu vào", "2.100×2.100; 30 m", "4.410.000 ô; EPSG:32649"],
 ["Đầu ra lịch sử", "2.099×2.099", "Bất nhất lưới của runtime cũ"],
 ["Gốc mép raster", "X=270.000; Y=1.415.000 m", "Phân biệt mép với tâm ô"],
 ["Bộ giải", "Kurganov–Petrova", "Central-upwind; nhánh Python V44"],
 ["CFL", "0,5", "Giá trị trong log"],
 ["Ngưỡng số", "eps=0,0001 m; GPU 0,001 m", "Phân biệt vai trò các ngưỡng"],
 ["Thời gian", "432.000 s", "120 giờ mô hình"],
 ["Bước xuất", "10.800 s", "Đọc thời gian frame thực tế"],
 ["Mưa", "Uniform; 144 bản ghi", "Header m/s; giới hạn mã cũ"],
 ["Nguồn lưu lượng", "4 vùng nguồn", "Ba Hạ; Bàn Thạch; Kỳ Lộ; Tuy An"],
 ["zsinit", "1,0 m", "Kiểm tra quy ước cao độ"],
 ["Mã biên", "left=0; right/top/bottom=2", "Chưa xác nhận có chuỗi triều quan trắc"]])
new_rows(tt[3],[["Mã lớp", "Tên lớp chính thức", "Vai trò khi thiết lập nhám"],
 ["10","Tree cover — cây gỗ","Kiểm tra bảng gán n và nguồn lớp phủ"],
 ["20","Shrubland — cây bụi","Phân biệt với grassland"],
 ["30","Grassland — đồng cỏ","Không gán tên cropland"],
 ["40","Cropland — cây trồng","Tham số n phải truy theo raster"],
 ["50","Built-up — xây dựng","Không suy ra n từ mã lớp đơn thuần"],
 ["80","Permanent water bodies — nước thường trực","Phân biệt với ngập sự kiện"]])
new_rows(tt[4],[["Nhánh","Lưới danh định","Đầu ra đã đọc","CSI tại 102 h","Runtime/đánh giá","Giới hạn so sánh"],
 ["V44","30 m; đầu ra 2.099²","40 frame; 120 h","20,64%","Nhánh cũ; mask lịch sử","SAR tham gia hiệu chỉnh; lưới/đơn vị"],
 ["V45","30 m; đầu ra 2.100²","35 frame; 105 h","32,40%","Engine R1; nhánh thăm dò 04/10","Khác runtime và phép đánh giá V44"],
 ["V48 R1","30 m; 2.100²","BASE đang chạy ở snapshot 05/10","Chưa có cặp kết quả","Runtime sửa mưa/lưới","Chưa nghiệm thu BASE/CONDITIONED"]])
new_rows(tt[5],[["Frame","Thời gian thực s","Giờ mô hình","CSI %","Recall %","Precision %"]]+[
 [str(r["frame"]),f'{r["time_s"]:.3f}'.replace('.',','),f'{r["time_s"]/3600:.6f}'.replace('.',',')]+[
 f'{v:.2f}'.replace('.',',') for v in r["computed_CSI_Recall_Precision_pct"]]
 for r in EVIDENCE["tables_all_40_frames"]])
tp,fp,fn,tn=72168,60698,216808,3130770
precision=tp/(tp+fp); recall=tp/(tp+fn)
total=tp+fp+fn+tn
pe=((tp+fp)*(tp+fn)+(tn+fn)*(tn+fp))/total**2
kappa=((tp+tn)/total-pe)/(1-pe)
mcc=(tp*tn-fp*fn)/np.sqrt(float((tp+fp)*(tp+fn)*(tn+fp)*(tn+fn)))
new_rows(tt[6],[["Đại lượng","Giá trị","Phạm vi diễn giải"],
 ["TP","72.168","Mô hình ngập; SAR ngập"],["FP","60.698","Mô hình ngập; SAR không ngập"],
 ["FN","216.808","Mô hình không ngập; SAR ngập"],["TN","3.130.770","Cả hai không ngập; miền hợp lệ"],
 ["CSI","20,64%","TP/(TP+FP+FN)"],["Recall","24,97%","TP/(TP+FN)"],
 ["Precision","54,32%","TP/(TP+FP)"],["FAR","45,68%","FP/(TP+FP)"],
 ["F1","34,22%","Trung bình điều hòa Recall/Precision"],
 ["Kappa",f"{kappa:.6f}".replace('.',','),"Tính từ cùng ma trận"],
 ["MCC",f"{mcc:.6f}".replace('.',','),"Tính từ cùng ma trận"],
 ["Ngập mô hình","119,5794 km²","Trong miền đánh giá"],
 ["Ngập SAR","260,0784 km²","Trong miền đánh giá"],
 ["Ô hợp lệ","3.480.444","Không tính ô loại khỏi mask"]])
hc=EVIDENCE["hand_frame33"]
new_rows(tt[7],[["Nhóm chỉ số","Số ô ngập","Tỷ lệ %","Diện tích km²"],
 ["≤3,5 m","120.153","90,43","108,1377"],
 [">3,5 tới 10 m","11.384","8,57","10,2456"],
 [">10 m","1.329","1,00","1,1961"],
 ["Tổng","132.866","100,00","119,5794"]])
new_rows(tt[8],[["Nội dung","Cần xác định","Cách đánh giá","Giới hạn"],
 ["Nhãn","Kết quả solver hoặc quan trắc","Nêu riêng mục tiêu xấp xỉ và thực địa","Không coi nhãn solver là chân lý"],
 ["Đầu vào","Biến có sẵn tại lúc dự báo","Kiểm tra độ sâu tiền nhiệm và thời gian","Tránh rò rỉ thông tin"],
 ["Chia tập","Sự kiện/khối không gian","Giữ tập kiểm tra độc lập","Random pixel chưa đủ"],
 ["Sai số","MAE, RMSE và phạm vi đo","Báo cáo theo tập và đơn vị","Không thay số từ thí nghiệm khác"],
 ["Thời gian","Máy, số mẫu, tiền/hậu xử lý","Đo lặp và mô tả phép đo","Chưa có latency được xác nhận"]])
tt[0].cell(11,2).text="Precision — tỷ lệ ô mô hình ngập trùng SAR (TP/(TP+FP))"
tt[0].cell(2,2).text="BG_Flood — tên phần mềm mô phỏng ngập hai chiều"
tt[0].cell(12,2).text="Height Above the Nearest Drainage — chênh cao tới kênh nhận nước theo đường dòng"
tt[0].cell(15,2).text="HLL — Harten–Lax–van Leer; HLLC khôi phục sóng tiếp xúc"

# Figures generated from audited data, rather than retaining obsolete claims.
def font(size,bold=False):
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts")/("arialbd.ttf" if bold else "arial.ttf")),size)
def txt(draw,xy,text,size=28,fill="#111111",bold=False): draw.text(xy,text,font=font(size,bold),fill=fill)
palette=np.array([[255,255,255],[236,238,240],[25,140,93],[204,76,52],[43,113,171],[132,156,178]],dtype=np.uint8)
classes=np.load(WORK/"v44_comparison_classes.npy")
map_image=Image.fromarray(palette[classes]).resize((1700,1700),Image.Resampling.NEAREST)
im=Image.new("RGB",(1900,2180),"white"); im.paste(map_image,(100,110)); dr=ImageDraw.Draw(im)
txt(dr,(100,35),"Đối chiếu V44 và SAR tại frame 33",42,bold=True)
dr.rectangle((100,110,1800,1810),outline="#929AA1",width=2)
txt(dr,(1680,150),"N",40,bold=True); dr.line((1700,240,1700,195),fill="#111111",width=6); dr.polygon([(1700,178),(1687,204),(1713,204)],fill="#111111")
labels=["TP  72.168 ô   Cả hai ngập","FP  60.698 ô   Chỉ mô hình ngập","FN  216.808 ô   Chỉ SAR ngập","TN  3.130.770 ô   Cả hai không ngập","Nước thường trực loại khỏi đánh giá"]
for j,(color,label) in enumerate(zip(palette[[2,3,4,1,5]],labels)):
    yy=1860+j*52; dr.rectangle((100,yy,137,yy+32),fill=tuple(map(int,color)),outline="#B3B9BD"); txt(dr,(156,yy-2),label,31)
txt(dr,(100,2120),"Miền và phép chồng xếp lịch sử   h > 0,3 m",27,fill="#444444")
im.save(WORK/"figure_3_1.png")
im=Image.new("RGB",(1800,1000),"white"); dr=ImageDraw.Draw(im)
txt(dr,(120,35),"Độ khớp với cùng mặt nạ SAR",40,bold=True)
left,top,right,bottom=160,170,1690,800
for value in range(0,71,10):
    yy=bottom-value/70*(bottom-top); dr.line((left,yy,right,yy),fill="#DFE3E6",width=2); txt(dr,(88,yy-16),str(value),28)
dr.line((left,top,left,bottom,right,bottom),fill="#333333",width=3)
for hour in [3,24,48,72,96,120]:
    xx=left+(hour-3)/117*(right-left); dr.line((xx,bottom,xx,bottom+10),fill="#333333",width=2); txt(dr,(xx-20,bottom+19),str(hour),28)
colors=["#173F68","#168278","#B57721"]
for k,label in enumerate(["CSI","Recall","Precision"]):
    xy=[(left+(r["time_s"]/3600-3)/117*(right-left),bottom-r["computed_CSI_Recall_Precision_pct"][k]/70*(bottom-top)) for r in EVIDENCE["tables_all_40_frames"]]
    dr.line(xy,fill=colors[k],width=5)
    dr.line((200+400*k,115,250+400*k,115),fill=colors[k],width=5); txt(dr,(265+400*k,94),label,29)
xx=left+99/117*(right-left)
for yy in range(top,bottom,22): dr.line((xx,yy,xx,min(yy+12,bottom)),fill="#626B72",width=2)
txt(dr,(xx-60,top-40),"102 giờ",27)
txt(dr,(120,890),"Chỉ số (%)",29); txt(dr,(1320,850),"Giờ mô hình",29)
txt(dr,(120,944),"Các khung ngoài mốc ảnh là chẩn đoán với một lớp tham chiếu cố định",27,fill="#444444")
im.save(WORK/"figure_3_2.png")
im=Image.new("RGB",(1800,900),"white"); dr=ImageDraw.Draw(im)
txt(dr,(110,40),"Phân bố ô ngập theo chỉ số địa hình",40,bold=True)
bar_left,bar_right=540,1490
items=[("≤ 3,5 m",120153,"#173F68"),("> 3,5 tới 10 m",11384,"#168278"),("> 10 m",1329,"#B57721")]
for j,(label,count,color) in enumerate(items):
    yy=190+j*150; pct=count/132866*100
    txt(dr,(110,yy+20),label,33); dr.rectangle((bar_left,yy,bar_right,yy+78),fill="#F0F2F3")
    dr.rectangle((bar_left,yy,bar_left+pct/100*(bar_right-bar_left),yy+78),fill=color)
    txt(dr,(1515,yy+20),f"{pct:.2f}%".replace('.',','),32)
    txt(dr,(bar_left,yy+90),f"{count:,} ô".replace(',','.'),28,fill="#444444")
txt(dr,(110,720),"132.866 ô ngập hợp lệ tại frame 33",32,bold=True)
txt(dr,(110,790),"Chỉ số Euclid của dự án   Chẩn đoán bổ trợ về địa hình",30,fill="#444444")
im.save(WORK/"figure_3_3.png")
for i,image_file,width in [(346,"figure_3_1.png",14.5),(348,"figure_3_2.png",15.5),(356,"figure_3_3.png",15.5)]:
    pp[i].clear(); pp[i].add_run().add_picture(str(WORK/image_file),width=Cm(width))
    pp[i].alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp[i].paragraph_format.keep_with_next=True
    pp[i].paragraph_format.first_line_indent=Cm(0)
    for shape in pp[i]._p.xpath('.//wp:docPr'): shape.set("descr", {346:"Bản đồ các lớp đối chiếu mô hình và SAR",348:"Chuỗi CSI Recall Precision",356:"Tỷ lệ ô ngập theo chênh cao tới mạng kênh"}[i])

# Citation identities are corrected first, then numbered by first appearance.
refs = {item["number"]:re.sub(r"^\[\d+\]\s*","",item["thesis_entry"]) for item in json.loads((ROOT/"Do_an_research_20261005/bibliography_registry_audit.json").read_text(encoding="utf-8"))}
refs.update({
1:"World Meteorological Organization (2021). WMO Atlas of Mortality and Economic Losses from Weather, Climate and Water Extremes (1970–2019). WMO-No. 1267. https://library.wmo.int/records/item/57608-wmo-atlas-of-mortality-and-economic-losses-from-weather-climate-and-water-extremes-1970-2019.",
2:"Ban Chỉ đạo Quốc gia về Phòng chống thiên tai. Thiên tai Việt Nam 2021, phần mưa lũ cuối tháng 11, tr. 150, 156. https://chinhsach.dmptc.gov.vn/Uploads/Thu%20vien%20tai%20lieu/Tai%20lieu%20truyen%20thong/Thien%20tai%20Viet%20Nam%202021%20fix%203.pdf.",
3:"Bosserelle, C. và cộng sự. BG_Flood source code and documentation, commit d835b94d39d2510c540414207106eea3ce9dd8dc. https://github.com/CyprienBosserelle/BG_Flood/tree/d835b94d39d2510c540414207106eea3ce9dd8dc. Truy cập 05/10/2026.",
6:"Nobre, A. D., Cuartas, L. A., Hodnett, M., Rennó, C. D., Rodrigues, G., Silveira, A., Waterloo, M. và Saleska, S. (2011). Height Above the Nearest Drainage — a hydrologically relevant new terrain model. Journal of Hydrology, 404(1–2), 13–29. https://doi.org/10.1016/j.jhydrol.2011.03.051.",
8:"Hawker, L., Uhe, P., Paulo, L., Sosa, J., Savage, J., Sampson, C. và Neal, J. (2022). A 30 m global map of elevation with forests and buildings removed. Environmental Research Letters, 17, 024016. https://doi.org/10.1088/1748-9326/ac4d4f.",
9:"Torres, R. và cộng sự (2012). GMES Sentinel-1 mission. Remote Sensing of Environment, 120, 9–24. https://doi.org/10.1016/j.rse.2011.05.028.",
10:"European Commission Joint Research Centre. Global flood monitoring. Report JRC119812, mục 5.1.1–5.1.5. https://publications.jrc.ec.europa.eu/repository/bitstream/JRC119812/global_flood_monitoring.pdf. Truy cập 05/10/2026.",
11:"Hersbach, H. và cộng sự (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146(730), 1999–2049. https://doi.org/10.1002/qj.3803.",
14:"Barnes, R., Lehman, C. và Mulla, D. J. (2014). Priority-flood: An optimal depression-filling and watershed-labeling algorithm for digital elevation models. Computers & Geosciences, 62, 117–127. https://doi.org/10.1016/j.cageo.2013.04.024.",
15:"Callow, J. N., Van Niel, K. P. và Boggs, G. S. (2007). How does modifying a DEM to reflect known hydrology affect subsequent terrain analysis? Journal of Hydrology, 332(1–2), 30–39. https://doi.org/10.1016/j.jhydrol.2006.06.020.",
16:"Saksena, S. và Merwade, V. (2015). Incorporating the effect of DEM resolution and accuracy for improved flood inundation mapping. Journal of Hydrology, 530, 180–194. https://doi.org/10.1016/j.jhydrol.2015.09.069.",
21:"European Space Agency. ESA WorldCover Data Access and Product Documentation. https://esa-worldcover.org/en/data-access. Truy cập 05/10/2026. Phân biệt sản phẩm 2020 v100 và 2021 v200.",
24:"Vacondio, R., Dal Palù, A. và Mignosa, P. (2014). GPU-enhanced Finite Volume Shallow Water solver for fast flood simulations. Environmental Modelling & Software, 57, 60–75. https://doi.org/10.1016/j.envsoft.2014.02.003.",
36:"Kurganov, A. và Petrova, G. (2007). A second-order well-balanced positivity preserving central-upwind scheme for the Saint-Venant system. Communications in Mathematical Sciences, 5(1), 133–160. https://doi.org/10.4310/CMS.2007.v5.n1.a6.",
37:"Copernicus Data Space Ecosystem. Copernicus DEM Product Collection Description. https://dataspace.copernicus.eu/explore-data/data-collections/copernicus-contributing-missions/collections-description/COP-DEM. Truy cập 05/10/2026.",
39:"Welles, E. (2004). Hydrology and Verification. NOAA National Weather Service, tr. 24. https://www.weather.gov/media/owp/oh/rfcdev/docs/dohVerif_edwin.pdf.",
40:"Lương Hữu Dũng, Lê Quốc Huy, Phan Văn Thành, Văn Thị Hằng, Trần Thanh Thủy và Nguyễn Văn Đại (2024). Ứng dụng mô hình HEC-RAS mô phỏng ngập lụt cho hạ lưu sông Ba. Tạp chí Khoa học Biến đổi khí hậu, 30, 65–76. https://imh.ac.vn/wp-content/uploads/2024/07/30.bai-6.pdf.",
})
aliases={12:31,17:40,18:40,19:40,20:2,26:16}
ordered=[]
def renumber(text):
    def sub(m):
        key=aliases.get(int(m.group(1)),int(m.group(1)))
        assert key in refs
        if key not in ordered: ordered.append(key)
        return f"[{ordered.index(key)+1}]"
    return re.sub(r"\[(\d+)\]",sub,text)
for p in doc.paragraphs:
    if p._p in [x._p for x in pp[398:]]: continue
    if re.search(r"\[\d+\]",p.text): p.text=renumber(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if re.search(r"\[\d+\]",cell.text): cell.text=renumber(cell.text)
for i in range(398,433): drop(i)
for j,key in enumerate(ordered,1):
    p=doc.add_paragraph(f"[{j}] {refs[key]}",style="Normal")
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.left_indent=Cm(.7)
    p.paragraph_format.first_line_indent=Cm(-.7)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.keep_together=True

# Native Word equations preserve editing and avoid literal programming notation.
equations={164:"ω₀(T) = Σᵢ₌₁ᵀ pᵢ ;   ω₁(T) = 1 − ω₀(T)",
166:"μ₀(T) = [Σ(i pᵢ), i=1…T]/ω₀(T) ;   μ₁(T) = [Σ(i pᵢ), i=T+1…L]/ω₁(T)",
168:"σ_B²(T) = ω₀(T) ω₁(T) [μ₀(T) − μ₁(T)]²",169:"T* = arg max σ_B²(T)",
183:"HAND(x) = Z(x) − Z(drainage(x))",199:"∂U/∂t + ∂F(U)/∂x + ∂G(U)/∂y = S(U)",
201:"U = [h, hu, hv]ᵀ",202:"F(U) = [hu, hu² + gh²/2, huv]ᵀ",203:"G(U) = [hv, huv, hv² + gh²/2]ᵀ",
204:"S(U) = [R − I, −gh ∂zᵦ/∂x − τᵦₓ/ρ, −gh ∂zᵦ/∂y − τᵦᵧ/ρ]ᵀ",
217:"τᵦₓ/ρ = g n² u √(u² + v²) / h¹ᐟ³",218:"τᵦᵧ/ρ = g n² v √(u² + v²) / h¹ᐟ³",
223:"Uᵢⱼⁿ⁺¹ = Uᵢⱼⁿ − (Δt/Δx) ΔFᵢⱼ − (Δt/Δy) ΔGᵢⱼ + Δt Sᵢⱼ",
227:"Δt ≤ CFL min[Δx/(|u| + √gh), Δy/(|v| + √gh)]",
241:"σ⁰(dB) = 10 log₁₀[σ⁰(linear)]",278:"CSI = TP/(TP + FP + FN) × 100%",
281:"Recall = TP/(TP + FN) × 100%",284:"Precision = TP/(TP + FP) × 100%",
287:"FAR = FP/(TP + FP) × 100% = 100% − Precision",
289:"F1 = 2 Precision Recall/(Precision + Recall)",
291:"MCC = (TP TN − FP FN)/√[(TP + FP)(TP + FN)(TN + FP)(TN + FN)]",
296:"P₃,₅ = N(ngập và chỉ số ≤3,5 m)/N(ngập hợp lệ) × 100%"}
def math_run(text):
    run=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=text
    rp=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Cambria Math');rf.set(qn('w:hAnsi'),'Cambria Math');rp.append(rf)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24' if MATH_INDEX in (217,218,291) else '20');rp.append(sz);run.append(rp)
    run.append(mt)
    return run

def math_sigma():
    node=OxmlElement('m:sSubSup')
    for tag,text in [('e','σ'),('sub','B'),('sup','2')]:
        child=OxmlElement('m:'+tag);child.append(math_run(text));node.append(child)
    return node

def math_sub(base,subscript):
    node=OxmlElement('m:sSub')
    for tag,text in [('e',base),('sub',subscript)]:
        child=OxmlElement('m:'+tag);child.append(math_run(text));node.append(child)
    return node

def math_fraction(numerator,denominator):
    node=OxmlElement('m:f')
    for tag,content in [('num',numerator),('den',denominator)]:
        child=OxmlElement('m:'+tag)
        child.append(math_run(content) if isinstance(content,str) else content)
        node.append(child)
    return node

def math_sum(lower,upper,body):
    node=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    char=OxmlElement('m:chr');char.set(qn('m:val'),'∑');pr.append(char)
    limit=OxmlElement('m:limLoc');limit.set(qn('m:val'),'undOvr');pr.append(limit);node.append(pr)
    for tag,text in [('sub',lower),('sup',upper),('e',body)]:
        child=OxmlElement('m:'+tag);child.append(math_run(text));node.append(child)
    return node

for index,text in equations.items():
    MATH_INDEX=index
    p=pp[index]; number=re.search(r"\(2\.(\d+)\)",p.text)
    if index==241:number=re.search(r"\(2\.(\d+)\)",'(2.10)')
    p.clear(); math=OxmlElement("m:oMath")
    if index==164:
        math.append(math_run('ω₀(T) = '));math.append(math_sum('i=1','T','pᵢ'));math.append(math_run(' ;  ω₁(T) = 1 − ω₀(T)'))
    elif index==166:
        for j,(lower,upper) in enumerate([('i=1','T'),('i=T+1','L')]):
            math.append(math_run((' ;  ' if j else '')+f'μ{["₀","₁"][j]}(T) = '))
            frac=OxmlElement('m:f');numr=OxmlElement('m:num');numr.append(math_sum(lower,upper,'i pᵢ'));frac.append(numr)
            den=OxmlElement('m:den');den.append(math_run(f'ω{["₀","₁"][j]}(T)'));frac.append(den);math.append(frac)
    elif index in (168,169):
        before,after=text.split('σ_B²');math.append(math_run(before));math.append(math_sigma());math.append(math_run(after))
    elif index==204:
        math.append(math_run('S(U) = [R − I, −gh ∂'));math.append(math_sub('z','b'))
        math.append(math_run('/∂x − '));math.append(math_sub('τ','bx'))
        math.append(math_run('/ρ, −gh ∂'));math.append(math_sub('z','b'))
        math.append(math_run('/∂y − '));math.append(math_sub('τ','by'));math.append(math_run('/ρ]ᵀ'))
    elif index in (217,218):
        axis='x' if index==217 else 'y';velocity='u' if axis=='x' else 'v'
        math.append(math_sub('τ','b'+axis));math.append(math_run('/ρ = '))
        power=OxmlElement('m:sSup');base=OxmlElement('m:e');base.append(math_run('h'));power.append(base)
        exponent=OxmlElement('m:sup');exponent.append(math_fraction('1','3'));power.append(exponent)
        math.append(math_fraction('g n² '+velocity+' √(u² + v²)',power))
    elif index==291:
        math.append(math_run('MCC = '))
        root=OxmlElement('m:rad');pr=OxmlElement('m:radPr');hidden=OxmlElement('m:degHide');hidden.set(qn('m:val'),'1');pr.append(hidden);root.append(pr)
        root.append(OxmlElement('m:deg'));body=OxmlElement('m:e');body.append(math_run('(TP + FP)(TP + FN)(TN + FP)(TN + FN)'));root.append(body)
        math.append(math_fraction('TP TN − FP FN',root))
    else:math.append(math_run(text))
    p._p.append(math)
    if number: p.add_run("    (2."+number.group(1)+")")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
    for r in p.runs: r.font.size=Pt(11)

# Real multilevel chapter numbering drives heading styles and the TOC.
numbering=doc.part.numbering_part.element
abs_ids=[int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
num_ids=[int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]
aid=max(abs_ids+[0])+1;nid=max(num_ids+[0])+1
anum=OxmlElement('w:abstractNum');anum.set(qn('w:abstractNumId'),str(aid))
for level,label in enumerate(['CHƯƠNG %1','%1.%2','%1.%2.%3']):
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),str(level))
    for tag,value in [('start','1'),('numFmt','decimal'),('lvlText',label),('suff','space')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);lvl.append(el)
    rp=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
    rp.append(rf);color=OxmlElement('w:color');color.set(qn('w:val'),'000000');rp.append(color);lvl.append(rp)
    anum.append(lvl)
numbering.append(anum)
num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),str(aid));num.append(ref);numbering.append(num)
def apply_number(p,level):
    pr=p._p.get_or_add_pPr();n=OxmlElement('w:numPr')
    for tag,value in [('ilvl',level),('numId',nid)]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),str(value));n.append(el)
    pr.append(n)
front_h1={7,12,21,27,33,35,52,62,378,397}
if 'Body Subheading' not in doc.styles:
    doc.styles.add_style('Body Subheading',1)
substyle=doc.styles['Body Subheading'];substyle.base_style=doc.styles['Normal']
substyle.font.bold=True;substyle.paragraph_format.keep_with_next=True;substyle.paragraph_format.first_line_indent=Cm(0)
toc_title.style='Title'
for i,p in enumerate(pp):
    if i in removed: continue
    text=p.text.strip()
    if i==3: p.style='Title';continue
    if i in (110,194,306):
        p.text=re.sub(r'^CHƯƠNG\s*\d+\s*:?\s*','',text);p.style='Heading 1';apply_number(p,0);p.paragraph_format.page_break_before=True
    elif 111<=i<=376 and re.match(r'^\d+\.\d+\.\d+\.',text):
        p.text=re.sub(r'^\d+\.\d+\.\d+\.\s*','',text);p.style='Heading 3';apply_number(p,2)
    elif 111<=i<=376 and re.match(r'^\d+\.\d+\.',text):
        p.text=re.sub(r'^\d+\.\d+\.\s*','',text);p.style='Heading 2';apply_number(p,1)
    elif i in front_h1:
        p.style='Heading 1';p.paragraph_format.page_break_before=True
    elif i in {63,72,80,87,93,101,379,390}:
        p.style='Heading 2'
    elif i in {114,116,118,120,332,334,363,372}:
        p.style='Body Subheading'

caption_indices=[269,318,321,330,339,345,352,371,347,349,357]
for i in caption_indices:
    p=pp[i]; text=p.text; match=re.match(r'^(Bảng|Hình)\s*(\d+)\.(\d+)\s*:?\s*(.*)',text)
    label,chapter,count,title=match.groups();p.clear();p.style='Caption'
    p.add_run(label+' '+chapter+'.')
    field(p,'SEQ '+('Bang' if label=='Bảng' else 'Hinh')+(' \\r 1' if count=='1' else '')+' \\* ARABIC',count)
    p.add_run(' '+title)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_with_next=label=='Bảng'
    p.paragraph_format.keep_together=True

# Clean typography while retaining A4 and the source margin pattern.
section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
section.left_margin=Cm(3);section.right_margin=Cm(2);section.top_margin=Cm(2.5);section.bottom_margin=Cm(3)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(13);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.first_line_indent=Cm(.75)
for name,size in [('Title',16),('Heading 1',15),('Heading 2',14),('Heading 3',13),('Caption',11)]:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.font.bold=name!='Caption';style.font.italic=False
    style.paragraph_format.first_line_indent=Cm(0)
    style.paragraph_format.space_before=Pt(12 if name!='Caption' else 6)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.2
    if name.startswith('Heading'):style.paragraph_format.keep_with_next=True;style.paragraph_format.keep_together=True
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    if style.name in {'Normal','Title','Heading 1','Heading 2','Heading 3','Caption','Body Subheading'}:
        for rf in style.element.xpath('.//w:rFonts'):
            for attr in list(rf.attrib):
                if attr.endswith('Theme'):del rf.attrib[attr]
            for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
for p in doc.paragraphs:
    for border in p._p.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    p.paragraph_format.widow_control=True
    if p.style.name=='Normal' and not p._p.xpath('.//w:drawing|.//m:oMath'):
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if re.match(r'^\[\d+\]',p.text):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if p.text.strip().endswith(':'):p.paragraph_format.keep_with_next=True
    for r in p.runs:
        r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
        if p.style.name=='Normal' and not p._p.xpath('.//w:drawing|.//m:oMath'):r.font.size=Pt(13)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
        if p.style.name.startswith('Heading') or p.style.name=='Title':r.font.bold=True;r.font.size=None
    if p.style.name.startswith('Heading') or p.style.name in {'Title','Caption'}:
        for r in p.runs:r.font.size=None
    if p.text.lstrip().startswith(('-', '1.', '2.', '3.', '4.', '5.', '6.')) and p.style.name=='Normal':
        p.paragraph_format.first_line_indent=Cm(0)
for i in [0,2,3,5]:
    pp[i].alignment=WD_ALIGN_PARAGRAPH.CENTER;pp[i].paragraph_format.first_line_indent=Cm(0)
pp[0].paragraph_format.space_before=Pt(12);pp[0].paragraph_format.space_after=Pt(36)
pp[2].paragraph_format.space_after=Pt(30)
pp[3].paragraph_format.space_after=Pt(40);pp[3].paragraph_format.line_spacing=1.35
pp[4].alignment=WD_ALIGN_PARAGRAPH.LEFT;pp[4].paragraph_format.first_line_indent=Cm(0);pp[4].paragraph_format.line_spacing=1.5
pp[5].paragraph_format.space_before=Pt(35)
for i in range(28,32):
    if i not in removed:pp[i].paragraph_format.line_spacing=1.3
for i in range(94,108):
    if i not in removed and pp[i].style.name=='Normal':pp[i].paragraph_format.line_spacing=1.3
for i in [325,326,327]:pp[i].alignment=WD_ALIGN_PARAGRAPH.LEFT
pp[361].text='barrier_mask = (\n    (sar_flood == 1) & plain_mask\n    & (local_relief >= relief_threshold)\n)'
pp[361].alignment=WD_ALIGN_PARAGRAPH.LEFT;pp[361].paragraph_format.first_line_indent=Cm(0)
pp[361].paragraph_format.line_spacing=1.0;pp[361].paragraph_format.keep_together=True
for r in pp[361].runs:r.font.name='Consolas';r.font.size=Pt(10)
for i in [10,19]:pp[i].alignment=WD_ALIGN_PARAGRAPH.RIGHT;pp[i].paragraph_format.first_line_indent=Cm(0)
for i in [6,11,20,26,32,34,51,61,109,193,305,377,396]:
    if i not in removed:drop(i)

def style_table(table,widths):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    pr=table._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+key),value)
        borders.append(e)
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader')
        if ri==0:trpr.append(repeat)
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        for ci,(cell,width) in enumerate(zip(row.cells,widths)):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8EDF2' if ri==0 else 'FFFFFF');tcpr.append(shd)
            margins=OxmlElement('w:tcMar')
            for side,value in [('top','70'),('bottom','70'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.keep_with_next=ri<len(table.rows)-1 if len(table.rows)<=7 else ri<2
                p.paragraph_format.widow_control=True
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 or (len(widths)>=4 and ri>0 and ci<4) else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5 if len(widths)>=5 else 11);r.font.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
widths=[[1.2,2.5,12.3],[5,5.5,5.5],[3.1,5.1,7.8],[2,5.8,8.2],[1.4,2,2.6,2,3.7,4.3],[1.6,3,3,2.8,2.8,2.8],[4,3.2,8.8],[5.4,3.4,3.6,3.6],[2.6,4.5,4.5,4.4]]
for table,w in zip(doc.tables,widths):style_table(table,w)

appendix=doc.add_paragraph('PHỤ LỤC HỒ SƠ DỮ LIỆU VÀ THỰC THI',style='Heading 1');appendix.paragraph_format.page_break_before=True
doc.add_paragraph('Các tệp dưới đây xác định bộ V44 được đối soát. Kết quả của nhánh khác phải có hồ sơ riêng và không thay thế dữ liệu lịch sử này.')
doc.add_paragraph('Danh mục tệp truy vết',style='Heading 2')
files=[EVIDENCE['config'],EVIDENCE['rain']['file']]+[x['file'] for x in EVIDENCE['rivers']]
table=doc.add_table(rows=1,cols=2);new_rows(table,[['Tệp','SHA256 tại đối soát 05/10/2026']]+[[Path(x['path']).name,x['sha256'][:32]+'\n'+x['sha256'][32:]] for x in files]);style_table(table,[6,10])
doc.add_paragraph('Đầu ra: PhuYen_30m_V44.nc; 40 frame; frame 33=367.200,28125 s; frame cuối=432.000 s. Runtime solver theo V44_run_new.log=253.274,39 s. Bộ đối soát lưu trong project_evidence_audit.json. Hash mã hiện lưu không xác nhận tuyệt đối từng byte runtime đã dùng lúc V44 chạy.')
doc.add_paragraph('Quy tắc đánh giá lịch sử',style='Heading 2')
doc.add_paragraph('Ngưỡng mô hình h>0,3 m; SAR được đưa về lưới đầu ra bằng nearest neighbour; nước thường trực và chỉ số địa hình được đưa về lưới theo phép nội suy của mã cũ. Miền hợp lệ: SAR khác NODATA, water density≤0,5 và chỉ số địa hình≥0. Các ô ngoài miền không được tính là TN. Diện tích tính với 900 m²/ô. Transform được giữ theo eval_v44.py để tái lập lịch sử, chưa được chứng nhận là phương pháp căn chỉnh tối ưu.')

# Add page number fields. Word updates these together with TOC/caption lists.
section.different_first_page_header_footer=True
footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Cm(0);field(footer,'PAGE')
for r in footer.runs:r.font.name='Times New Roman';r.font.size=Pt(11)
# Front matter uses roman page numbers; the main text restarts at Arabic 1.
front_section=deepcopy(section._sectPr)
pgtype=OxmlElement('w:pgNumType');pgtype.set(qn('w:fmt'),'lowerRoman');pgtype.set(qn('w:start'),'1');front_section.append(pgtype)
section_type=front_section.find(qn('w:type'))
if section_type is None:section_type=OxmlElement('w:type');front_section.insert(0,section_type)
section_type.set(qn('w:val'),'nextPage')
boundary=pp[62].insert_paragraph_before('');boundary._p.get_or_add_pPr().append(front_section)
boundary.paragraph_format.space_after=Pt(0);boundary.paragraph_format.space_before=Pt(0)
boundary.paragraph_format.line_spacing=1;boundary.paragraph_format.first_line_indent=Cm(0)
for r in boundary.runs:r.font.size=Pt(1)
main_section=doc.sections[-1]._sectPr
for tag in ['w:titlePg','w:pgNumType']:
    old=main_section.find(qn(tag))
    if old is not None:main_section.remove(old)
pgtype=OxmlElement('w:pgNumType');pgtype.set(qn('w:fmt'),'decimal');pgtype.set(qn('w:start'),'1');main_section.append(pgtype)
pp[62].paragraph_format.page_break_before=False
settings=doc.settings.element;uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
doc.core_properties.title='Ứng dụng viễn thám GIS và mô hình thủy lực số để mô phỏng phạm vi ngập tại hạ lưu sông Ba Phú Yên'
doc.core_properties.subject='Kết quả V44 và phương pháp đối chiếu SAR'
doc.core_properties.comments=''
doc.save(DEST)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_HASH
(WORK/'authoring_audit.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':SOURCE_HASH,'output':str(DEST),
 'changed_source_paragraphs':sorted(set(changes)),'removed_source_paragraphs':sorted(removed),'references':{str(i+1):{'old_id':key,'entry':refs[key]} for i,key in enumerate(ordered)},
 'equations':len(equations),'figures':3,'table_count':len(doc.tables)},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(DEST),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'references':len(ordered),'equations':len(equations)},ensure_ascii=False))
