"""Verify the saved Word artifact against the previously extracted evidence."""
from pathlib import Path
import hashlib
import json
import re
import sys
from zipfile import ZipFile
from lxml import etree
from docx import Document
import pypdfium2 as pdfium

sys.stdout.reconfigure(encoding='utf-8')
root=Path(r'D:\VTP')
work=root/'Do_an_edit_20261005'
source=root/'Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx'
output=root/'Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_chinh_sua_20261005.docx'
qa=work/sys.argv[1]
inspected=list(map(int,sys.argv[2].split(','))) if len(sys.argv)>2 else []
evidence=json.loads((root/'Do_an_research_20261005/project_evidence_audit.json').read_text(encoding='utf-8'))
registry=json.loads((root/'Do_an_research_20261005/bibliography_registry_audit.json').read_text(encoding='utf-8'))
doc=Document(output)
with ZipFile(output) as package:
    xml=etree.fromstring(package.read('word/document.xml'))
    styles=etree.fromstring(package.read('word/styles.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
alltext='\n'.join(xml.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
body='\n'.join(p.text for p in doc.paragraphs if not re.match(r'^\[\d+\]',p.text))
refparas=[p.text for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
refnumbers=[int(re.match(r'^\[(\d+)\]',t).group(1)) for t in refparas]
cited=sorted(set(map(int,re.findall(r'\[(\d+)\]',body))))
table=doc.tables[5]
checked=[]
for row,expected in zip(table.rows[1:],evidence['tables_all_40_frames']):
    values=[cell.text for cell in row.cells]
    wanted=[str(expected['frame']),f"{expected['time_s']:.3f}".replace('.',','),
            f"{expected['time_s']/3600:.6f}".replace('.',',')]
    wanted.extend(f'{v:.2f}'.replace('.',',') for v in expected['computed_CSI_Recall_Precision_pct'])
    checked.append({'frame':expected['frame'],'match':values==wanted})
pdf=pdfium.PdfDocument(str(next(qa.glob('*.pdf'))))
pages=[]
for index,page in enumerate(pdf):
    textpage=page.get_textpage()
    txt=textpage.get_text_range()
    lines=[s.strip() for s in txt.splitlines() if s.strip()]
    pages.append({'page':index+1,'characters':len(txt),'first_lines':lines[:3],'last_lines':lines[-3:]})
    textpage.close();page.close()
pdf.close()
(qa/'page_text_inventory.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
wrong_dois=[x['doi'] for x in registry if x['number'] in [3,6,10,14,15,16,24,26] and 'doi' in x]
instructions=xml.xpath('//w:instrText/text()',namespaces=ns)
report={
 'output':str(output),'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
 'original_unchanged':hashlib.sha256(source.read_bytes()).hexdigest()=='9d6d2255d065940c4d43ba4ad82030db877ad82aaa9c7414640ea7978e524d2d',
 'all_40_frames_match_evidence':len(table.rows)==41 and len(checked)==40 and all(x['match'] for x in checked),
 'frame_rows':checked,
 'numbered_references':refnumbers,'body_citations':cited,
 'citations_and_bibliography_complete':refnumbers==list(range(1,24)) and cited==refnumbers,
 'wrong_original_dois_absent':not any(doi in alltext for doi in wrong_dois),
 'native_equations':len(xml.xpath('//m:oMath',namespaces=ns)),
 'chapter_2_equation_numbers_complete':sorted(map(int,re.findall(r'\(2\.(\d+)\)',alltext)))==list(range(1,17)),
 'actual_figures':len(doc.inline_shapes),
 'caption_sequences':{'figures':sum('SEQ Hinh' in t for t in instructions),'tables':sum('SEQ Bang' in t for t in instructions)},
 'automatic_lists':sum(t.strip().startswith('TOC ') for t in instructions),
 'external_reference_hyperlinks':sum(rel.reltype.endswith('/hyperlink') and rel.is_external for rel in doc.part.rels.values()),
 'no_paragraph_rules':len(xml.xpath('//w:pBdr',namespaces=ns))+len(styles.xpath('//w:pBdr',namespaces=ns))==0,
 'no_field_errors':not any(s in alltext for s in ['Error!','Lỗi! Không tìm','Reference source not found','Bookmark not defined']),
 'pages':len(pages),'no_empty_pages':all(p['characters']>20 for p in pages),
 'inspected_pages_at_full_size':inspected,
 'all_final_pages_visually_inspected':inspected==list(range(1,len(pages)+1)),
 'remaining_personal_information':sorted(set(re.findall(r'\[(?:Họ và tên[^\]]*|Mã số sinh viên)\]',alltext))),
 'scope_limitations':['V44 uses the historical alignment and mask; reproducing its metrics does not certify geospatial accuracy.',
                      'V45 and V48 are explicitly dated snapshots; the thesis does not claim completed paired V48 results.',
                      'Student name, student ID and signing date are left for the author to confirm.']
}
checks=['original_unchanged','all_40_frames_match_evidence','citations_and_bibliography_complete',
        'wrong_original_dois_absent','chapter_2_equation_numbers_complete','no_paragraph_rules','no_field_errors','no_empty_pages']
report['structural_checks_passed']=all(report[k] for k in checks) and report['native_equations']==22 and report['actual_figures']==3 and report['automatic_lists']==3
report['status']='PASS' if report['structural_checks_passed'] and report['all_final_pages_visually_inspected'] else 'PENDING_VISUAL_QA' if report['structural_checks_passed'] else 'FAIL'
(work/'final_delivery_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['frame_rows','scope_limitations','inspected_pages_at_full_size']},ensure_ascii=False,indent=2))
