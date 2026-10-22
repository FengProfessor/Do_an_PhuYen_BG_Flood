# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r'D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx')

para_words = sum(len(p.text.split()) for p in doc.paragraphs)
table_words = 0
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            table_words += sum(len(p.text.split()) for p in c.paragraphs)

total_words = para_words + table_words

print('=== FINAL VERIFICATION SUMMARY ===')
print(f'Paragraphs count: {len(doc.paragraphs)}')
print(f'Tables count: {len(doc.tables)}')
print(f'Inline figures: {len(doc.inline_shapes)}')
print(f'Paragraph words: {para_words}')
print(f'Table words: {table_words}')
print(f'Total words (Vietnamese): {total_words}')

# List all tables and captions
print('\n=== TABLE CAPTIONS IN DOCX ===')
for p in doc.paragraphs:
    if p.text.startswith('Bảng'):
        print(' ', p.text)

print('\n=== FIGURE CAPTIONS IN DOCX ===')
for p in doc.paragraphs:
    if p.text.startswith('Hình'):
        print(' ', p.text)

# Check references count
refs = [p.text for p in doc.paragraphs if p.text.startswith('[') and ']' in p.text]
print(f'\nReferences count: {len(refs)}')
