import docx
import zipfile
import xml.etree.ElementTree as ET
import os

docx_path = r"D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx"

if not os.path.exists(docx_path):
    print(f"Error: {docx_path} does not exist!")
    exit(1)

file_size_bytes = os.path.getsize(docx_path)
print(f"File size: {file_size_bytes:,} bytes ({file_size_bytes / (1024*1024):.2f} MB)")

# 1. Zip inspection (media, app.xml)
with zipfile.ZipFile(docx_path, 'r') as z:
    file_list = z.namelist()
    media_files = [f for f in file_list if f.startswith('word/media/')]
    print(f"\nMedia files in docx ({len(media_files)}):")
    for mf in media_files:
        info = z.getinfo(mf)
        print(f"  - {mf}: {info.file_size:,} bytes")
        
    app_xml = None
    if 'docProps/app.xml' in file_list:
        app_xml_data = z.read('docProps/app.xml').decode('utf-8', errors='ignore')
        print(f"\napp.xml content:")
        root = ET.fromstring(app_xml_data)
        for child in root:
            tag = child.tag.split('}')[-1]
            if child.text:
                print(f"  {tag}: {child.text}")

# 2. python-docx inspection
doc = docx.Document(docx_path)

total_paragraphs = len(doc.paragraphs)
total_words = 0
total_chars = 0

all_text = []

for p in doc.paragraphs:
    txt = p.text.strip()
    if txt:
        all_text.append(txt)
        words = txt.split()
        total_words += len(words)
        total_chars += len(txt)

print(f"\nParagraphs: {total_paragraphs}")
print(f"Words in paragraphs: {total_words}")

# Tables
tables = doc.tables
print(f"\nTables count: {len(tables)}")
total_table_words = 0
for idx, tbl in enumerate(tables, 1):
    rows = len(tbl.rows)
    cols = len(tbl.columns) if rows > 0 else 0
    tbl_text = []
    for r in tbl.rows:
        for c in r.cells:
            t = c.text.strip()
            if t:
                tbl_text.append(t)
    # unique cell words or raw split
    # Note: merged cells might repeat, but let's count words in each cell
    tbl_words = sum(len(c.text.split()) for r in tbl.rows for c in r.cells)
    total_table_words += tbl_words
    print(f"  Table {idx}: {rows} rows x {cols} cols, ~{tbl_words} words")

grand_total_words = total_words + total_table_words
print(f"\nTotal words (paragraphs + table cells): {grand_total_words}")

# Section page break estimate
page_breaks = 0
for p in doc.paragraphs:
    for run in p.runs:
        if '<w:br w:type="page"/>' in run._r.xml:
            page_breaks += 1
print(f"Explicit page breaks: {page_breaks}")

# Check COM automation for true Word page count if available
try:
    import win32com.client
    import pythoncom
    print("\nAttempting Word COM automation for page count...")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    doc_com = word.Documents.Open(docx_path, ReadOnly=True)
    wdStatisticPages = 2
    wdStatisticWords = 0
    wdStatisticCharacters = 3
    wdStatisticLines = 1
    wdStatisticParagraphs = 4
    pages = doc_com.ComputeStatistics(wdStatisticPages)
    words_com = doc_com.ComputeStatistics(wdStatisticWords)
    chars_com = doc_com.ComputeStatistics(wdStatisticCharacters)
    lines_com = doc_com.ComputeStatistics(wdStatisticLines)
    paras_com = doc_com.ComputeStatistics(wdStatisticParagraphs)
    doc_com.Close(False)
    word.Quit()
    print(f"Word COM Statistics:")
    print(f"  Pages: {pages}")
    print(f"  Words: {words_com}")
    print(f"  Characters: {chars_com}")
    print(f"  Lines: {lines_com}")
    print(f"  Paragraphs: {paras_com}")
except Exception as e:
    print(f"Word COM automation note: {e}")
