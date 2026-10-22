from docx import Document
from pathlib import Path
import json, re, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
src = Path(r'D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx')
doc = Document(src)
out = []
out.append(f'PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}')
for i,p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        out.append(f'P{i:04d}\t{p.style.name}\t{t}')
for ti,table in enumerate(doc.tables):
    out.append(f'\nTABLE {ti} rows={len(table.rows)} cols={len(table.columns)}')
    for ri,row in enumerate(table.rows):
        vals = [' '.join(c.text.split()) for c in row.cells]
        out.append(f'T{ti}R{ri}\t' + ' | '.join(vals))
Path(r'D:\VTP\_extract.txt').write_text('\n'.join(out), encoding='utf-8')
print('\n'.join(out[:260]))
print(f'\nEXTRACT_SAVED={Path(r"D:\\VTP\\_extract.txt")}')

