# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_PATH = r"D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx"

doc = docx.Document()

# 1. Page Setup A4 & Margins
section = doc.sections[0]
section.page_width = Inches(8.27)   # 21.0 cm
section.page_height = Inches(11.69) # 29.7 cm
section.top_margin = Inches(0.98)    # 2.5 cm
section.bottom_margin = Inches(1.18) # 3.0 cm
section.left_margin = Inches(1.18)   # 3.0 cm (đóng gáy)
section.right_margin = Inches(0.79)  # 2.0 cm

# 2. Styles
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(13)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.25
normal_style.paragraph_format.space_before = Pt(0)
normal_style.paragraph_format.space_after = Pt(6)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="{top}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:left w:val="{left}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:bottom w:val="{bottom}" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:right w:val="{right}" w:sz="4" w:space="0" w:color="B0B0B0"/>
        </w:tcBorders>
    '''.format(
        top=kwargs.get('top', 'single'),
        left=kwargs.get('left', 'none'),
        bottom=kwargs.get('bottom', 'single'),
        right=kwargs.get('right', 'none')
    ))
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def add_h(text, level=1, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
    elif level == 2:
        run.font.size = Pt(13.5)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(13)
        run.italic = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p

def add_p(text, bold_prefix="", indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    return p

def add_caption(text, is_table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    if is_table:
        run.bold = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.italic = True
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)

def add_fig(img_path, caption_text, width_inches=5.8):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_inches))
        add_caption(caption_text, is_table=False)
    else:
        add_caption(f"[{caption_text}]", is_table=False)

def add_formula(text, formula_id):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.italic = True
    run.font.name = 'Times New Roman'
    
    # tab to right for formula_id
    r_id = p.add_run(f"\t\t\t\t\t({formula_id})")
    r_id.bold = True
    r_id.font.name = 'Times New Roman'

print("Base setup completed.")
