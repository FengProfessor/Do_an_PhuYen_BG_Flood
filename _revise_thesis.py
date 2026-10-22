from docx import Document
from pathlib import Path
from copy import deepcopy
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

src = Path(r'D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR.docx')
out = Path(r'D:\VTP\Do_an_tot_nghiep_PhuYen_BG_Flood_SAR_da_kiem_chung.docx')
doc = Document(src)

def setp(i, text, expect=None):
    p = doc.paragraphs[i]
    old = p.text
    if expect and expect.lower() not in old.lower():
        raise RuntimeError(f'P{i:04d} mismatch: expected {expect!r}, got {old[:100]!r}')
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)

# Front matter / abstract
setp(16, 'Em cũng xin gửi lời cảm ơn tới các chuyên gia và các bạn đồng môn tại phòng thí nghiệm đã chia sẻ tài nguyên tính toán GPU, thảo luận chuyên môn và hỗ trợ em kiểm tra, vận hành quy trình mô phỏng BG_Flood.', 'Godunov')
setp(22, 'Lũ lụt là một trong những loại hình thiên tai nguy hiểm và có tần suất cao tại duyên hải miền Trung Việt Nam. Trận lũ cuối tháng 11/2021 tại Phú Yên xảy ra trong bối cảnh mưa lớn diện rộng và vận hành xả lũ lớn của hệ thống hồ chứa; tài liệu chính thức ghi nhận Thủy điện Sông Ba Hạ xả khoảng 9.400 m³/s lúc 15 giờ ngày 30/11/2021. Đồ án xây dựng quy trình mô phỏng ngập 2D trên lưới 30 m bằng BG_Flood/GPU, kết hợp Sentinel-1 SAR và GIS để phân tích trường ngập hạ lưu sông Ba.', 'Lũ lụt')
setp(23, 'Nghiên cứu tập trung vào ảnh hưởng của địa hình raster và các rào cản nhân tạo đối với kết nối thủy lực giữa lòng sông và đồng bằng. Các phiên bản V1–V44 lần lượt thử nghiệm cấu hình nhám, xử lý hố trũng bằng Priority-Flood và hiệu chỉnh cục bộ cao độ tại một số vị trí cản dòng. Phương án san gạt hình thái học (Morphological Carving) được xây dựng cho V45 ở mức thử nghiệm, chưa được xem là kết quả đã kiểm định. Kịch bản V44 hoàn thành chuỗi mô phỏng 120 giờ mà không phát sinh NaN; đây là chỉ báo ổn định tính toán, không đồng nghĩa với việc đã chứng minh độ chính xác vật lý.', 'V1')
setp(24, 'Tại khung thời gian mô hình gần thời điểm quan sát Sentinel-1 lúc khoảng 05:36 ngày 30/11/2021, phép đối chiếu V44 với mặt nạ ngập SAR cho CSI = 20,64%, Recall = 24,97% và Precision = 54,32%. Các chỉ số này được xem là độ khớp với sản phẩm SAR trong cấu hình V44, không phải kiểm định độc lập hoàn toàn vì dữ liệu SAR đã tham gia vào một bước hiệu chỉnh địa hình của kịch bản. Phân tích HAND cho thấy 90,43% diện tích ngập mô phỏng nằm trong vùng HAND ≤ 3,5 m; chỉ số này được sử dụng như kiểm tra tính hợp lý địa hình, không phải ngưỡng chuẩn để khẳng định độ chính xác thủy lực. Kết quả cho thấy quy trình có tiềm năng nhưng vẫn cần bổ sung kiểm định bằng mực nước, vết lũ và dữ liệu độc lập.', 'CSI')
setp(28, 'Flooding is a recurrent hazard in Central Vietnam. The late-November 2021 flood in Phu Yen occurred under widespread heavy rainfall and large reservoir releases; official records report a Song Ba Ha Hydropower release of about 9,400 m³/s at 15:00 on 30 November 2021. This thesis develops a 30 m, GPU-based 2D flood-modelling workflow using BG_Flood together with Sentinel-1 SAR and GIS for the lower Ba River.', 'Flooding')
setp(29, 'The study examines how raster topography and artificial barriers affect hydraulic connectivity between the river channel and floodplain. Model versions V1–V44 tested roughness configuration, Priority-Flood depression treatment and local terrain conditioning. Morphological carving is retained as an exploratory V45 design and is not treated as a validated result. V44 completed the 120-hour run without NaN divergence; this demonstrates computational robustness for that run but does not by itself establish physical accuracy.', 'This study')
setp(30, 'At the model frame nearest the Sentinel-1 observation at about 05:36 local time on 30 November 2021, V44 achieved CSI = 20.64%, Recall = 24.97% and Precision = 54.32% against the SAR-derived flood mask. These values represent agreement with the SAR product in the V44 configuration and should not be interpreted as fully independent validation because SAR information was also used in a terrain-conditioning step. HAND analysis showed that 90.43% of simulated inundation fell within HAND ≤ 3.5 m; this is reported as a topographic-plausibility diagnostic rather than a universal hydraulic-accuracy criterion. Independent water-level and flood-mark observations are still required for stronger validation.', 'Quantitative validation')
setp(39, 'Hình 2.1: Sơ đồ khái quát lưới thể tích hữu hạn (FVM) và thông lượng central-upwind qua mặt ô', 'HLLC')

# Lists/captions
setp(48, 'Hình 3.5: Bản đồ đối chiếu không gian 4 màu giữa mô phỏng V44 và mặt nạ ngập Sentinel-1 SAR tại khung thời gian gần 102 giờ', 'V44')
setp(58, 'Bảng 3.5: Thống kê chỉ số định lượng tại khung thời gian đối chiếu SAR gần 102 giờ', 'Bảng 3.5')
setp(60, '', 'Bảng 3.7')

# Introduction / scope
setp(64, 'Theo WMO [1], giai đoạn 1970–2019 ghi nhận hơn 11.000 thiên tai liên quan đến thời tiết, khí hậu và nước, gây hơn 2 triệu ca tử vong và khoảng 3,64 nghìn tỷ USD thiệt hại kinh tế. Các con số này bao quát nhiều loại hiểm họa, không chỉ riêng lũ lụt. Đối với nghiên cứu ngập lụt, xu thế gia tăng phơi nhiễm, đô thị hóa và biến đổi khí hậu làm tăng nhu cầu về các công cụ mô phỏng và cảnh báo có khả năng lượng hóa bất định.', 'WMO')
setp(67, 'Đợt mưa lũ cuối tháng 11/2021 tại Phú Yên diễn biến rất nhanh. Báo cáo và bản tin địa phương ghi nhận mưa lớn nhiều ngày, với tổng mưa phổ biến từ khoảng 219 đến 698 mm trong giai đoạn 26/11–02/12 ở Phú Yên; tại 15 giờ ngày 30/11, Thủy điện Sông Ba Hạ xả khoảng 9.000–9.400 m³/s và Sông Hinh xả khoảng 2.000 m³/s. Lũ trên sông Ba đạt đỉnh tại Củng Sơn 38,72 m lúc 19 giờ và tại Phú Lâm 4,63 m lúc 23 giờ ngày 30/11/2021 [2], [20]. Những mốc này được dùng để kiểm tra lại cách diễn giải thời gian của ảnh SAR và chuỗi mô phỏng.', 'Trận lũ')
setp(69, 'BG_Flood là mô hình ngập lụt mã nguồn mở chạy trên GPU, được phát triển bởi Cyprien Bosserelle, Emily M. Lane, Alice Harang và cộng sự; kho mã chính thức mô tả lõi phương trình nước nông lấy cảm hứng từ bộ giải Saint-Venant của Basilisk và cấu trúc dữ liệu GPU dạng block/quadtree [3]. Sentinel-1 SAR cung cấp quan sát radar xuyên mây có giá trị để đối chiếu phạm vi ngập, nhưng bản thân sản phẩm SAR vẫn có sai số phân loại và vùng quan sát kém trong đô thị, thảm thực vật hoặc bóng radar [9], [10].', 'BG_Flood')
setp(76, '2. Thiết lập và vận hành mô hình BG_Flood trên GPU; theo dõi ổn định số học, bước thời gian và hiệu suất tính toán trong chuỗi mô phỏng 120 giờ, đồng thời tách riêng tiêu chí “không phát sinh NaN” khỏi đánh giá độ đúng vật lý.', '120 giờ')
setp(77, '3. Xây dựng quy trình xử lý Sentinel-1 SAR để tạo mặt nạ ngập tham chiếu phục vụ đối chiếu không gian. Mặt nạ SAR được xem là dữ liệu quan sát gián tiếp có sai số, không được coi là “ground truth” tuyệt đối.', 'Otsu')
setp(90, '- Quan sát Trái Đất từ không gian (Earth Observation): sử dụng Sentinel-1 SAR như nguồn quan sát độc lập đối với trường thủy lực tại thời điểm chụp. Tuy nhiên, nếu chính mặt nạ SAR tham gia hiệu chỉnh địa hình thì phép so sánh lại với cùng mặt nạ được diễn giải là độ khớp hiệu chỉnh, không phải kiểm định độc lập hoàn toàn.', 'Earth Observation')
setp(92, '- Thực nghiệm số học kế tiếp (Iterative Computational Experimentation): kế thừa và so sánh các kịch bản V1–V44; V45 được giữ ở trạng thái thiết kế/thử nghiệm cho đến khi dữ liệu địa hình công trình, mốc cao độ và nguồn biên được kiểm chứng đầy đủ.', 'V1')
setp(96, '- Phân tích hạn chế của cách chọn ô hạ đê dựa trên mặt nạ SAR trong V44 và đề xuất cách xử lý địa hình cho V45. Các can thiệp địa hình V45 chỉ được coi là giả thuyết thực nghiệm cho đến khi có số liệu hình học đê/cống hoặc mặt cắt độc lập.', 'SAR')
setp(105, 'Chương 2: Phương pháp nghiên cứu và công nghệ áp dụng. Trình bày hệ phương trình nước nông 2D, phương pháp thể tích hữu hạn và thông lượng central-upwind Kurganov–Petrova trong phiên bản mã Python được sử dụng, quy trình xử lý Sentinel-1 SAR, các phép tiền xử lý DEM và hệ thống chỉ số đối chiếu.', 'HLLC')
setp(106, 'Chương 3: Kết quả thực nghiệm và thảo luận. Trình bày cơ sở dữ liệu đầu vào, lịch sử các kịch bản V1–V44, đối chiếu V44 với mặt nạ SAR, phân tích HAND, chẩn đoán rào cản địa hình và giới hạn của phương án V45. Các kết quả GeoAI không thuộc miền Phú Yên được loại khỏi phần kết quả định lượng.', 'V45')
setp(108, 'Phần Tài liệu tham khảo: Danh mục nguồn được rà soát lại theo tiêu đề, tác giả, tạp chí/đơn vị phát hành và DOI hoặc đường dẫn chính thức khi có thể.', 'chuẩn chỉ')

# Literature review
setp(119, 'Vào đầu những năm 2000, cùng với sự gia tăng năng lực tính toán và sự phổ biến của DEM số, các mô hình thủy lực 2D ngày càng được ứng dụng rộng rãi trong mô phỏng ngập. Khi kích thước ô lưới giảm, chi phí tính toán tăng mạnh do số ô tăng và bước thời gian của các lược đồ hiện chịu ràng buộc bởi điều kiện CFL. Vì vậy, hiệu quả tính toán là một trong các lý do thúc đẩy phát triển các bộ giải song song CPU/GPU cho bài toán ngập diện rộng.', '15 đến 30 ngày')
setp(121, 'Kiến trúc GPU cho phép song song hóa nhiều phép tính trên lưới và đã được ứng dụng hiệu quả trong các bộ giải nước nông [24]. Đối với chính bộ dữ liệu của đồ án, nhật ký các kịch bản cho thấy thời gian chạy phụ thuộc mạnh vào độ phân giải, bước thời gian và cấu hình; vì vậy không sử dụng một hệ số tăng tốc GPU/CPU chung cho mọi trường hợp. Hiệu suất của V44 được báo cáo riêng theo thời gian thực thi thực tế của kịch bản.', '50 đến hơn 120 lần')
setp(125, 'Các nghiên cứu gần đây tại Việt Nam đã sử dụng mô hình 1D–2D để mô phỏng ngập hạ lưu sông Ba. Đáng chú ý, Lương Hữu Dũng và cộng sự (2024) [17] ứng dụng HEC-HMS/HEC-RAS, hiệu chỉnh và kiểm định bằng mực nước các trận lũ 2009, 2016, 2021 tại Phú Lâm và kết hợp tài liệu vết lũ năm 2021. Nghiên cứu này cung cấp đối chứng quan trọng cho việc kiểm tra thời gian đỉnh lũ và nhu cầu sử dụng số liệu mực nước/vết lũ độc lập khi đánh giá mô hình ngập tại Phú Yên.', 'Trịnh Thu Phương')
setp(140, 'Thứ ba, vận hành chuỗi hồ chứa có thể làm thay đổi nhanh đường quá trình dòng chảy về hạ du. Trong trận lũ 30/11/2021, lưu lượng xả Sông Ba Hạ tăng mạnh trong ngày và đạt khoảng 9.000–9.400 m³/s vào 15 giờ theo các báo cáo công khai. Việc phân tích đóng góp của hồ chứa đối với ngập hạ du cần dựa trên toàn bộ hydrograph, mưa khu giữa và điều kiện biên hạ lưu, không thể suy ra chỉ từ một giá trị lưu lượng đỉnh.', '9.000')

# BG_Flood numerical method
setp(190, 'BG_Flood là mô hình mã nguồn mở mô phỏng thủy động lực nước nông trên GPU. Tài liệu chính thức của dự án ghi nhận Cyprien Bosserelle, Emily M. Lane và Alice Harang là các tác giả của công bố BG-Flood năm 2021 [3]. Kho mã chính thức mô tả cấu trúc block/quadtree trên GPU và lõi SWE có nguồn gốc/cảm hứng từ bộ giải Saint-Venant của Basilisk.', 'Bosserelle')
setp(191, 'Trong phiên bản Python/GPU lưu cùng đồ án, các hàm thông lượng được đặt tên `kurganov_x_vec` và `kurganov_y_vec`, tức sử dụng dạng central-upwind Kurganov–Petrova thay vì bộ giải Riemann HLLC. Vì vậy bản hiệu đính mô tả phương pháp theo mã thực tế được sử dụng, không gán HLLC/MUSCL nếu chưa có bằng chứng từ phiên bản mã chạy mô phỏng [18].', 'HLLC' if 'HLLC' in doc.paragraphs[191].text else 'Godunov')
setp(220, '2.1.3. Phương pháp thể tích hữu hạn và thông lượng central-upwind Kurganov–Petrova', 'HLLC')
setp(222, 'Miền tính toán được rời rạc hóa theo thể tích hữu hạn. Trong phiên bản mã Python/GPU của đồ án, thông lượng theo hai phương được tính bằng các hàm Kurganov central-upwind và cập nhật từ phân kỳ thông lượng. Họ lược đồ central-upwind Kurganov–Petrova được phát triển cho hệ Saint-Venant với mục tiêu duy trì cân bằng và tính không âm trong các điều kiện phù hợp [18]. Các tính chất hình thức của bài báo gốc không được mặc nhiên xem là đã chứng minh cho mọi sửa đổi trong fork cục bộ.', 'MUSCL')
setp(224, 'Tại mỗi mặt ô, mã thực tế tái dựng trạng thái trái/phải và ước lượng tốc độ truyền sóng để tính thông lượng central-upwind. Do đó, các mô tả trước đây về cấu trúc ba sóng HLLC và “khôi phục sóng tiếp xúc” được loại khỏi bản hiệu đính vì không phù hợp với `flow_gpu.py` đang lưu trong dự án. Độ ổn định và bảo toàn của phiên bản fork cần được đánh giá bằng các bài kiểm thử cân bằng khối lượng, mặt nước tĩnh và lan truyền ướt/khô.', 'HLLC')
setp(233, '- Tham số `eps` mặc định trong `param.py` là 0,0001 m; đồng thời một số bước cập nhật trong `flow_gpu.py` sử dụng ngưỡng 0,001 m để đặt vận tốc bằng 0 và xác định ô đủ ướt cho bước thời gian.', '0.001')
setp(234, '- Trong bước cập nhật, các ô có h < 0,001 m được đặt u = v = 0 nhằm hạn chế vận tốc giả ở lớp nước rất mỏng. Phần ma sát của fork còn sử dụng `h_fric_min = 0,01 m`; đây là lựa chọn triển khai cần được ghi rõ khi đánh giá độ nhạy.', 'hu = 0')
setp(235, '- Tính chất well-balanced/positivity-preserving là đặc trưng của một số lược đồ Kurganov–Petrova trong tài liệu [18], nhưng với fork cục bộ có thêm clamp và ngưỡng ướt/khô, bản hiệu đính không tuyên bố “cân bằng chính xác tới sai số máy” nếu chưa có bài kiểm thử riêng chứng minh.', 'Well-Balanced')

# SAR / DEM methods
setp(237, 'Đồ án xây dựng mặt nạ ngập tham chiếu từ Sentinel-1 SAR để đối chiếu không gian với mô hình. SAR có ưu điểm quan sát xuyên mây nhưng vẫn có thể bỏ sót/ngộ nhận nước do đô thị, thảm thực vật, nhiễu speckle, góc tới và bóng radar; vì vậy thuật ngữ “ground truth” được thay bằng “SAR-derived reference flood mask” [10], [27], [28].', 'Ground-truth')
setp(238, 'Giai đoạn 1: Thu nhận và kiểm tra siêu dữ liệu. Hồ sơ catalog hiện tìm thấy cảnh Sentinel-1A ứng viên chụp lúc 22:36:38 UTC ngày 29/11/2021, tương ứng 05:36:38 ICT ngày 30/11/2021. Raster SAR nội bộ có kích thước 1.911 × 2.402 ô và độ phân giải khoảng 90 m. Cần lưu scene ID và chuỗi xử lý từ sản phẩm gốc đến raster phân loại để khóa hoàn toàn provenance; thời điểm 05:36 không được gọi là đỉnh lũ hạ lưu vì Phú Lâm đạt đỉnh lúc 23 giờ cùng ngày [20].', '05:36')
setp(250, 'FABDEM 30 m có thể còn các depression/pit và sai số địa hình cục bộ. Trong các kịch bản 30 m của dự án, lỗi NaN xuất hiện gần t ≈ 160.000 s; sau khi áp dụng Priority-Flood để xử lý depression, các lần chạy tiếp theo hoàn thành chuỗi dài hơn. Quan sát này cho thấy mối liên hệ thực nghiệm giữa tiền xử lý DEM và ổn định số học, nhưng không đủ để chứng minh depression là nguyên nhân duy nhất của mọi NaN.', 'NaN')
setp(260, '- Bước 3: V44 hạ cao độ cục bộ một số ô được thuật toán nhận diện để tạo kết nối qua rào cản raster. Vì mặt nạ SAR tham gia điều kiện chọn ô, phép chỉnh này là một bước hiệu chỉnh hướng theo quan sát; không được diễn giải như vị trí cống thực địa nếu chưa có dữ liệu công trình độc lập.', '2.0m')
setp(262, 'Phân tích V44 cho thấy điều kiện `sar_flood == 1` có thể bỏ qua đỉnh đê khô dù vùng phía sau đê xuất hiện ngập trên SAR. Vì vậy V45 đề xuất thử nghiệm morphological carving không phụ thuộc trực tiếp vào việc pixel đỉnh đê được SAR phân loại là nước. Đây là giả thuyết hiệu chỉnh địa hình, cần ràng buộc bằng mặt cắt, cao trình đê/cống và hệ cao độ thống nhất trước khi dùng cho mô phỏng khoa học.', 'V44')
setp(297, 'HAND được sử dụng ở đây như chỉ báo mô tả quan hệ giữa trường ngập và địa hình tương đối so với mạng thoát nước. Không có ngưỡng phổ quát “P_HAND ≥ 85–90%” trong tài liệu gốc [6] để kết luận một mô hình thủy lực đạt hay không đạt. Vì vậy tỷ lệ P_HAND chỉ được báo cáo như phép kiểm tra tính hợp lý địa hình bổ sung và phải được xem cùng mực nước, vết lũ, SAR và cân bằng nước.', '85 - 90')
setp(300, '1. Dữ liệu đầu vào: Sentinel-1 SAR; DEM FABDEM 30 m; ESA WorldCover 10 m; chuỗi mưa ERA5/ERA5-Land đã qua tiền xử lý; chuỗi lưu lượng tại các biên sông và bản đồ nhám/thấm. Hồ sơ kiểm toán cho thấy provenance của một số chuỗi mưa và lưu lượng (bao gồm hệ số biến đổi nội bộ) chưa được khóa đầy đủ, vì vậy các chuỗi này được xem là đầu vào mô hình cần tiếp tục xác minh.', '145')

# Chapter 3 flood chronology/data
setp(312, 'Đợt mưa lũ từ cuối tháng 11 đến đầu tháng 12/2021 là một trong những trận lũ lớn gần đây tại Phú Yên. Để tránh nhầm giữa thời điểm ảnh vệ tinh và đỉnh lũ thủy văn, bản hiệu đính sử dụng các mốc quan trắc công khai tại Củng Sơn và Phú Lâm làm đối chứng thời gian [2], [20].', '26 tháng 11')
setp(314, 'Các nguồn công khai ghi nhận lưu lượng xả Sông Ba Hạ tăng nhanh trong ngày 30/11 và đạt khoảng 9.000–9.400 m³/s vào 15 giờ; Sông Hinh cùng thời điểm ở khoảng 2.000 m³/s. Các con số 3.200 m³/s cho Sông Hinh và tổng >12.600 m³/s trong bản cũ không được giữ như số liệu quan trắc vì chưa tìm thấy nguồn gốc chính thức tương ứng. Khi mô hình sử dụng chuỗi forcing khác các mốc công khai, cần ghi rõ đó là chuỗi đầu vào đã xử lý và nêu provenance của phép biến đổi.', '12.600')
setp(315, 'Theo Ban Chỉ huy PCTT-TKCN tỉnh Phú Yên, đỉnh lũ sông Ba tại Củng Sơn đạt 38,72 m lúc 19 giờ ngày 30/11/2021; tại Phú Lâm đạt 4,63 m lúc 23 giờ cùng ngày [20]. Các giá trị 34,50 m và 3,75 m trong bản cũ không phải đỉnh lũ của sự kiện và đã được sửa. Điều này cũng cho thấy ảnh Sentinel-1 lúc khoảng 05:36 ngày 30/11 được chụp trước đỉnh Phú Lâm khoảng 17 giờ.', '+34.50')
setp(326, '- Nhánh sông Bàn Thạch/Sông Hinh (River 2, vị trí x = 285.200, y = 1.425.000): cấu hình mô hình hiện có chứa chuỗi forcing với đỉnh xấp xỉ 3.200 m³/s. Giá trị này được giữ như thông tin về tệp đầu vào mô hình, không coi là số đo quan trắc đã xác minh; bản tin công khai tại 15 giờ ngày 30/11 ghi khoảng 2.000 m³/s cho Sông Hinh.', '3.200')
setp(335, 'Khi chuyển sang lưới 30 m, một số kịch bản bị dừng gần t ≈ 160.000 s do NaN. Khảo sát địa hình cho thấy các depression cục bộ có thể tạo gradient lớn và bước thời gian rất nhỏ. Sau khi áp dụng Priority-Flood [14] ở V41, chuỗi mô phỏng hoàn thành 120 giờ trong cấu hình thử nghiệm. Kết quả này hỗ trợ giả thuyết rằng tiền xử lý depression cải thiện độ ổn định, nhưng không chứng minh đây là nguyên nhân duy nhất; các clamp vận tốc/độ sâu và lựa chọn bước thời gian trong fork cũng cần được xem xét.', 'triệt để')
setp(337, 'Kịch bản V44 kết hợp DEM đã xử lý depression, các chuỗi inflow được cấu hình và thuật toán hạ cục bộ một số rào cản địa hình. Mô hình hoàn thành 120 giờ. Tuy nhiên, hồ sơ provenance của một số chuỗi forcing vẫn cần bổ sung; vì vậy bản hiệu đính không gọi toàn bộ các chuỗi này là “nguồn xả lũ thực tế” nếu chưa có tài liệu gốc đi kèm.', 'đầy đủ 3 nguồn')
setp(339, 'Bảng 3.4: Diễn biến các chỉ số đối chiếu V44 với cùng mặt nạ SAR theo các khung thời gian mô phỏng', 'Bảng 3.4')
setp(343, '3. Giai đoạn ngập rộng đồng bằng (40h–120h): CSI tăng dần và đạt 20,64% tại khung 102h; sau đó tiếp tục tăng nhẹ trong bảng tính. Khung 102h được chọn vì gần thời điểm ảnh SAR theo giả định thời gian mô phỏng hiện tại, không phải vì đây là đỉnh lũ quan trắc. Đỉnh Phú Lâm thực tế xảy ra lúc 23 giờ ngày 30/11/2021 [20].', 'đạt đỉnh lũ')
setp(344, 'Precision trong Bảng 3.4 dao động khoảng 52–63% và tại khung 102h đạt 54,32%. Con số này có nghĩa là 54,32% các ô được mô hình phân loại ngập cũng nằm trong mặt nạ ngập SAR tại phép đối chiếu đó; nó không phải “độ chính xác tổng thể” của mô hình và chịu ảnh hưởng bởi sai số của chính mặt nạ SAR.', 'Độ chính xác Precision')
setp(345, 'Bảng 3.5: Thống kê chi tiết tại khung thời gian đối chiếu SAR gần 102 giờ (Frame 33)', 'đỉnh lũ')
setp(347, 'Hình 3.1: Bản đồ đối chiếu không gian 4 màu giữa V44 và mặt nạ Sentinel-1 SAR tại khung gần 102 giờ (TP: xanh lá; FP: đỏ; FN: xanh dương; TN: xám)', '102')
setp(353, 'Phân tích Bảng 3.6 cho thấy 90,43% diện tích ngập V44 nằm trong vùng HAND ≤ 3,5 m và khoảng 1,00% nằm ở HAND > 10 m. Kết quả này cho thấy trường ngập tập trung chủ yếu ở địa hình thấp theo định nghĩa HAND, nhưng không phải là kiểm định độc lập về mực nước hay độ sâu. Ngưỡng 3,5 m là lựa chọn phân lớp của nghiên cứu và cần được báo cáo như một ngưỡng phân tích, không phải chuẩn chấp nhận phổ quát.', '90.43')
setp(354, '3.6. Chẩn đoán rào cản địa hình hạ lưu sông Ba và thiết kế thử nghiệm V45', 'V45')
setp(367, 'Tệp `PhuYen_DEM_30m_V45.asc` hiện chỉ là DEM thử nghiệm. Audit cục bộ ghi nhận 41.907 ô bị thay đổi và mức hạ lớn nhất tới khoảng 171,67 m, trong khi chưa có bộ dữ liệu hình học đê/cống độc lập đủ để chứng minh các thay đổi này là thực địa. Các chạy V45 hiện mới phục vụ kiểm tra kỹ thuật ngắn; chưa có mô phỏng V45 120 giờ và chưa có CSI V45 được kiểm định. Vì vậy V45 không được dùng như bằng chứng rằng độ phủ ngập đã “cải thiện đáng kể”.', 'Kết quả mô phỏng thử nghiệm')

# GeoAI section: keep as future work only; remove unsupported Phu Yen metrics/table
setp(368, '3.7. Định hướng mô hình học máy thay thế (Surrogate Model)', '3.7')
setp(369, 'Kho mã có các script và artifact XGBoost/ANN/LSTM lịch sử. Tuy nhiên, dataset ML hiện lưu trong `results_ml_BG` có hệ tọa độ/miền không gian khác khu vực Phú Yên và README của module mô tả bài toán tại Huế. Vì vậy các metric của bộ artifact này không được sử dụng làm kết quả Phú Yên trong bản hiệu đính.', 'Surrogate Model')
setp(370, 'Để xây dựng surrogate cho Phú Yên, cần sinh bộ dữ liệu từ các kịch bản BG_Flood đã khóa provenance; chia tập huấn luyện/kiểm tra theo không gian và theo sự kiện; giữ riêng các kịch bản chưa thấy để kiểm định; và báo cáo R², MAE/RMSE cùng thời gian suy luận trên phần cứng xác định. Việc chia ngẫu nhiên theo từng ô lưới có thể làm kết quả lạc quan do tự tương quan không gian.', '4,4 triệu')
setp(371, '', 'Bảng 3.7')
setp(372, 'Các yêu cầu phương pháp chính đối với nghiên cứu surrogate trong tương lai gồm:', 'Đánh giá khoa học')
setp(373, '1. Thời gian suy luận chỉ được báo cáo sau khi đo trực tiếp trên cùng phần cứng và cùng kích thước miền; không suy rộng các con số latency từ artifact thuộc miền khác.', 'độ trễ')
setp(374, '2. Tự tương quan không gian có thể gây rò rỉ giữa tập train/test nếu chia ngẫu nhiên theo ô. Nên dùng spatial block cross-validation và, nếu có thể, tách theo trận lũ hoặc kịch bản thủy văn để đánh giá khả năng tổng quát hóa.', 'Spatial Autocorrelation')
setp(375, '3. Mô hình ML thuần túy không tự bảo đảm bảo toàn khối lượng/động lượng và có thể suy giảm mạnh ngoài phân bố huấn luyện. Cần áp đặt kiểm tra vật lý sau dự báo và đánh giá riêng các kịch bản cực đoan hoặc thay đổi hình học công trình.', 'Thiếu tính bảo toàn')
setp(376, 'Trong phạm vi đồ án này, GeoAI được giữ ở mức định hướng phát triển; các chỉ số định lượng không thuộc miền Phú Yên đã được loại khỏi phần kết quả.', 'Do đó')

# Conclusions
setp(381, '1. Xây dựng được quy trình tích hợp mô hình thủy lực 2D chạy trên GPU, Sentinel-1 SAR và GIS trên lưới 30 m hơn 4,4 triệu ô. Kịch bản V44 hoàn thành 120 giờ và cho phép xuất trường ngập để đối chiếu không gian. Hiệu suất tính toán được báo cáo theo nhật ký chạy của từng kịch bản; không dùng hiệu suất như bằng chứng thay thế cho kiểm định vật lý.', '1.83')
setp(384, '- Sau khi áp dụng Priority-Flood ở V41, mô hình hoàn thành chuỗi 120 giờ trong cấu hình thử nghiệm. Kết quả này cho thấy tiền xử lý depression có liên quan đến việc cải thiện ổn định số học, nhưng các sửa đổi khác trong fork và điều kiện CFL cũng có thể ảnh hưởng.', 'triệt để')
setp(387, '- Tại khung thời gian gần quan sát Sentinel-1, V44 đạt CSI 20,64%, Recall 24,97% và Precision 54,32% so với mặt nạ SAR. Đây là độ khớp với sản phẩm SAR trong cấu hình đã hiệu chỉnh, không phải kiểm định độc lập hoàn toàn.', 'đỉnh lũ')
setp(388, '4. Phân tích HAND cho thấy 90,43% diện tích ngập V44 nằm trong vùng HAND ≤ 3,5 m. Kết quả này được dùng như chỉ báo tính hợp lý địa hình bổ sung; không có ngưỡng chuẩn 85–90% để kết luận độ chính xác thủy lực chỉ từ HAND.', 'Kiểm chứng độc lập')
setp(389, '5. Xác định hạn chế logic của V44 khi chỉ hạ các ô cản dòng thỏa `sar_flood == 1`. Một DEM V45 thử nghiệm đã được tạo để nghiên cứu kết nối qua rào cản, nhưng audit cho thấy thay đổi địa hình lớn và thiếu hình học công trình độc lập; do đó V45 chưa được coi là kết quả đã kiểm định và chưa được dùng để tuyên bố cải thiện CSI.', '41.907')
setp(392, '1. Về dữ liệu địa hình: trước khi chạy V45 dài hạn, cần thống nhất datum cao độ giữa DEM, trạm mực nước và biên biển; đối chiếu đê/cống bằng LiDAR, UAV, mặt cắt hoặc hồ sơ công trình; chỉ hạ địa hình tại vị trí có bằng chứng. DEM V45 hiện tại nên được xem là nhánh thử nghiệm, không phải DEM chuẩn.', 'PhuYen_DEM_30m_V45.asc')

# Bibliography corrections
setp(399, '[2] Ban Chỉ đạo Quốc gia về Phòng, chống thiên tai (2022). Thiên tai Việt Nam 2021. Báo cáo tổng hợp; xem các trang về đợt mưa lũ miền Trung cuối tháng 11/2021 và vận hành hồ chứa. [Online]. https://chinhsach.dmptc.gov.vn/Uploads/Thu%20vien%20tai%20lieu/Tai%20lieu%20truyen%20thong/Thien%20tai%20Viet%20Nam%202021%20fix%203.pdf', '[2]')
setp(400, '[3] Bosserelle, C., Lane, E. M., & Harang, A. (2021). BG-Flood: A GPU adaptive, open-source, general inundation hazard model. Proceedings of the Australasian Coasts & Ports 2021 Conference. Project repository and citation information: https://github.com/CyprienBosserelle/BG_Flood', '[3]')
setp(403, '[6] Nobre, A. D., Cuartas, L. A., Hodnett, M., Rennó, C. D., Rodrigues, G., Silveira, A., Waterloo, M., & Saleska, S. (2011). Height Above the Nearest Drainage – a hydrologically relevant new terrain model. Journal of Hydrology, 404(1–2), 13–29. https://doi.org/10.1016/j.jhydrol.2011.03.051', '[6]')
setp(405, '[8] Hawker, L., Uhe, P., Paulo, L., Sosa, J., Savage, J., Sampson, C., & Neal, J. (2022). A 30 m global map of elevation with forests and buildings removed. Environmental Research Letters, 17(2), 024016. https://doi.org/10.1088/1748-9326/ac4d4f', '[8]')
setp(407, '[10] Schumann, G. J.-P., & Moller, D. K. (2015). Microwave remote sensing of flood inundation. Physics and Chemistry of the Earth, Parts A/B/C, 83–84, 84–95. https://doi.org/10.1016/j.pce.2015.05.002', '[10]')
setp(409, '[12] Toro, E. F. (2001). Shock-Capturing Methods for Free-Surface Shallow Flows. John Wiley & Sons. ISBN 978-0-471-98766-6.', '[12]')
setp(411, '[14] Barnes, R., Lehman, C., & Mulla, D. J. (2014). Priority-flood: An optimal depression-filling and watershed-labeling algorithm for digital elevation models. Computers & Geosciences, 62, 117–127. https://doi.org/10.1016/j.cageo.2013.04.024', '[14]')
setp(412, '[15] Callow, J. N., Van Niel, K. P., & Boggs, G. S. (2007). How does modifying a DEM to reflect known hydrology affect subsequent terrain analysis? Journal of Hydrology, 332(1–2), 30–39. https://doi.org/10.1016/j.jhydrol.2006.06.020', '[15]')
setp(413, '[16] Saksena, S., & Merwade, V. (2015). Incorporating the effect of DEM resolution and accuracy for improved flood inundation mapping. Journal of Hydrology, 530, 180–194. https://doi.org/10.1016/j.jhydrol.2015.09.069', '[16]')
setp(414, '[17] Lương Hữu Dũng, Lê Quốc Huy, Phan Văn Thành, Văn Thị Hằng, Trần Thanh Thủy, & Nguyễn Văn Đại (2024). Ứng dụng mô hình HEC-RAS mô phỏng ngập lụt cho hạ lưu sông Ba. Tạp chí Khoa học Biến đổi khí hậu, 30, 65–76. https://doi.org/10.55659/2525-2496/30.99719', '[17]')
setp(415, '[18] Kurganov, A., & Petrova, G. (2007). A second-order well-balanced positivity preserving central-upwind scheme for the Saint-Venant system. Communications in Mathematical Sciences, 5(1), 133–160. https://doi.org/10.4310/CMS.2007.v5.n1.a6', '[18]')
setp(416, '[19] European Space Agency / Copernicus (2021–2026). Sentinel-1 mission and product documentation, accessed through the Copernicus Data Space Ecosystem. https://dataspace.copernicus.eu/', '[19]')
setp(417, '[20] Ban Chỉ huy PCTT-TKCN tỉnh Phú Yên (2021). Mưa ngớt, lũ trên các sông đang giảm: số liệu đỉnh lũ Củng Sơn 38,72 m lúc 19 giờ và Phú Lâm 4,63 m lúc 23 giờ ngày 30/11/2021. Phú Yên Online, 01/12/2021. https://phuyen.baodaklak.vn/141/268069/mua-ngot-lu-tren-cac-song-dang-giam.html', '[20]')
setp(421, '[24] Vacondio, R., Dal Palù, A., & Mignosa, P. (2014). GPU-enhanced finite volume shallow water solver for fast flood simulations. Environmental Modelling & Software, 57, 60–75. https://doi.org/10.1016/j.envsoft.2014.02.003', '[24]')
setp(423, '[26] Sanders, B. F. (2007). Evaluation of on-line DEM products for flood inundation modeling. Advances in Water Resources, 30(8), 1831–1843. https://doi.org/10.1016/j.advwatres.2007.02.005', '[26]')
setp(424, '[27] Pulvirenti, L., Pierdicca, N., Chini, M., & Guerriero, L. (2011). An algorithm for operational flood mapping from Synthetic Aperture Radar (SAR) data using fuzzy logic. Natural Hazards and Earth System Sciences, 11, 529–540. https://doi.org/10.5194/nhess-11-529-2011', '[27]')
setp(429, '[32] Cunge, J. A., Holly, F. M., & Verwey, A. (1980). Practical Aspects of Computational River Hydraulics. Pitman Publishing, London. ISBN 0-273-08442-9.', '[32]')
setp(432, '[35] Copernicus Data Space Ecosystem (2026). Sentinel-1 data access and mission products. https://dataspace.copernicus.eu/', '[35]')

# Update table labels/contents that overstate validation or numerical method
# Table 2 = BG_Flood config table (0-based index 2 from extraction)
for table in doc.tables:
    for row in table.rows:
        vals = [c.text.strip() for c in row.cells]
        if vals and vals[0] == 'Phương pháp giải số':
            row.cells[1].text = 'Finite Volume / central-upwind (fork cục bộ)'
            row.cells[2].text = 'Thông lượng Kurganov–Petrova theo `flow_gpu.py`; không mô tả là HLLC nếu chưa có bằng chứng từ mã chạy'
        if len(vals) > 1 and vals[1] == 'HLL / HLLC':
            row.cells[1].text = 'CU'
            row.cells[2].text = 'Central-Upwind scheme (Lược đồ thông lượng central-upwind)'
        if vals and vals[0] == 'Ngưỡng ranh giới ướt/khô':
            row.cells[1].text = 'eps mặc định = 0,0001 m; ngưỡng vận tốc = 0,001 m'
            row.cells[2].text = 'Fork cục bộ còn dùng h_fric_min = 0,01 m trong tính ma sát'
        if vals and vals[0] == '33 (Đỉnh lũ SAR)':
            row.cells[0].text = '33 (Khung đối chiếu SAR)'
        if vals and vals[0] == 'Critical Success Index (CSI)':
            row.cells[2].text = 'Độ khớp với mặt nạ SAR trong cấu hình V44; không phải kiểm định độc lập hoàn toàn'
        if vals and vals[0] == 'F1-Score / Cohen\'s Kappa / MCC':
            row.cells[2].text = 'Các chỉ số mô tả độ khớp phân loại; không gán mức “khá” nếu không nêu quy ước đánh giá'

# Remove unsupported ML result table (last table in current document, identified by first cell)
for table in list(doc.tables):
    if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip().startswith('Phương pháp / Mô hình'):
        tbl = table._element
        tbl.getparent().remove(tbl)
        break

# Core properties
props = doc.core_properties
props.comments = 'Bản hiệu đính 30/09/2026: rà soát nguồn, số liệu, phương pháp số, SAR/HAND, V45 và GeoAI; loại các khẳng định không đủ bằng chứng.'

doc.save(out)
print(f'SAVED {out} size={out.stat().st_size}')
print('tables_after=', len(Document(out).tables), 'paras=', len(Document(out).paragraphs))
